#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Motor Universal de Geração do Cronograma Físico-Financeiro, Linha de Base (Baseline 01),
Planilha Executiva Excel (.xlsx), Exportação MS Project (.xml) e Dashboard HTML (.html).

Uso:
    python scripts/gerar_cronograma.py --obra OBRA_TMULT
    python scripts/gerar_cronograma.py --obra RESIDENCIAL_ALPHA
    python scripts/gerar_cronograma.py --dir projetos/OBRA_TMULT
"""

import os
import sys
import json
import argparse
import datetime
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment
import plotly.graph_objects as go
from plotly.subplots import make_subplots

# Raiz do repositório
ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

from scripts.common.obra_io import resolver_obra_dir, carregar_config_obra
from scripts.common.excel_theme import (
    NAVY, BLUE_DARK, GREEN_FILL, GRAY_LIGHT,
    FONT_TITLE, FONT_HEADER, FONT_BOLD, FONT_REGULAR,
    THIN_BORDER, DOUBLE_BOTTOM_BORDER,
    ALIGN_CENTER, ALIGN_LEFT,
)
from scripts.common.eap import determinar_distribuicao_canonica, classificar_codigo_eap
from scripts.common.cpm import calcular_cpm_detalhado


if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass


def parse_val(v):
    if isinstance(v, str):
        v = v.replace('R$', '').strip().replace('.', '').replace(',', '.')
    return float(v)


def parse_args():
    parser = argparse.ArgumentParser(description="Motor Universal de Cronograma Físico-Financeiro.")
    parser.add_argument("--obra", type=str, default="OBRA_TMULT", help="Nome da pasta da obra em projetos/")
    parser.add_argument("--dir", type=str, default=None, help="Caminho completo da pasta da obra")
    return parser.parse_args()


def carregar_e_distribuir_orcamento(orcamento_csv, prazo_meses=6, regras_custom=None):
    df = pd.read_csv(orcamento_csv, sep=';', encoding='utf-8-sig')

    eap_col = [c for c in df.columns if 'digo EAP' in c or 'EAP' in c][0]
    desc_col = [c for c in df.columns if 'Descricao' in c or 'Item' in c][0]
    df['Preco_Total'] = df['Custo Total (R$)'].apply(parse_val)
    df['Preco_Unitario'] = df['Preço Unitário (R$)'].apply(parse_val)

    for m in range(1, prazo_meses + 1):
        df[f'P_M{m}'] = 0.0

    for idx, r in df.iterrows():
        eap = str(r[eap_col]).strip()
        desc = str(r[desc_col]).strip()

        # Custom rules override if provided
        if regras_custom and eap in regras_custom:
            pesos = regras_custom[eap]
            for m in range(1, prazo_meses + 1):
                df.at[idx, f'P_M{m}'] = float(pesos.get(f'P_M{m}', 0.0))
        else:
            pesos, aviso = determinar_distribuicao_canonica(eap, prazo_meses=prazo_meses, descricao=desc)
            if aviso:
                print(f"[AVISO EAP] Linha {idx}: {aviso}")
            for m in range(1, prazo_meses + 1):
                df.at[idx, f'P_M{m}'] = float(pesos.get(f'P_M{m}', 0.0))

    # Calcular valores monetários
    for m in range(1, prazo_meses + 1):
        df[f'V_M{m}'] = (df['Preco_Total'] * df[f'P_M{m}']).round(2)

    # Ajuste de arredondamento para garantir fechamento exato
    for idx, r in df.iterrows():
        soma_v = sum(r[f'V_M{m}'] for m in range(1, prazo_meses + 1))
        diff = round(r['Preco_Total'] - soma_v, 2)
        if abs(diff) > 0.0001:
            for m in range(prazo_meses, 0, -1):
                if df.at[idx, f'P_M{m}'] > 0:
                    df.at[idx, f'V_M{m}'] = round(df.at[idx, f'V_M{m}'] + diff, 2)
                    break

    return df, eap_col, desc_col


def gerar_csv(df, eap_col, desc_col, output_dir, sigla, prazo_meses=6):
    out_csv = os.path.join(output_dir, f"CRONOGRAMA_FISICO_FINANCEIRO_{sigla}.csv")

    cols_export = [eap_col, desc_col, 'Disciplina', 'Qtd Projeto', 'Unidade Proj', 'Preço Unitário (R$)', 'Custo Total (R$)']
    for m in range(1, prazo_meses + 1):
        cols_export.extend([f'P_M{m}', f'V_M{m}'])

    df_export = df[cols_export].copy()

    header_names = ['Código EAP', 'Item / Descrição', 'Disciplina', 'Qtd Projeto', 'Unidade', 'Preço Unitário (R$)', 'Preço Total Turnkey (R$)']
    for m in range(1, prazo_meses + 1):
        header_names.extend([f'% Mês {m}', f'R$ Mês {m}'])

    df_export.columns = header_names
    df_export.to_csv(out_csv, sep=';', index=False, encoding='utf-8-sig')
    print(f"-> CSV gerado: {out_csv}")
    return out_csv


def gerar_excel(df, eap_col, desc_col, output_dir, sigla, titulo_obra, prazo_meses=6, fisico_acum=None):
    out_xlsx = os.path.join(output_dir, f"CRONOGRAMA_FISICO_FINANCEIRO_{sigla}.xlsx")
    wb = openpyxl.Workbook()

    if fisico_acum is None or len(fisico_acum) < prazo_meses:
        fisico_acum = [round(100.0 * (m / prazo_meses), 2) for m in range(1, prazo_meses + 1)]

    # Aba 1: Resumo Executivo & Curva S
    ws1 = wb.active
    ws1.title = "Resumo Executivo & Curva S"
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:G2")
    ws1["A1"] = f"{titulo_obra.upper()} — CRONOGRAMA FÍSICO-FINANCEIRO & CURVA S"
    ws1["A1"].fill = NAVY
    ws1["A1"].font = FONT_TITLE
    ws1["A1"].alignment = ALIGN_CENTER

    ws1.merge_cells("A3:G3")
    ws1["A3"] = f"Linha de Base Oficial 01 (Baseline 01) | Prazo: {prazo_meses} Meses ({prazo_meses * 30} Dias) | Base Orçamentária: Turnkey"
    ws1["A3"].fill = BLUE_DARK
    ws1["A3"].font = FONT_HEADER
    ws1["A3"].alignment = ALIGN_CENTER

    headers_curva = ["Mês / Período", "Faturamento Mensal (R$)", "% Mensal", "Faturamento Acumulado (R$)", "% Financeiro Acumulado", "% Físico Acumulado"]
    row_curva = 5
    for col_idx, h in enumerate(headers_curva, start=2):
        cell = ws1.cell(row=row_curva, column=col_idx, value=h)
        cell.fill = NAVY
        cell.font = FONT_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    totais_mes = [df[f'V_M{m}'].sum() for m in range(1, prazo_meses + 1)]
    tot_geral = sum(totais_mes) if sum(totais_mes) > 0 else 1.0

    acum_val = 0.0
    for m in range(1, prazo_meses + 1):
        r_num = row_curva + m
        v = totais_mes[m-1]
        acum_val += v
        pct_m = v / tot_geral
        pct_fin_acum = acum_val / tot_geral

        ws1.cell(row=r_num, column=2, value=f"Mês {m}").alignment = Alignment(horizontal="center")
        ws1.cell(row=r_num, column=3, value=v).number_format = '"R$ "#,##0.00'
        ws1.cell(row=r_num, column=4, value=pct_m).number_format = '0.00%'
        ws1.cell(row=r_num, column=5, value=acum_val).number_format = '"R$ "#,##0.00'
        ws1.cell(row=r_num, column=6, value=pct_fin_acum).number_format = '0.00%'
        ws1.cell(row=r_num, column=7, value=fisico_acum[m-1] / 100.0).number_format = '0.00%'

        for c_idx in range(2, 8):
            cell = ws1.cell(row=r_num, column=c_idx)
            cell.font = FONT_REGULAR
            cell.border = THIN_BORDER
            if m % 2 == 0:
                cell.fill = GRAY_LIGHT

    r_total = row_curva + prazo_meses + 1
    ws1.cell(row=r_total, column=2, value="TOTAL GLOBAL").alignment = Alignment(horizontal="center")
    ws1.cell(row=r_total, column=3, value=tot_geral).number_format = '"R$ "#,##0.00'
    ws1.cell(row=r_total, column=4, value=1.0).number_format = '0.00%'
    ws1.cell(row=r_total, column=5, value=tot_geral).number_format = '"R$ "#,##0.00'
    ws1.cell(row=r_total, column=6, value=1.0).number_format = '0.00%'
    ws1.cell(row=r_total, column=7, value=1.0).number_format = '0.00%'

    for c_idx in range(2, 8):
        cell = ws1.cell(row=r_total, column=c_idx)
        cell.fill = GREEN_FILL
        cell.font = FONT_BOLD
        cell.border = DOUBLE_BOTTOM_BORDER

    # Gráfico Curva S
    chart = LineChart()
    chart.title = "Curva S de Avanço Planejado (Físico vs Financeiro)"
    chart.style = 13
    chart.y_axis.title = "Avanço Acumulado (%)"
    chart.x_axis.title = "Período (Meses)"
    chart.width = 18
    chart.height = 10

    data_ref = Reference(ws1, min_col=6, min_row=row_curva, max_col=7, max_row=row_curva + prazo_meses)
    cats_ref = Reference(ws1, min_col=2, min_row=row_curva + 1, max_row=row_curva + prazo_meses)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    ws1.add_chart(chart, f"B{row_curva + prazo_meses + 4}")

    col_widths_ws1 = {1: 4, 2: 18, 3: 25, 4: 16, 5: 25, 6: 22, 7: 20}
    for col_idx, width in col_widths_ws1.items():
        ws1.column_dimensions[get_column_letter(col_idx)].width = width

    # Aba 2: Cronograma Analítico
    ws2 = wb.create_sheet(title="Cronograma Analítico")
    ws2.views.sheetView[0].showGridLines = True

    total_cols = 7 + (prazo_meses * 2) + 1
    last_col_letter = get_column_letter(total_cols)
    ws2.merge_cells(f"A1:{last_col_letter}2")
    ws2["A1"] = f"{sigla} — MATRIZ ANALÍTICA DE DISTRIBUIÇÃO MENSAL ({len(df)} ITENS DA EAP)"
    ws2["A1"].fill = NAVY
    ws2["A1"].font = FONT_TITLE
    ws2["A1"].alignment = ALIGN_CENTER

    headers_analitico = [
        "Código EAP", "Descrição do Pacote / Serviço", "Disciplina", "Qtd Proj", "Unid",
        "Preço Unit (R$)", "Preço Total Turnkey (R$)"
    ]
    for m in range(1, prazo_meses + 1):
        headers_analitico.extend([f"% M{m}", f"R$ Mês {m}"])
    headers_analitico.append("Total %")

    row_hdr = 4
    for col_idx, h in enumerate(headers_analitico, start=1):
        cell = ws2.cell(row=row_hdr, column=col_idx, value=h)
        cell.fill = NAVY
        cell.font = FONT_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    row_idx = 5
    for i, r in df.iterrows():
        ws2.cell(row=row_idx, column=1, value=str(r[eap_col]).strip()).alignment = Alignment(horizontal="center")
        ws2.cell(row=row_idx, column=2, value=str(r[desc_col]).strip()).alignment = Alignment(horizontal="left")
        ws2.cell(row=row_idx, column=3, value=str(r.get('Disciplina', '')).strip()).alignment = Alignment(horizontal="center")
        ws2.cell(row=row_idx, column=4, value=r['Qtd Projeto']).number_format = '#,##0.00'
        ws2.cell(row=row_idx, column=5, value=str(r.get('Unidade Proj', '')).strip()).alignment = Alignment(horizontal="center")
        ws2.cell(row=row_idx, column=6, value=r['Preco_Unitario']).number_format = '"R$ "#,##0.00'
        ws2.cell(row=row_idx, column=7, value=r['Preco_Total']).number_format = '"R$ "#,##0.00'

        curr_col = 8
        pct_col_letters = []
        for m in range(1, prazo_meses + 1):
            ws2.cell(row=row_idx, column=curr_col, value=r[f'P_M{m}']).number_format = '0.00%'
            pct_col_letters.append(get_column_letter(curr_col))
            curr_col += 1
            ws2.cell(row=row_idx, column=curr_col, value=r[f'V_M{m}']).number_format = '"R$ "#,##0.00'
            curr_col += 1

        sum_formula = "=" + "+".join([f"{ltr}{row_idx}" for ltr in pct_col_letters])
        ws2.cell(row=row_idx, column=curr_col, value=sum_formula).number_format = '0.00%'

        for c in range(1, total_cols + 1):
            cell = ws2.cell(row=row_idx, column=c)
            cell.font = FONT_REGULAR
            cell.border = THIN_BORDER
            if i % 2 == 1:
                cell.fill = GRAY_LIGHT
        row_idx += 1

    # Linha Totalizador
    ws2.cell(row=row_idx, column=1, value="TOTAL").alignment = Alignment(horizontal="center")
    ws2.cell(row=row_idx, column=2, value="TOTAL GERAL DO EMPREENDIMENTO (TURNKEY)").alignment = Alignment(horizontal="left")
    ws2.cell(row=row_idx, column=7, value=f"=SUM(G5:G{row_idx-1})").number_format = '"R$ "#,##0.00'

    for m_idx in range(prazo_meses):
        v_col = 9 + (m_idx * 2)
        p_col = 8 + (m_idx * 2)
        v_col_letter = get_column_letter(v_col)
        p_col_letter = get_column_letter(p_col)
        ws2.cell(row=row_idx, column=v_col, value=f"=SUM({v_col_letter}5:{v_col_letter}{row_idx-1})").number_format = '"R$ "#,##0.00'
        ws2.cell(row=row_idx, column=p_col, value=f"={v_col_letter}{row_idx}/G{row_idx}").number_format = '0.00%'

    ws2.cell(row=row_idx, column=total_cols, value=1.0).number_format = '0.00%'

    for c in range(1, total_cols + 1):
        cell = ws2.cell(row=row_idx, column=c)
        cell.fill = GREEN_FILL
        cell.font = FONT_BOLD
        cell.border = DOUBLE_BOTTOM_BORDER

    col_widths = {1: 14, 2: 45, 3: 20, 4: 12, 5: 10, 6: 16, 7: 22}
    for c_idx in range(8, total_cols):
        col_widths[c_idx] = 16 if c_idx % 2 == 1 else 10
    col_widths[total_cols] = 12
    for col_idx, width in col_widths.items():
        ws2.column_dimensions[get_column_letter(col_idx)].width = width

    wb.save(out_xlsx)
    print(f"-> Excel Executivo gerado: {out_xlsx}")
    return out_xlsx


def gerar_ms_project_xml(cpm_file, output_dir, sigla, titulo_obra, data_inicio_iso="2026-10-01", duracao_dias=180):
    out_xml = os.path.join(output_dir, f"CRONOGRAMA_{sigla}_MSPROJECT.xml")

    if not os.path.exists(cpm_file):
        raise FileNotFoundError(
            f"[ERRO AUD-009] Arquivo CPM '{cpm_file}' não encontrado. "
            "A exportação MS Project exige o cronograma CPM validado e não aceita tarefas sintéticas em modo produtivo."
        )

    with open(cpm_file, "r", encoding="utf-8") as f:
        cpm_raw = json.load(f)
    tasks_input = cpm_raw.get("atividades", [])
    if not tasks_input:
        raise ValueError(f"[ERRO AUD-009] Nenhuma atividade encontrada no arquivo CPM '{cpm_file}'.")

    # Cálculo rigoroso do CPM (Forward Pass, Backward Pass, folgas e criticidade)
    cpm_dto = calcular_cpm_detalhado(tasks_input, data_inicio_iso=data_inicio_iso)
    cpm_atividades = cpm_dto["atividades"]
    duracao_total_cpm = cpm_dto["duracao_total_dias"]

    start_date = datetime.date.fromisoformat(data_inicio_iso)
    finish_date = datetime.date.fromisoformat(cpm_dto["data_fim"])

    xml_content = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Project xmlns="http://schemas.microsoft.com/project">
    <Name>{sigla}_BASELINE_01</Name>
    <Title>{titulo_obra}</Title>
    <StartDate>{start_date.isoformat()}T08:00:00</StartDate>
    <FinishDate>{finish_date.isoformat()}T17:00:00</FinishDate>
    <CalendarUID>1</CalendarUID>
    <DefaultStartTime>08:00:00</DefaultStartTime>
    <DefaultFinishTime>17:00:00</DefaultFinishTime>
    <MinutesPerDay>480</MinutesPerDay>
    <MinutesPerWeek>2400</MinutesPerWeek>
    <DaysPerMonth>20</DaysPerMonth>
    <Tasks>
        <Task>
            <UID>0</UID>
            <ID>0</ID>
            <Name>{titulo_obra.upper()} ({duracao_total_cpm} DIAS CPM)</Name>
            <Type>1</Type>
            <CreateDate>{datetime.date.today().isoformat()}T08:00:00</CreateDate>
            <Start>{start_date.isoformat()}T08:00:00</Start>
            <Finish>{finish_date.isoformat()}T17:00:00</Finish>
            <Duration>PT{duracao_total_cpm * 8}H0M0S</Duration>
            <Summary>1</Summary>
            <Critical>1</Critical>
        </Task>
"""
    task_uid_map = {t["id"]: idx for idx, t in enumerate(tasks_input, start=1)}

    for idx, t in enumerate(tasks_input, start=1):
        aid = t["id"]
        cpm_info = cpm_atividades.get(aid, {})
        dur_days = cpm_info.get("duracao_dias", t.get("duracao_dias", 1))
        dur_hours = dur_days * 8
        is_critica = 1 if cpm_info.get("critica", False) else 0
        folga_total = cpm_info.get("folga_total", 0)
        total_slack_minutes = folga_total * 480
        t_start_iso = cpm_info.get("data_inicio", f"{start_date.isoformat()}T08:00:00")
        t_finish_iso = cpm_info.get("data_fim", f"{finish_date.isoformat()}T17:00:00")
        if "T" not in t_start_iso:
            t_start_iso = f"{t_start_iso}T08:00:00"
        if "T" not in t_finish_iso:
            t_finish_iso = f"{t_finish_iso}T17:00:00"

        preds_xml = ""
        for p in t.get("predecessoras", []):
            if p in task_uid_map:
                preds_xml += f"""            <PredecessorLink>
                <PredecessorUID>{task_uid_map[p]}</PredecessorUID>
                <Type>1</Type>
                <CrossProject>0</CrossProject>
                <LinkLag>0</LinkLag>
                <LagFormat>7</LagFormat>
            </PredecessorLink>
"""
        xml_content += f"""        <Task>
            <UID>{idx}</UID>
            <ID>{idx}</ID>
            <Name>{aid.replace('_', ' ')}</Name>
            <Active>1</Active>
            <Duration>PT{dur_hours}H0M0S</Duration>
            <Start>{t_start_iso}</Start>
            <Finish>{t_finish_iso}</Finish>
            <Critical>{is_critica}</Critical>
            <TotalSlack>{total_slack_minutes}</TotalSlack>
{preds_xml}        </Task>
"""
    xml_content += """    </Tasks>
</Project>
"""
    with open(out_xml, "w", encoding="utf-8") as f:
        f.write(xml_content)
    print(f"-> XML MS Project gerado: {out_xml}")
    return out_xml


def gerar_dashboard_html(df, output_dir, sigla, titulo_obra, prazo_meses=6, fisico_acum=None):
    out_html = os.path.join(output_dir, "CRONOGRAMA_DASHBOARD_INTERATIVO.html")

    if fisico_acum is None or len(fisico_acum) < prazo_meses:
        fisico_acum = [round(100.0 * (m / prazo_meses), 2) for m in range(1, prazo_meses + 1)]

    totais_mes = [df[f'V_M{m}'].sum() for m in range(1, prazo_meses + 1)]
    tot_geral = sum(totais_mes) if sum(totais_mes) > 0 else 1.0
    acum_val = 0.0
    pct_fin_acum = []
    for v in totais_mes:
        acum_val += v
        pct_fin_acum.append(round(acum_val / tot_geral * 100, 2))

    meses = [f"Mês {m}" for m in range(1, prazo_meses + 1)]

    fig = make_subplots(
        rows=2, cols=1,
        subplot_titles=("Curva S Planejada: Avanço Físico vs Financeiro Acumulado", "Desembolso e Faturamento Mensal (R$)"),
        vertical_spacing=0.15
    )

    fig.add_trace(
        go.Scatter(x=meses, y=fisico_acum, name="% Físico Acumulado", mode='lines+markers', line=dict(color='#1B365D', width=3), marker=dict(size=8)),
        row=1, col=1
    )
    fig.add_trace(
        go.Scatter(x=meses, y=pct_fin_acum, name="% Financeiro Acumulado", mode='lines+markers', line=dict(color='#D99B26', width=3, dash='dot'), marker=dict(size=8)),
        row=1, col=1
    )
    fig.add_trace(
        go.Bar(x=meses, y=totais_mes, name="Faturamento Mensal (R$)", marker_color='#284B78', text=[f"R$ {v:,.0f}" for v in totais_mes], textposition='auto'),
        row=2, col=1
    )

    fig.update_layout(
        title=f"<b>PMO VIRTUAL: DASHBOARD EXECUTIVO DO CRONOGRAMA — {titulo_obra.upper()}</b><br><sup>Baseline 01 | R$ {tot_geral:,.2f} Turnkey | {prazo_meses * 30} Dias Corridos</sup>",
        height=750,
        showlegend=True,
        template="plotly_white",
        font=dict(family="Arial, sans-serif")
    )
    fig.write_html(out_html)
    print(f"-> Dashboard HTML Interativo gerado: {out_html}")
    return out_html


def main():
    args = parse_args()
    obra_dir = resolver_obra_dir(args)
    config = carregar_config_obra(obra_dir)

    # Suporte a estrutura aninhada (dados_obra) e plana (campo direto)
    dados_obra = config.get("dados_obra", config)
    sigla = dados_obra.get("sigla", os.path.basename(obra_dir).replace("OBRA_", ""))
    titulo = dados_obra.get("nome_obra", config.get("nome_obra", f"Obra {sigla}"))
    prazo_meses = int(dados_obra.get("prazo_meses", config.get("prazo_meses", 6)))
    data_inicio = dados_obra.get("data_inicio", config.get("data_inicio", "2026-10-01"))
    # Normalizar formato de data para ISO (YYYY-MM-DD)
    if "/" in str(data_inicio):
        partes = data_inicio.split("/")
        data_inicio = f"{partes[2]}-{partes[1]}-{partes[0]}"
    curva_fisico = dados_obra.get("curva_fisico_acumulado", config.get("curva_fisico_acumulado", [12.05, 30.20, 54.10, 69.85, 88.40, 100.00]))

    orcamento_csv = os.path.join(obra_dir, "02_ORCAMENTO_BASE_E_CONTRATOS", "ORCAMENTO_BASE_CONSOLIDADO.csv")
    output_dir = os.path.join(obra_dir, "03_PLANEJAMENTO_E_CRONOGRAMA")
    cpm_file = os.path.join(output_dir, "dados_cpm.json")

    os.makedirs(output_dir, exist_ok=True)

    if not os.path.exists(orcamento_csv):
        print(f"[ERRO] Orçamento consolidado não encontrado em: {orcamento_csv}")
        print("Execute primeiro: python scripts/precificar_obra.py --obra <NOME>")
        sys.exit(1)

    print(f"=== INICIANDO MOTOR UNIVERSAL DE CRONOGRAMA: {titulo} ({sigla}) ===")
    df, eap_col, desc_col = carregar_e_distribuir_orcamento(orcamento_csv, prazo_meses=prazo_meses)
    gerar_csv(df, eap_col, desc_col, output_dir, sigla, prazo_meses=prazo_meses)
    gerar_excel(df, eap_col, desc_col, output_dir, sigla, titulo, prazo_meses=prazo_meses, fisico_acum=curva_fisico)
    gerar_ms_project_xml(cpm_file, output_dir, sigla, titulo, data_inicio_iso=data_inicio, duracao_dias=prazo_meses * 30)
    gerar_dashboard_html(df, output_dir, sigla, titulo, prazo_meses=prazo_meses, fisico_acum=curva_fisico)
    print("=== CRONOGRAMA FÍSICO-FINANCEIRO GERADO COM SUCESSO! ===")


if __name__ == "__main__":
    main()
