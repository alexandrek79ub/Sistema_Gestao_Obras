import os
import zipfile
import csv
import io
import re
import openpyxl

def main():
    zip_path = os.path.abspath("apoio/sinapi_sp/SINAPI-2026-07-formato-xlsx.zip")
    output_dir = os.path.abspath("apoio/sinapi_sp")
    os.makedirs(output_dir, exist_ok=True)
    
    print(f"[INFO] Abrindo {zip_path}...")
    with zipfile.ZipFile(zip_path, 'r') as z:
        excel_name = "SINAPI_Referência_2026_07.xlsx"
        print(f"[INFO] Carregando {excel_name} (modo read_only)...")
        excel_bytes = io.BytesIO(z.read(excel_name))
        wb = openpyxl.load_workbook(excel_bytes, read_only=True, data_only=False)
        
        # 1. Extrair Composições SP (CSD - Sem Desoneração)
        print("[INFO] Processando aba CSD (Composições SP)...")
        ws_csd = wb['CSD']
        csv_csd_path = os.path.join(output_dir, "SINAPI_SP_COMPOSICOES_2026_07.csv")
        
        count_csd = 0
        with open(csv_csd_path, 'w', newline='', encoding='utf-8') as f_out:
            writer = csv.writer(f_out, delimiter=';')
            writer.writerow(["GRUPO", "CODIGO_COMPOSICAO", "DESCRICAO", "UNIDADE", "CUSTO_TOTAL_SP_RS"])
            
            for row in ws_csd.iter_rows(min_row=11, values_only=True):
                grupo = row[0]
                raw_cod = row[1]
                desc = row[2]
                unid = row[3]
                custo_sp = row[54] if len(row) > 54 else None
                
                # Extrai o código da composição da fórmula HYPERLINK ou string/int
                cod = None
                if raw_cod:
                    m = re.search(r'(\d{4,7})', str(raw_cod))
                    if m:
                        cod = m.group(1)
                    else:
                        cod = str(raw_cod).strip()
                
                if cod and desc:
                    writer.writerow([grupo, cod, desc, unid, custo_sp])
                    count_csd += 1
        print(f"[OK] {count_csd} composições salvas em: {csv_csd_path}")
        
        # 2. Extrair Insumos SP (ISD - Sem Desoneração)
        print("[INFO] Processando aba ISD (Insumos SP)...")
        ws_isd = wb['ISD']
        csv_isd_path = os.path.join(output_dir, "SINAPI_SP_INSUMOS_2026_07.csv")
        
        count_isd = 0
        with open(csv_isd_path, 'w', newline='', encoding='utf-8') as f_out:
            writer = csv.writer(f_out, delimiter=';')
            writer.writerow(["GRUPO", "CODIGO_INSUMO", "DESCRICAO", "UNIDADE", "PRECO_UNIT_SP_RS"])
            
            for row in ws_isd.iter_rows(min_row=11, values_only=True):
                grupo = row[0]
                cod = row[1]
                desc = row[2]
                unid = row[3]
                preco_sp = row[30] if len(row) > 30 else None
                
                if cod and desc:
                    writer.writerow([grupo, cod, desc, unid, preco_sp])
                    count_isd += 1
        print(f"[OK] {count_isd} insumos salvos em: {csv_isd_path}")

    print("\n[SUCESSO] Base SINAPI SP 07/2026 totalmente extraída e pronta para uso!")

if __name__ == "__main__":
    main()
