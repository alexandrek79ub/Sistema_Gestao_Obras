#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Tema Visual Corporativo — OpenPyXL.

Centraliza 100% das cores, fontes, bordas e funções de formatação de planilhas
usadas pelos Motores Universais. Elimina ~2.250 linhas de código duplicado em 15+ scripts.

IMPORTANTE: Este módulo é puramente visual. NÃO contém nenhuma lógica de cálculo
            de engenharia. Os valores nas células NÃO são alterados por este módulo.
"""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# =============================================================================
# PALETA DE CORES CORPORATIVA (Hexadecimal sem #)
# =============================================================================
# Azuis institucionais
NAVY        = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
BLUE_DARK   = PatternFill(start_color="284B78", end_color="284B78", fill_type="solid")
BLUE_LIGHT  = PatternFill(start_color="E8EEF5", end_color="E8EEF5", fill_type="solid")
BLUE_ACCENT = PatternFill(start_color="C2D7EF", end_color="C2D7EF", fill_type="solid")

# Destaques e alertas
GOLD_ACCENT    = PatternFill(start_color="D99B26", end_color="D99B26", fill_type="solid")
YELLOW_LIGHT   = PatternFill(start_color="FEF7E0", end_color="FEF7E0", fill_type="solid")
YELLOW_ACCENT  = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

# Status operacional
GREEN_FILL  = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
GREEN_DARK  = PatternFill(start_color="137333", end_color="137333", fill_type="solid")
RED_FILL    = PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid")

# Fundos neutros
HEADER_GRAY = PatternFill(start_color="EAEEF3", end_color="EAEEF3", fill_type="solid")
ZEBRA_LIGHT = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
GRAY_LIGHT  = PatternFill(start_color="F4F6F9", end_color="F4F6F9", fill_type="solid")

# =============================================================================
# FONTES TIPADAS
# =============================================================================
FONT_TITLE   = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
FONT_HEADER  = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_SUBHDR  = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_BOLD    = Font(name="Calibri", size=11, bold=True, color="1B365D")
FONT_REGULAR = Font(name="Calibri", size=11, color="333333")
FONT_SMALL   = Font(name="Calibri", size=9,  color="555555")
FONT_RED     = Font(name="Calibri", size=11, bold=True, color="C5221F")
FONT_GREEN   = Font(name="Calibri", size=11, bold=True, color="137333")
FONT_GOLD    = Font(name="Calibri", size=11, bold=True, color="D99B26")

# =============================================================================
# BORDAS
# =============================================================================
THIN_BORDER = Border(
    left=Side(style='thin', color='DDDDDD'),
    right=Side(style='thin', color='DDDDDD'),
    top=Side(style='thin', color='DDDDDD'),
    bottom=Side(style='thin', color='DDDDDD')
)

DOUBLE_BOTTOM_BORDER = Border(
    left=Side(style='thin', color='DDDDDD'),
    right=Side(style='thin', color='DDDDDD'),
    top=Side(style='thin', color='DDDDDD'),
    bottom=Side(style='double', color='1B365D')
)

# =============================================================================
# ALINHAMENTOS
# =============================================================================
ALIGN_CENTER       = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_LEFT         = Alignment(horizontal='left',   vertical='center', wrap_text=True)
ALIGN_RIGHT        = Alignment(horizontal='right',  vertical='center')
ALIGN_CENTER_NOWRAP = Alignment(horizontal='center', vertical='center')


# =============================================================================
# FUNÇÕES UTILITÁRIAS DE FORMATAÇÃO
# =============================================================================

def aplicar_cabecalho_tabela(ws, colunas, row_num=1, fill=None, font=None, altura=30):
    """
    Escreve e formata uma linha de cabeçalho de tabela.

    Args:
        ws       : Worksheet OpenPyXL.
        colunas  : Lista de strings com os títulos das colunas.
        row_num  : Número da linha onde o cabeçalho será inserido (padrão: 1).
        fill     : PatternFill a usar (padrão: NAVY).
        font     : Font a usar (padrão: FONT_HEADER).
        altura   : Altura da linha em pontos (padrão: 30).
    """
    fill = fill or NAVY
    font = font or FONT_HEADER
    ws.row_dimensions[row_num].height = altura
    for col_idx, titulo in enumerate(colunas, start=1):
        c = ws.cell(row=row_num, column=col_idx, value=titulo)
        c.fill = fill
        c.font = font
        c.alignment = ALIGN_CENTER
        c.border = THIN_BORDER


def aplicar_linha_titulo(ws, texto, row_num=1, col_span=10, fill=None, font=None, altura=35):
    """
    Mescla células e aplica um título principal à planilha.

    Args:
        ws       : Worksheet OpenPyXL.
        texto    : Texto do título.
        row_num  : Linha do título (padrão: 1).
        col_span : Número de colunas a mesclar (padrão: 10).
        fill     : PatternFill (padrão: NAVY).
        font     : Font (padrão: FONT_TITLE).
        altura   : Altura da linha em pontos (padrão: 35).
    """
    fill = fill or NAVY
    font = font or FONT_TITLE
    ultima_col = get_column_letter(col_span)
    ws.merge_cells(f"A{row_num}:{ultima_col}{row_num}")
    c = ws.cell(row=row_num, column=1, value=texto)
    c.fill = fill
    c.font = font
    c.alignment = ALIGN_CENTER
    ws.row_dimensions[row_num].height = altura


def formatar_linha_dados(ws, row_num, num_cols, zebra=True, fill_override=None):
    """
    Aplica zebrado e borda a uma linha de dados.

    Args:
        ws            : Worksheet OpenPyXL.
        row_num       : Número da linha.
        num_cols      : Número total de colunas.
        zebra         : Se True, aplica ZEBRA_LIGHT em linhas pares (padrão: True).
        fill_override : PatternFill específico para sobrescrever o zebrado.
    """
    fill = fill_override
    if fill is None and zebra and row_num % 2 == 0:
        fill = ZEBRA_LIGHT
    for col_idx in range(1, num_cols + 1):
        c = ws.cell(row=row_num, column=col_idx)
        if fill:
            c.fill = fill
        if not c.font or not c.font.bold:
            c.font = FONT_REGULAR
        c.border = THIN_BORDER
        if c.alignment is None or not c.alignment.horizontal:
            c.alignment = ALIGN_LEFT


def formatar_linha_total(ws, row_num, num_cols, fill=None, font=None):
    """
    Formata a linha de totais/rodapé de uma tabela.

    Args:
        ws       : Worksheet OpenPyXL.
        row_num  : Número da linha de total.
        num_cols : Número total de colunas.
        fill     : PatternFill (padrão: GOLD_ACCENT).
        font     : Font (padrão: FONT_BOLD).
    """
    fill = fill or GOLD_ACCENT
    font = font or FONT_BOLD
    ws.row_dimensions[row_num].height = 22
    for col_idx in range(1, num_cols + 1):
        c = ws.cell(row=row_num, column=col_idx)
        c.fill = fill
        c.font = font
        c.border = DOUBLE_BOTTOM_BORDER
        c.alignment = ALIGN_CENTER


def auto_ajustar_colunas(ws, min_width=10, max_width=60, padding=4):
    """
    Ajusta automaticamente a largura de todas as colunas com base no conteúdo.

    Args:
        ws        : Worksheet OpenPyXL.
        min_width : Largura mínima em caracteres (padrão: 10).
        max_width : Largura máxima em caracteres (padrão: 60).
        padding   : Caracteres extras de folga (padrão: 4).
    """
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        width = max(min_width, min(max_width, max_len + padding))
        ws.column_dimensions[col_letter].width = width


def formatar_moeda(val):
    """
    Formata um float como moeda brasileira: R$ 1.234.567,89
    Conforme Manual de Boas Práticas §3.2.
    """
    try:
        val = float(val)
    except (TypeError, ValueError):
        return "R$ 0,00"
    if abs(val) < 0.001:
        return "R$ 0,00"
    s = f"R$ {abs(val):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"-{s}" if val < 0 else s
