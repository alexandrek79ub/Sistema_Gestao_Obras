import http.client
import json
import sqlite3
import tempfile
import threading
import unittest
from http.server import ThreadingHTTPServer
from pathlib import Path

from api_pmo import Handler
from motor_quantitativos.importadores.pdf_json_importer import importar_json_inicial
from motor_quantitativos.domain.elementos import QuantifiedItem
from motor_quantitativos.repositorio.sqlite_repository import (
    atualizar_quantitativo, connect, garantir_obra, gravar_orcamento, 
    recalcular_orcamento, substituir_quantitativos, persistir_itens_quantificados
)
from motor_quantitativos.calculo.avaliador_expressoes import calcular_expressao
from motor_quantitativos.importadores.roteador import processar_prancha
from motor_quantitativos.exportadores import exportar_artefatos
from motor_quantitativos.auditoria.trilha_revisoes import registrar_revisao


class MotorQuantitativosTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.base = Path(self.tmp.name) / "obra"
        self.db = connect(Path(self.tmp.name) / "teste.sqlite")
        self.obra = garantir_obra(self.db, "TESTE", "Obra Teste", str(self.base))
        revisao = registrar_revisao(self.db, self.obra, "QUANTITATIVO", "teste")
        self.item = substituir_quantitativos(self.db, self.obra, [{
            "cod_eap": "1.3.11", "descricao": "Impermeabilização", "disciplina": "Fundação",
            "unidade": "m²", "quantidade_liquida": 12.5, "expressao_matematica": "5 * 2.5",
            "prancha_referencia": "F-01", "status": "LEVANTADO",
        }], revisao)[0]
        revisao_orc = registrar_revisao(self.db, self.obra, "ORCAMENTO", "teste")
        gravar_orcamento(self.db, self.obra, [{
            "cod_eap": "1.3.11", "prancha_referencia": "F-01", "preco_unitario": 20,
            "bdi_pct": 10, "codigo_sinapi": "123", "centro_custo": "CC-01", "fonte_preco": "SINAPI",
        }], revisao_orc)
        self.db.commit()

    def tearDown(self):
        self.db.close()
        self.tmp.cleanup()

    def test_schema_mantem_quantidade_liquida_sem_perdas(self):
        row = self.db.execute("SELECT quantidade_liquida,status FROM itens_quantitativo WHERE id=?", (self.item,)).fetchone()
        self.assertEqual((row[0], row[1]), (12.5, "LEVANTADO"))
        self.assertFalse(self.db.execute("SELECT 1 FROM pragma_table_info('itens_quantitativo') WHERE name='taxa_perda'").fetchone())
        self.assertEqual(self.db.execute("SELECT MAX(versao) FROM schema_migrations").fetchone()[0], 3)

    def test_edicao_auditavel_recalcula_orcamento_e_exporta(self):
        revisao = atualizar_quantitativo(self.db, self.obra, self.item, {"quantidade_liquida": 14.0}, "teste-web", "Correção da prancha", 1)
        recalcular_orcamento(self.db, self.obra, revisao)
        arquivos = exportar_artefatos(self.db, self.obra)
        self.db.commit()
        linha_atualizada = self.db.execute("SELECT quantidade_liquida,versao FROM itens_quantitativo WHERE id=?", (self.item,)).fetchone()
        self.assertEqual((linha_atualizada["quantidade_liquida"], linha_atualizada["versao"]), (14.0, 2))
        self.assertEqual(self.db.execute("SELECT custo_total FROM itens_orcamento WHERE quantitativo_id=?", (self.item,)).fetchone()[0], 308.0)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM auditoria_eventos").fetchone()[0], 1)
        self.assertTrue((self.base / "QUANTITATIVO_MESTRE.csv").exists())
        self.assertTrue((self.base / "ORCAMENTO_BASE_CONSOLIDADO.csv").exists())
        self.assertTrue(any(path.name.startswith("MEMORIA_CALCULO_") for path in arquivos))
        self.assertIn("14.0", (self.base / "ORCAMENTO_BASE_CONSOLIDADO.csv").read_text(encoding="utf-8-sig"))

    def test_rejeita_conflito_sem_justificativa_ou_versao_atual(self):
        with self.assertRaises(ValueError):
            atualizar_quantitativo(self.db, self.obra, self.item, {"quantidade_liquida": 14.0}, "teste", "", 1)
        atualizar_quantitativo(self.db, self.obra, self.item, {"quantidade_liquida": 14.0}, "teste", "Ajuste confirmado", 1)
        with self.assertRaises(RuntimeError):
            atualizar_quantitativo(self.db, self.obra, self.item, {"quantidade_liquida": 15.0}, "teste", "Tentativa antiga", 1)

    def test_importacao_inicial_nao_sobrescreve_banco_sem_flag_explicita(self):
        arquivo_json = Path(self.tmp.name) / "entrada.json"
        destino = Path(self.tmp.name) / "exportacoes"
        arquivo_json.write_text(json.dumps({"projeto": "Importação Teste", "base_dir": str(destino), "disciplinas": {
            "fundacao": {"titulo": "Fundação", "pranchas_ref": "F-01", "itens_orcamento": [{
                "codigo_eap": "1.3.11", "descricao": "Impermeabilização", "unidade": "m²",
                "equacoes": [{"expressao_matematica": "5 * 2.5"}],
            }]}}}, ensure_ascii=False), encoding="utf-8")
        banco = Path(self.tmp.name) / "importacao.sqlite"
        saidas = importar_json_inicial(str(arquivo_json), str(banco))
        self.assertTrue(any(arquivo.name == "ORCAMENTO_BASE_CONSOLIDADO.csv" for arquivo in saidas))
        self.assertTrue(any(arquivo.name == "ORCAMENTO_BASE_CONSOLIDADO.xlsx" for arquivo in saidas))
        
        arquivo_xlsx = next(a for a in saidas if a.name == "ORCAMENTO_BASE_CONSOLIDADO.xlsx")
        import openpyxl
        wb = openpyxl.load_workbook(str(arquivo_xlsx))
        self.assertIn("01_Orçamento_Base", wb.sheetnames)
        self.assertIn("02_BDI_Analítico", wb.sheetnames)
        self.assertIn("03_Composições_CCU", wb.sheetnames)
        self.assertIn("04_Curva_ABC_Serviços", wb.sheetnames)
        self.assertIn("05_Memória_Quantitativos", wb.sheetnames)
        self.assertIn("06_Boletim_Medição", wb.sheetnames)
        ws = wb["01_Orçamento_Base"]
        self.assertEqual(ws["H5"].value, "=ROUND(F5*(1+G5/100), 2)")
        self.assertEqual(ws["I5"].value, "=ROUND(E5*H5, 2)")
        wb.close()

        db_import = connect(banco)
        try:
            self.assertEqual(db_import.execute("SELECT COUNT(*) FROM itens_orcamento").fetchone()[0], 0)
        finally:
            db_import.close()


        with self.assertRaises(ValueError):
            importar_json_inicial(str(arquivo_json), str(banco))


    def test_importacao_fisica_rejeita_campos_de_orcamento(self):
        arquivo = Path(self.tmp.name) / "com_preco.json"
        arquivo.write_text(json.dumps({"projeto": "Teste", "disciplinas": {"fundacao": {"itens_orcamento": [{
            "codigo_eap": "1.3.9", "descricao": "Concreto", "unidade": "m3", "preco_unitario": 1,
            "equacoes": [{"expressao_matematica": "1 * 1"}],
        }]}}}), encoding="utf-8")
        with self.assertRaises(ValueError):
            importar_json_inicial(str(arquivo), str(Path(self.tmp.name) / "preco.sqlite"))

    def test_persistencia_preserva_elementos_distintos_da_mesma_prancha(self):
        revisao = registrar_revisao(self.db, self.obra, "QUANTITATIVO", "teste")
        itens = [
            QuantifiedItem("1.3.9", "Sapata S1", "m3", 1.0, "1 * 1", "FUN.SAPATA_CONCRETO", 1,
                            ("S1",), ("E1",), cia="FUN-GER-S1", element_type="SAPATA", source_file="F-01.pdf", source_revision="R1"),
            QuantifiedItem("1.3.9", "Sapata S2", "m3", 2.0, "2 * 1", "FUN.SAPATA_CONCRETO", 1,
                            ("S2",), ("E2",), cia="FUN-GER-S2", element_type="SAPATA", source_file="F-01.pdf", source_revision="R1"),
        ]
        persistir_itens_quantificados(self.db, self.obra, itens, revisao)
        rows = self.db.execute("SELECT element_id,cia,source_revision FROM itens_quantitativo WHERE rule_id='FUN.SAPATA_CONCRETO' ORDER BY element_id").fetchall()
        self.assertEqual([(row["element_id"], row["cia"], row["source_revision"]) for row in rows], [("S1", "FUN-GER-S1", "R1"), ("S2", "FUN-GER-S2", "R1")])

    def test_avaliador_bloqueia_expressao_invalida(self):
        with self.assertRaises(ValueError):
            calcular_expressao("__import__('os').system('x')")

    def test_migracao_v3_preserva_orcamento_de_schema_anterior(self):
        banco = Path(self.tmp.name) / "v2.sqlite"
        legado = sqlite3.connect(banco)
        legado.executescript("""
        CREATE TABLE schema_migrations (versao INTEGER PRIMARY KEY, aplicada_em TEXT NOT NULL);
        INSERT INTO schema_migrations VALUES (2, '2026-01-01T00:00:00Z');
        CREATE TABLE obras (id INTEGER PRIMARY KEY, codigo TEXT NOT NULL UNIQUE, nome TEXT NOT NULL, diretorio_base TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
        CREATE TABLE revisoes (id INTEGER PRIMARY KEY, obra_id INTEGER NOT NULL, tipo TEXT NOT NULL, usuario TEXT NOT NULL, justificativa TEXT NOT NULL, origem TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE itens_quantitativo (id INTEGER PRIMARY KEY, obra_id INTEGER NOT NULL, revisao_id INTEGER, cod_eap TEXT NOT NULL, descricao TEXT NOT NULL, disciplina TEXT NOT NULL, unidade TEXT NOT NULL, quantidade_liquida REAL NOT NULL, expressao_matematica TEXT NOT NULL, prancha_referencia TEXT NOT NULL, status TEXT NOT NULL, rfi TEXT NOT NULL DEFAULT '', observacao TEXT NOT NULL DEFAULT '', cia TEXT NOT NULL DEFAULT '', element_type TEXT NOT NULL DEFAULT '', element_id TEXT NOT NULL DEFAULT '', rule_id TEXT NOT NULL DEFAULT '', rule_version INTEGER NOT NULL DEFAULT 0, evidence_json TEXT NOT NULL DEFAULT '[]', source_revision TEXT NOT NULL DEFAULT '', versao INTEGER NOT NULL DEFAULT 1, updated_at TEXT NOT NULL, UNIQUE(obra_id, cod_eap, prancha_referencia));
        CREATE TABLE itens_orcamento (id INTEGER PRIMARY KEY, obra_id INTEGER NOT NULL, quantitativo_id INTEGER NOT NULL, revisao_id INTEGER, codigo_sinapi TEXT NOT NULL DEFAULT '', centro_custo TEXT NOT NULL DEFAULT '', fonte_preco TEXT NOT NULL DEFAULT '', custo_material REAL NOT NULL DEFAULT 0, custo_mao_obra REAL NOT NULL DEFAULT 0, custo_equipamento REAL NOT NULL DEFAULT 0, bdi_pct REAL NOT NULL DEFAULT 0, preco_unitario REAL NOT NULL DEFAULT 0, custo_total REAL NOT NULL DEFAULT 0, updated_at TEXT NOT NULL, UNIQUE(obra_id, quantitativo_id));
        INSERT INTO obras VALUES (1, 'OBRA', 'Obra', '', 't', 't');
        INSERT INTO itens_quantitativo VALUES (1, 1, NULL, '1.3.9', 'Concreto', 'Fundacao', 'm3', 1, '1', 'F-01', 'LEVANTADO', '', '', '', '', '', '', 0, '[]', '', 1, 't');
        INSERT INTO itens_orcamento VALUES (1, 1, 1, NULL, '', '', '', 0, 0, 0, 0, 10, 10, 't');
        """)
        legado.commit()
        legado.close()
        migrado = connect(banco)
        try:
            self.assertEqual(migrado.execute("SELECT MAX(versao) FROM schema_migrations").fetchone()[0], 3)
            self.assertEqual(migrado.execute("SELECT custo_total FROM itens_orcamento WHERE quantitativo_id=1").fetchone()[0], 10)
        finally:
            migrado.close()

    def test_fluxo_pdf_contratual_exige_confirmacao_e_persiste_rastreabilidade(self):
        import fitz
        pdf = Path(self.tmp.name) / "F-01.pdf"
        documento = fitz.open()
        documento.new_page().insert_text((72, 72), "2x S1 (100x150x40 cm)")
        documento.save(pdf)
        documento.close()
        banco = Path(self.tmp.name) / "fluxo.sqlite"
        saida = Path(self.tmp.name) / "artefatos"
        args = dict(obra_codigo="OBRA_TESTE", obra_nome="Obra Teste", source_revision="R1",
                    disciplina="FUNDACOES", db_path=banco, diretorio_obra=saida)
        with self.assertRaises(ValueError):
            processar_prancha(pdf, **args)
        resultado = processar_prancha(pdf, evidencias_confirmadas=True, **args)
        self.assertEqual((resultado["elementos"], resultado["itens"]), (1, 1))
        db = connect(banco)
        try:
            row = db.execute("SELECT quantidade_liquida,element_id,rule_id,source_revision,evidence_json FROM itens_quantitativo").fetchone()
            self.assertEqual((row["quantidade_liquida"], row["element_id"], row["rule_id"], row["source_revision"]), (1.2, "S1", "FUN.SAPATA_CONCRETO", "R1"))
            self.assertIn("EVD-", row["evidence_json"])
        finally:
            db.close()

    def test_api_exige_chave_versao_e_regenera_exportacoes(self):
        anteriores = Handler.db_path, Handler.api_key, Handler.usuario, Handler.backup_dir
        Handler.db_path, Handler.api_key, Handler.usuario, Handler.backup_dir = str(Path(self.tmp.name) / "teste.sqlite"), "segredo", "engenheiro-web", str(Path(self.tmp.name) / "backups")
        servidor = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=servidor.serve_forever, daemon=True)
        thread.start()
        try:
            porta = servidor.server_address[1]
            conexao = http.client.HTTPConnection("127.0.0.1", porta, timeout=3)
            conexao.request("GET", f"/api/quantitativos?obra_id={self.obra}")
            self.assertEqual(conexao.getresponse().status, 401)
            payload = json.dumps({"alteracoes": {"quantidade_liquida": 14}, "justificativa": "Revisão conferida", "versao_esperada": 1})
            conexao.request("PATCH", f"/api/quantitativo/{self.obra}/{self.item}", body=payload,
                             headers={"Content-Type": "application/json", "X-PMO-API-Key": "segredo"})
            self.assertEqual(conexao.getresponse().status, 200)
            self.assertTrue((self.base / "ORCAMENTO_BASE_CONSOLIDADO.csv").exists())
            self.assertEqual(len(list((Path(self.tmp.name) / "backups").glob("*.sqlite"))), 1)
        finally:
            servidor.shutdown()
            servidor.server_close()
            thread.join(timeout=3)
            Handler.db_path, Handler.api_key, Handler.usuario, Handler.backup_dir = anteriores


if __name__ == "__main__":
    unittest.main()
