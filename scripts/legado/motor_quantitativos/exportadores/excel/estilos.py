import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Paleta Corporativa PMO Virtual (Slate / Navy)
COR_TOPO = "1E293B"       # Slate 800
COR_CABECALHO = "0F172A"  # Slate 900
COR_SUB = "334155"        # Slate 700
COR_ZEBRADO = "F8FAFC"    # Slate 50
COR_TOTAL = "F1F5F9"      # Slate 100
COR_BORDA = "CBD5E1"      # Slate 300


def borda_padrao() -> Border:
    fina = Side(style="thin", color=COR_BORDA)
    return Border(left=fina, right=fina, top=fina, bottom=fina)


def borda_total() -> Border:
    fina = Side(style="thin", color="64748B")
    dupla = Side(style="double", color="0F172A")
    return Border(top=fina, bottom=dupla, left=fina, right=fina)


def ajustar_larguras(ws) -> None:
    ws.views.sheetView[0].showGridLines = True
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if val_str.startswith("="):
                val_str = "R$ 999.999,99"
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
