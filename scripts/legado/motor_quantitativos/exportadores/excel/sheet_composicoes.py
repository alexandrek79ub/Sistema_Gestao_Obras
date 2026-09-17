from typing import Any
from openpyxl.styles import Alignment, Font, PatternFill
from motor_quantitativos.exportadores.excel.estilos import (
    COR_CABECALHO, COR_TOPO, COR_ZEBRADO,
    ajustar_larguras, borda_padrao
)


def adicionar_aba_composicoes(wb, nome_obra: str, itens: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title="03_Composições_CCU")
    ws.freeze_panes = "A4"

    ws.merge_cells("A1:J1")
    ws["A1"].value = f"COMPOSIÇÃO ANALÍTICA DE CUSTOS UNITÁRIOS (CCU) — {nome_obra.upper()}"
    ws["A1"].font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color=COR_TOPO, end_color=COR_TOPO, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    colunas = [
        ("Item EAP", "center"), ("Descrição do Serviço", "left"), ("Unid", "center"),
        ("Material (R$)", "right"), ("Mão de Obra (R$)", "right"), ("Equipamento (R$)", "right"),
        ("Custo Direto (R$)", "right"), ("BDI (%)", "right"), ("Preço Unit. c/ BDI (R$)", "right"),
        ("Fonte / Referência", "left"),
    ]

    for col_idx, (nome_col, alin) in enumerate(colunas, start=1):
        cell = ws.cell(row=3, column=col_idx, value=nome_col)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type="solid")
        cell.alignment = Alignment(horizontal=alin, vertical="center")
        cell.border = borda_padrao()
    ws.row_dimensions[3].height = 24

    for idx, item in enumerate(itens):
        r_num = 4 + idx
        ws.row_dimensions[r_num].height = 20
        
        c_mat = float(item.get("custo_material") or 0)
        c_mo = float(item.get("custo_mao_obra") or 0)
        c_eq = float(item.get("custo_equipamento") or 0)
        bdi = float(item.get("bdi_pct") or 0)

        valores = [
            (item.get("cod_eap", ""), "@", "center"),
            (item.get("descricao", ""), "@", "left"),
            (item.get("unidade", ""), "@", "center"),
            (c_mat, "R$ #,##0.00", "right"),
            (c_mo, "R$ #,##0.00", "right"),
            (c_eq, "R$ #,##0.00", "right"),
            (f"=ROUND(SUM(D{r_num}:F{r_num}), 2)", "R$ #,##0.00", "right"),
            (bdi, "0.00", "right"),
            (f"=ROUND(G{r_num}*(1+H{r_num}/100), 2)", "R$ #,##0.00", "right"),
            (item.get("fonte_preco", ""), "@", "left"),
        ]

        fill_cor = PatternFill(start_color=COR_ZEBRADO, end_color=COR_ZEBRADO, fill_type="solid") if idx % 2 == 1 else None

        for col_idx, (val, fmt, alin) in enumerate(valores, start=1):
            cell = ws.cell(row=r_num, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal=alin, vertical="center")
            cell.border = borda_padrao()
            if fill_cor:
                cell.fill = fill_cor

    ajustar_larguras(ws)
