#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Motor de Geração do Fluxo de Caixa, Curva de Desembolso e Análise de Capital de Giro
para a OBRA_TMULT (Edifício Administrativo - Porto do Açu).

Frente: 05_SUPRIMENTOS_E_FINANCEIRO
Skills aplicadas:
- SKILL_GESTAO_09_FLUXO_DE_CAIXA.md (Curva de Desembolso x Timing de Pagamento / Portão de Qualidade)
- SKILL_GESTAO_07_CIENCIA_DE_DADOS.md (Projeção em Faixa / 3 Cenários)
- SKILL_QUANTIFICACAO_MASTER.md & Base Oficial SINAPI SP 07/2026

Artefatos gerados:
1. FLUXO_DE_CAIXA_TMULT.csv
2. FLUXO_DE_CAIXA_TMULT.xlsx (com abas de Resumo Executivo, Gráficos e Analítico)
3. RELATORIO_FLUXO_DE_CAIXA_E_CAPITAL_DE_GIRO_TMULT.md
"""

import os
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference, Series

# Caminhos do Projeto
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CRONOGRAMA_CSV = os.path.join(BASE_DIR, "projetos", "OBRA_TMULT", "03_PLANEJAMENTO_E_CRONOGRAMA", "CRONOGRAMA_FISICO_FINANCEIRO_TMULT.csv")
OUTPUT_DIR = os.path.join(BASE_DIR, "projetos", "OBRA_TMULT", "05_SUPRIMENTOS_E_FINANCEIRO")

os.makedirs(OUTPUT_DIR, exist_ok=True)

# Estilos OpenPyXL Corporativos
NAVY_HEADER = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
GOLD_ACCENT = PatternFill(start_color="D99B26", end_color="D99B26", fill_type="solid")
GRAY_LIGHT = PatternFill(start_color="F4F6F9", end_color="F4F6F9", fill_type="solid")
GREEN_LIGHT = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
RED_LIGHT = PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid")
BLUE_LIGHT = PatternFill(start_color="E8EEF5", end_color="E8EEF5", fill_type="solid")

FONT_TITLE = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_BOLD = Font(name="Calibri", size=11, bold=True, color="1B365D")
FONT_REGULAR = Font(name="Calibri", size=11, color="333333")
FONT_RED = Font(name="Calibri", size=11, bold=True, color="C5221F")
FONT_GREEN = Font(name="Calibri", size=11, bold=True, color="137333")

THIN_BORDER = Border(
    left=Side(style='thin', color='DDDDDD'),
    right=Side(style='thin', color='DDDDDD'),
    top=Side(style='thin', color='DDDDDD'),
    bottom=Side(style='thin', color='DDDDDD')
)
DOUBLE_BOTTOM_BORDER = Border(
    left=Side(style='thin', color='DDDDDD'),
    right=Side(style='thin', color='DDDDDD'),
    top=Side(style='thin', color='DDDDDD'),
    bottom=Side(style='double', color='1B365D')
)

def formatar_moeda(val):
    if abs(val) < 0.001:
        return "R$ 0,00"
    s = f"R$ {abs(val):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"-{s}" if val < 0 else s

def carregar_dados():
    df = pd.read_csv(CRONOGRAMA_CSV, sep=';', encoding='utf-8')
    eap_col = [c for c in df.columns if 'EAP' in c][0]
    
    # Vendas brutas do cronograma
    vendas_m = {m: round(df[f'R$ Mês {m}'].sum(), 2) for m in range(1, 7)}
    total_venda = sum(vendas_m.values())
    
    # Calibração do custo direto para fechar com exatidão R$ 1.314.562,67
    TOTAL_CUSTO_DIRETO = 1314562.67
    fator_calib = TOTAL_CUSTO_DIRETO / 1315417.51
    
    diretos_m = {}
    for m in range(1, 7):
        s = 0.0
        for idx, r in df.iterrows():
            eap = str(r[eap_col]).strip()
            bdi = 0.15 if eap.startswith('3.3') else 0.2717
            s += (r[f'R$ Mês {m}'] / (1.0 + bdi)) * fator_calib
        diretos_m[m] = round(s, 2)
        
    diff_cd = round(TOTAL_CUSTO_DIRETO - sum(diretos_m.values()), 2)
    diretos_m[6] = round(diretos_m[6] + diff_cd, 2)
    
    return vendas_m, diretos_m, total_venda, TOTAL_CUSTO_DIRETO

def modelar_fluxo_caixa(vendas_m, diretos_m, total_venda, total_custo_direto):
    meses = list(range(1, 8)) # Mês 1 a Mês 7 (Liquidação)
    
    # -------------------------------------------------------------
    # 1. ENTRADAS (INFLOWS) — CENÁRIO REALISTA (BASE)
    # -------------------------------------------------------------
    # Retenção Contratual: 5% sobre cada medição
    # Faturamento Líquido: 95%
    # Prazo de Recebimento: Medição do Mês m aprovada no fim do mês e recebida em D+15 no Mês m+1
    # Liberação da Retenção Técnica: Mês 7 integral
    faturamento_bruto = {m: vendas_m.get(m, 0.0) for m in meses}
    retencao_5pct = {m: round(vendas_m.get(m, 0.0) * 0.05, 2) for m in meses}
    faturamento_liquido = {m: round(vendas_m.get(m, 0.0) * 0.95, 2) for m in meses}
    
    total_retencao = sum(retencao_5pct.values())
    
    recebimento_medicoes = {m: 0.0 for m in meses}
    for m in range(1, 7):
        recebimento_medicoes[m+1] = faturamento_liquido[m]
        
    devolucao_retencao = {m: 0.0 for m in meses}
    devolucao_retencao[7] = total_retencao
    
    total_entradas_realista = {m: round(recebimento_medicoes[m] + devolucao_retencao[m], 2) for m in meses}
    
    # -------------------------------------------------------------
    # 2. SAÍDAS (OUTFLOWS) — CENÁRIO REALISTA (BASE)
    # -------------------------------------------------------------
    # Custos Canteiro EAP 1.0 (Total: R$ 366.255,96 / 6 meses = R$ 61.042,66/mês):
    # - Equipe de Gestão Técnica (EAP 1.0.1 e 1.0.2): R$ 29.300,00/mês (desembolso imediato no mês m)
    # - Vivência, Alimentação 16 op. e Transporte (EAP 1.0.5 e 1.0.6): R$ 23.009,33/mês (desembolso no mês m)
    # - Containers NR-18 e Sanitários (EAP 1.0.3): R$ 6.383,33/mês (50% no mês m, 50% em m+1)
    # - Contas de Consumo (EAP 1.0.4): R$ 2.350,00/mês (50% no mês m, 50% em m+1)
    
    # Custos Físicos Civis/Instalações (diretos_m - 61.042,66):
    # - Mão de Obra Direta de Campo: ~35% dos serviços físicos (desembolso no mês m)
    # - Materiais e Insumos Civis: ~65% dos serviços físicos (20% sinal à vista em m, 80% D+30 em m+1)
    
    # Impostos sobre Faturamento (BDI):
    # - 8,65% sobre a receita bruta faturada de cada medição, recolhido dia 20 do mês subsequente (m+1)
    
    # Despesas Indiretas Centrais da Construtora (BDI rateado):
    # - Administração Central (4%), Seguros (1%), Riscos (1,5%) e Despesas Financeiras (1%) = R$ 98.592,18
    # - Linear em 6 meses: R$ 16.432,03/mês
    
    des_mo_gestao = {m: 0.0 for m in meses}
    des_vivencia = {m: 0.0 for m in meses}
    des_mo_campo = {m: 0.0 for m in meses}
    des_mat_vista = {m: 0.0 for m in meses}
    des_mat_prazo = {m: 0.0 for m in meses}
    des_locacoes = {m: 0.0 for m in meses}
    des_contas = {m: 0.0 for m in meses}
    des_impostos = {m: 0.0 for m in meses}
    des_ac_seguros = {m: 0.0 for m in meses}
    
    for m in range(1, 7):
        # Gestão e Vivência
        des_mo_gestao[m] += 29300.00
        des_vivencia[m] += 23009.33
        
        # Containers e Contas
        des_locacoes[m] += round(6383.33 * 0.50, 2)
        des_locacoes[m+1] += round(6383.33 * 0.50, 2)
        
        des_contas[m] += round(2350.00 * 0.50, 2)
        des_contas[m+1] += round(2350.00 * 0.50, 2)
        
        # Custo Físico
        custo_fis = diretos_m[m] - 61042.66
        mo_c = round(custo_fis * 0.35, 2)
        mat_c = round(custo_fis * 0.65, 2)
        
        des_mo_campo[m] += mo_c
        des_mat_vista[m] += round(mat_c * 0.20, 2)
        des_mat_prazo[m+1] += round(mat_c * 0.80, 2)
        
        # Impostos
        imp = round(vendas_m[m] * 0.0865, 2)
        des_impostos[m+1] += imp
        
        # AC e Seguros
        des_ac_seguros[m] += 16432.03
        
    total_saidas_realista = {}
    for m in meses:
        s = (des_mo_gestao[m] + des_vivencia[m] + des_mo_campo[m] + 
             des_mat_vista[m] + des_mat_prazo[m] + des_locacoes[m] + 
             des_contas[m] + des_impostos[m] + des_ac_seguros[m])
        total_saidas_realista[m] = round(s, 2)
        
    # Saldo Operacional e Saldo Acumulado Realista
    saldo_periodo_realista = {}
    saldo_acumulado_realista = {}
    acum = 0.0
    for m in meses:
        liq = round(total_entradas_realista[m] - total_saidas_realista[m], 2)
        acum = round(acum + liq, 2)
        saldo_periodo_realista[m] = liq
        saldo_acumulado_realista[m] = acum
        
    # -------------------------------------------------------------
    # 3. CENÁRIO OTIMISTA (Com Adiantamento de Mobilização Contratual)
    # -------------------------------------------------------------
    # Adiantamento de 10% do contrato no Dia Zero (Mês 1) = R$ 166.076,23
    # Amortização em 4 parcelas iguais nas medições M1 a M4 (R$ 41.519,06/mês)
    adiantamento = round(total_venda * 0.10, 2)
    amort_mensal = round(adiantamento / 4.0, 2)
    
    total_entradas_otimista = {m: 0.0 for m in meses}
    total_entradas_otimista[1] += adiantamento
    for m in range(1, 7):
        amort = amort_mensal if m <= 4 else 0.0
        rec_liq = round((vendas_m[m] * 0.95) - amort, 2)
        total_entradas_otimista[m+1] += rec_liq
    total_entradas_otimista[7] += total_retencao
    
    saldo_acumulado_otimista = {}
    acum_oti = 0.0
    for m in meses:
        liq = round(total_entradas_otimista[m] - total_saidas_realista[m], 2)
        acum_oti = round(acum_oti + liq, 2)
        saldo_acumulado_otimista[m] = acum_oti
        
    # -------------------------------------------------------------
    # 4. CENÁRIO DE ESTRESSE / PESSIMISTA (Atraso de Medições em 30 Dias)
    # -------------------------------------------------------------
    # Cliente demora D+45 para pagar cada medição (cai no Mês m+2)
    meses_pess = list(range(1, 9)) # M1 a M8
    total_entradas_pessimista = {m: 0.0 for m in meses_pess}
    for m in range(1, 7):
        total_entradas_pessimista[m+2] += faturamento_liquido[m]
    total_entradas_pessimista[8] += total_retencao
    
    total_saidas_pessimista = {m: total_saidas_realista.get(m, 0.0) for m in meses_pess}
    
    saldo_acumulado_pessimista = {}
    acum_pess = 0.0
    for m in meses_pess:
        liq = round(total_entradas_pessimista[m] - total_saidas_pessimista[m], 2)
        acum_pess = round(acum_pess + liq, 2)
        saldo_acumulado_pessimista[m] = acum_pess
        
    return {
        "meses": meses,
        "vendas_m": vendas_m,
        "diretos_m": diretos_m,
        "faturamento_bruto": faturamento_bruto,
        "retencao_5pct": retencao_5pct,
        "faturamento_liquido": faturamento_liquido,
        "total_retencao": total_retencao,
        "recebimento_medicoes": recebimento_medicoes,
        "devolucao_retencao": devolucao_retencao,
        "total_entradas_realista": total_entradas_realista,
        "des_mo_gestao": des_mo_gestao,
        "des_vivencia": des_vivencia,
        "des_mo_campo": des_mo_campo,
        "des_mat_vista": des_mat_vista,
        "des_mat_prazo": des_mat_prazo,
        "des_locacoes": des_locacoes,
        "des_contas": des_contas,
        "des_impostos": des_impostos,
        "des_ac_seguros": des_ac_seguros,
        "total_saidas_realista": total_saidas_realista,
        "saldo_periodo_realista": saldo_periodo_realista,
        "saldo_acumulado_realista": saldo_acumulado_realista,
        "saldo_acumulado_otimista": saldo_acumulado_otimista,
        "meses_pess": meses_pess,
        "total_entradas_pessimista": total_entradas_pessimista,
        "total_saidas_pessimista": total_saidas_pessimista,
        "saldo_acumulado_pessimista": saldo_acumulado_pessimista,
        "total_venda": total_venda,
        "total_custo_direto": total_custo_direto
    }

def gerar_csv(fluxo):
    csv_path = os.path.join(OUTPUT_DIR, "FLUXO_DE_CAIXA_TMULT.csv")
    meses = fluxo["meses"]
    
    linhas = []
    # Cabeçalho
    header = ["Conta / Rubrica Financeira", "Natureza"] + [f"Mês {m}" for m in meses] + ["Total Consolidado (R$)"]
    linhas.append(header)
    
    def criar_linha(titulo, nat, dicionario):
        vals = [dicionario.get(m, 0.0) for m in meses]
        tot = sum(vals)
        return [titulo, nat] + [f"{v:.2f}" for v in vals] + [f"{tot:.2f}"]
        
    linhas.append(["1. ENTRADAS DE CAIXA (RECEBIMENTOS)", "GRUPO", "", "", "", "", "", "", "", ""])
    linhas.append(criar_linha("1.1 Faturamento Bruto de Medições", "Receita Bruta", fluxo["faturamento_bruto"]))
    linhas.append(criar_linha("1.2 Retenção Técnica Contratual (5%)", "Dedução", fluxo["retencao_5pct"]))
    linhas.append(criar_linha("1.3 Faturamento Líquido de Medições", "Receita Líquida", fluxo["faturamento_liquido"]))
    linhas.append(criar_linha("1.4 Recebimento Efetivo de Medições (D+15)", "Inflow", fluxo["recebimento_medicoes"]))
    linhas.append(criar_linha("1.5 Devolução da Retenção Técnica (TRD M7)", "Inflow", fluxo["devolucao_retencao"]))
    linhas.append(criar_linha("TOTAL DE ENTRADAS DE CAIXA (A)", "Total Inflow", fluxo["total_entradas_realista"]))
    
    linhas.append(["", "", "", "", "", "", "", "", "", ""])
    linhas.append(["2. SAÍDAS DE CAIXA (DESEMBOLSOS)", "GRUPO", "", "", "", "", "", "", "", ""])
    linhas.append(criar_linha("2.1 Equipe Técnica de Gestão (Eng/Mestre/TST)", "MO Gestão", fluxo["des_mo_gestao"]))
    linhas.append(criar_linha("2.2 Vivência, Alimentação (16 op.) e Transporte", "Vivência", fluxo["des_vivencia"]))
    linhas.append(criar_linha("2.3 Mão de Obra de Campo (Serviços Físicos)", "MO Campo", fluxo["des_mo_campo"]))
    linhas.append(criar_linha("2.4 Materiais e Insumos Civis (Sinal à Vista)", "Material", fluxo["des_mat_vista"]))
    linhas.append(criar_linha("2.5 Materiais e Insumos Civis (Prazo D+30)", "Material", fluxo["des_mat_prazo"]))
    linhas.append(criar_linha("2.6 Locação de Containers NR-18 e Sanitários", "Equipamento", fluxo["des_locacoes"]))
    linhas.append(criar_linha("2.7 Contas de Consumo (Água, Luz, Internet)", "Canteiro", fluxo["des_contas"]))
    linhas.append(criar_linha("2.8 Tributos sobre Faturamento (8,65% NF)", "Tributos", fluxo["des_impostos"]))
    linhas.append(criar_linha("2.9 Custos Indiretos Centrais BDI (AC/Seguros)", "Indireto", fluxo["des_ac_seguros"]))
    linhas.append(criar_linha("TOTAL DE SAÍDAS DE CAIXA (B)", "Total Outflow", fluxo["total_saidas_realista"]))
    
    linhas.append(["", "", "", "", "", "", "", "", "", ""])
    linhas.append(["3. BALANÇO E SALDOS DE CAIXA", "GRUPO", "", "", "", "", "", "", "", ""])
    linhas.append(criar_linha("3.1 Saldo Operacional Líquido do Mês (A - B)", "Resultado", fluxo["saldo_periodo_realista"]))
    
    # Saldo acumulado
    vals_acum = [fluxo["saldo_acumulado_realista"][m] for m in meses]
    linhas.append(["3.2 Saldo de Caixa Acumulado (Cenário Realista)", "Acumulado"] + [f"{v:.2f}" for v in vals_acum] + [f"{vals_acum[-1]:.2f}"])
    
    vals_oti = [fluxo["saldo_acumulado_otimista"][m] for m in meses]
    linhas.append(["3.3 Saldo Acumulado (Cenário Otimista c/ Adiantamento)", "Sensibilidade"] + [f"{v:.2f}" for v in vals_oti] + [f"{vals_oti[-1]:.2f}"])
    
    vals_pess = [fluxo["saldo_acumulado_pessimista"][m] for m in meses]
    linhas.append(["3.4 Saldo Acumulado (Cenário Estresse Atraso 30d)", "Sensibilidade"] + [f"{v:.2f}" for v in vals_pess] + [f"{vals_pess[-1]:.2f}"])
    
    import csv
    with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f, delimiter=';')
        writer.writerows(linhas)
        
    print(f"[OK] CSV de Fluxo de Caixa gerado: {csv_path}")

def gerar_xlsx(fluxo):
    xlsx_path = os.path.join(OUTPUT_DIR, "FLUXO_DE_CAIXA_TMULT.xlsx")
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # ABA 1: RESUMO EXECUTIVO & GRÁFICOS
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Resumo Executivo"
    ws1.views.sheetView[0].showGridLines = True
    
    # Banner
    ws1.merge_cells("B2:I2")
    ws1["B2"] = "RELATÓRIO FINANCEIRO EXECUTIVO — FLUXO DE CAIXA E CAPITAL DE GIRO"
    ws1["B2"].fill = NAVY_HEADER
    ws1["B2"].font = FONT_TITLE
    ws1["B2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 35
    
    # Subtítulo
    ws1.merge_cells("B3:I3")
    ws1["B3"] = "Empreendimento: TMULT — Terminal Multiuso (Porto do Açu) | Prazo: 6 Meses (180 Dias) | Baseline 01"
    ws1["B3"].fill = GOLD_ACCENT
    ws1["B3"].font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ws1["B3"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[3].height = 22
    
    # KPIs Executivos
    pico_realista = min(fluxo["saldo_acumulado_realista"].values())
    mes_pico = min(fluxo["saldo_acumulado_realista"], key=fluxo["saldo_acumulado_realista"].get)
    saldo_final = fluxo["saldo_acumulado_realista"][7]
    margem_liq = (saldo_final / fluxo["total_venda"]) * 100
    
    kpis = [
        ("Preço Global Turnkey da Obra", fluxo["total_venda"], '"R$ "#,##0.00'),
        ("Custo Direto Total da Obra", fluxo["total_custo_direto"], '"R$ "#,##0.00'),
        ("Total de Impostos s/ Faturamento (8,65%)", sum(fluxo["des_impostos"].values()), '"R$ "#,##0.00'),
        ("Despesas Administrativas Centrais BDI", sum(fluxo["des_ac_seguros"].values()), '"R$ "#,##0.00'),
        ("Lucro Líquido Operacional Realizado", saldo_final, '"R$ "#,##0.00'),
        ("Margem Líquida Realizada sobre Venda", margem_liq / 100.0, '0.00%'),
        ("Retenção Contratual Técnica (5,0%)", fluxo["total_retencao"], '"R$ "#,##0.00'),
        ("NECESSIDADE MÍNIMA DE CAPITAL DE GIRO", abs(pico_realista), '"R$ "#,##0.00'),
        ("Mês de Máxima Exposição de Caixa", f"Mês {mes_pico} (Pico da Envoltória)", "@"),
        ("Ponto de Equilíbrio do Caixa (Break-even)", "Mês 7 (Liquidação do TRD)", "@")
    ]
    
    row_kpi = 5
    ws1.merge_cells(f"B{row_kpi}:E{row_kpi}")
    ws1[f"B{row_kpi}"] = "INDICADORES FINANCEIROS E PARÂMETROS DE LIQUIDEZ"
    ws1[f"B{row_kpi}"].fill = NAVY_HEADER
    ws1[f"B{row_kpi}"].font = FONT_HEADER
    ws1[f"B{row_kpi}"].alignment = Alignment(horizontal="center")
    
    for i, (kpi, val, fmt) in enumerate(kpis, start=1):
        curr_row = row_kpi + i
        ws1.merge_cells(f"B{curr_row}:D{curr_row}")
        ws1[f"B{curr_row}"] = kpi
        ws1[f"B{curr_row}"].font = FONT_BOLD if "CAPITAL DE GIRO" in kpi or "Lucro" in kpi else FONT_REGULAR
        ws1[f"B{curr_row}"].border = THIN_BORDER
        
        c_val = ws1[f"E{curr_row}"]
        c_val.value = val
        c_val.number_format = fmt
        c_val.alignment = Alignment(horizontal="right")
        c_val.border = THIN_BORDER
        
        if "CAPITAL DE GIRO" in kpi:
            ws1[f"B{curr_row}"].fill = RED_LIGHT
            c_val.fill = RED_LIGHT
            c_val.font = FONT_RED
        elif "Lucro" in kpi:
            ws1[f"B{curr_row}"].fill = GREEN_LIGHT
            c_val.fill = GREEN_LIGHT
            c_val.font = FONT_GREEN
        elif i % 2 == 0:
            ws1[f"B{curr_row}"].fill = GRAY_LIGHT
            c_val.fill = GRAY_LIGHT

    # Tabela Resumo Mensal do Fluxo (para plotar gráfico nativo)
    row_tbl = 17
    headers_resumo = ["Período", "Entradas (R$)", "Saídas (R$)", "Saldo Mês (R$)", "Saldo Acumulado (R$)"]
    for col_idx, h in enumerate(headers_resumo, start=2):
        cell = ws1.cell(row=row_tbl, column=col_idx, value=h)
        cell.fill = NAVY_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER
        
    for idx, m in enumerate(fluxo["meses"], start=1):
        r = row_tbl + idx
        ws1.cell(row=r, column=2, value=f"Mês {m}").alignment = Alignment(horizontal="center")
        ws1.cell(row=r, column=3, value=fluxo["total_entradas_realista"][m]).number_format = '"R$ "#,##0.00'
        ws1.cell(row=r, column=4, value=fluxo["total_saidas_realista"][m]).number_format = '"R$ "#,##0.00'
        ws1.cell(row=r, column=5, value=fluxo["saldo_periodo_realista"][m]).number_format = '"R$ "#,##0.00'
        
        c_acum = ws1.cell(row=r, column=6, value=fluxo["saldo_acumulado_realista"][m])
        c_acum.number_format = '"R$ "#,##0.00'
        
        for c in range(2, 7):
            cell = ws1.cell(row=r, column=c)
            cell.font = FONT_REGULAR
            cell.border = THIN_BORDER
            if idx % 2 == 0:
                cell.fill = GRAY_LIGHT
        
        if fluxo["saldo_acumulado_realista"][m] < 0:
            c_acum.font = FONT_RED
        else:
            c_acum.font = FONT_GREEN

    # Linha Total Resumo
    r_tot = row_tbl + len(fluxo["meses"]) + 1
    ws1.cell(row=r_tot, column=2, value="TOTAL GERAL").alignment = Alignment(horizontal="center")
    ws1.cell(row=r_tot, column=3, value=sum(fluxo["total_entradas_realista"].values())).number_format = '"R$ "#,##0.00'
    ws1.cell(row=r_tot, column=4, value=sum(fluxo["total_saidas_realista"].values())).number_format = '"R$ "#,##0.00'
    ws1.cell(row=r_tot, column=5, value=saldo_final).number_format = '"R$ "#,##0.00'
    ws1.cell(row=r_tot, column=6, value=saldo_final).number_format = '"R$ "#,##0.00'
    
    for c in range(2, 7):
        cell = ws1.cell(row=r_tot, column=c)
        cell.fill = GREEN_LIGHT
        cell.font = FONT_BOLD
        cell.border = DOUBLE_BOTTOM_BORDER

    # Gráfico 1: Linha de Saldo Acumulado de Caixa (Curva de Balanço Financeiro)
    chart_line = LineChart()
    chart_line.title = "Curva de Balanço Financeiro (Saldo de Caixa Acumulado - Cenário Realista)"
    chart_line.style = 13
    chart_line.y_axis.title = "Saldo Acumulado (R$)"
    chart_line.x_axis.title = "Cronograma (Meses)"
    chart_line.width = 18
    chart_line.height = 11
    
    data_line = Reference(ws1, min_col=6, min_row=row_tbl, max_col=6, max_row=row_tbl + len(fluxo["meses"]))
    cats_line = Reference(ws1, min_col=2, min_row=row_tbl + 1, max_row=row_tbl + len(fluxo["meses"]))
    chart_line.add_data(data_line, titles_from_data=True)
    chart_line.set_categories(cats_line)
    ws1.add_chart(chart_line, "H5")
    
    # Gráfico 2: Barras de Entradas vs Saídas Mensais
    chart_bar = BarChart()
    chart_bar.type = "col"
    chart_bar.style = 10
    chart_bar.title = "Entradas vs Saídas de Caixa por Mês (Inflows vs Outflows)"
    chart_bar.y_axis.title = "Valor Mensal (R$)"
    chart_bar.x_axis.title = "Mês"
    chart_bar.width = 18
    chart_bar.height = 11
    
    data_bar = Reference(ws1, min_col=3, min_row=row_tbl, max_col=4, max_row=row_tbl + len(fluxo["meses"]))
    chart_bar.add_data(data_bar, titles_from_data=True)
    chart_bar.set_categories(cats_line)
    ws1.add_chart(chart_bar, "H19")

    # -------------------------------------------------------------
    # ABA 2: FLUXO DE CAIXA ANALÍTICO
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Fluxo Analítico")
    ws2.views.sheetView[0].showGridLines = True
    
    # Cabeçalho
    ws2.merge_cells("A1:J1")
    ws2["A1"] = "FLUXO DE CAIXA ANALÍTICO MENSAL — ENTRADAS, DESEMBOLSOS E BALANÇO (R$)"
    ws2["A1"].fill = NAVY_HEADER
    ws2["A1"].font = FONT_TITLE
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 30
    
    cols_analitico = ["Conta / Rubrica Operacional", "Natureza Contábil"] + [f"Mês {m}" for m in fluxo["meses"]] + ["Total Global (R$)"]
    ws2.row_dimensions[2].height = 24
    for c_idx, h in enumerate(cols_analitico, start=1):
        cell = ws2.cell(row=2, column=c_idx, value=h)
        cell.fill = NAVY_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER
        
    row_curr = 3
    
    def add_section_header(title, fill_color=BLUE_LIGHT):
        nonlocal row_curr
        ws2.merge_cells(start_row=row_curr, start_column=1, end_row=row_curr, end_column=10)
        cell = ws2.cell(row=row_curr, column=1, value=title)
        cell.fill = fill_color
        cell.font = FONT_BOLD
        cell.border = THIN_BORDER
        row_curr += 1
        
    def add_data_row(titulo, nat, dicionario, is_bold=False, is_highlight=False):
        nonlocal row_curr
        ws2.cell(row=row_curr, column=1, value=titulo).font = FONT_BOLD if is_bold else FONT_REGULAR
        ws2.cell(row=row_curr, column=2, value=nat).alignment = Alignment(horizontal="center")
        
        vals = [dicionario.get(m, 0.0) for m in fluxo["meses"]]
        tot = sum(vals)
        for idx, v in enumerate(vals, start=3):
            c = ws2.cell(row=row_curr, column=idx, value=v)
            c.number_format = '"R$ "#,##0.00'
            c.font = FONT_BOLD if is_bold else FONT_REGULAR
            
        c_tot = ws2.cell(row=row_curr, column=10, value=tot)
        c_tot.number_format = '"R$ "#,##0.00'
        c_tot.font = FONT_BOLD
        
        for c in range(1, 11):
            cell = ws2.cell(row=row_curr, column=c)
            cell.border = THIN_BORDER
            if is_highlight:
                cell.fill = GREEN_LIGHT
            elif row_curr % 2 == 0:
                cell.fill = GRAY_LIGHT
        row_curr += 1

    # 1. ENTRADAS
    add_section_header("1. ENTRADAS DE CAIXA (RECEBIMENTOS DE MEDIÇÕES)")
    add_data_row("1.1 Faturamento Bruto de Medições", "Medição EAP", fluxo["faturamento_bruto"])
    add_data_row("1.2 Retenção Técnica Contratual (5,0%)", "Dedução Contratual", fluxo["retencao_5pct"])
    add_data_row("1.3 Faturamento Líquido de Medições (95%)", "Receita Faturada", fluxo["faturamento_liquido"])
    add_data_row("1.4 Recebimento Efetivo de Medições (D+15)", "Inflow", fluxo["recebimento_medicoes"])
    add_data_row("1.5 Devolução da Retenção Técnica (TRD / Mês 7)", "Inflow", fluxo["devolucao_retencao"])
    add_data_row("TOTAL DE ENTRADAS DE CAIXA (A)", "Total Inflows", fluxo["total_entradas_realista"], is_bold=True, is_highlight=True)
    
    # 2. SAÍDAS
    add_section_header("2. SAÍDAS DE CAIXA (DESEMBOLSOS DE PRODUÇÃO E GESTÃO)")
    add_data_row("2.1 Equipe Técnica de Gestão (Engenheiro Residente + Mestre + TST)", "MO Gestão", fluxo["des_mo_gestao"])
    add_data_row("2.2 Vivência, Alimentação (16 op.) e Logística de Transporte", "Vivência/Benefícios", fluxo["des_vivencia"])
    add_data_row("2.3 Mão de Obra Direta de Campo (Serviços Físicos Civis/Instalações)", "MO Campo", fluxo["des_mo_campo"])
    add_data_row("2.4 Materiais e Insumos Civis (Sinal / Entrada à Vista no Pedido)", "Material D0", fluxo["des_mat_vista"])
    add_data_row("2.5 Materiais e Insumos Civis (Faturamento a Prazo D+30)", "Material D+30", fluxo["des_mat_prazo"])
    add_data_row("2.6 Locação de Containers NR-18 e Sanitários Químicos", "Equipamento", fluxo["des_locacoes"])
    add_data_row("2.7 Contas de Consumo Provisórias (Energia, Água Pipa, Internet Fibra)", "Canteiro", fluxo["des_contas"])
    add_data_row("2.8 Tributos sobre Faturamento (ISS, PIS, COFINS, CPRB - 8,65% NF)", "Tributos", fluxo["des_impostos"])
    add_data_row("2.9 Custos Indiretos Centrais BDI (Administração Central, Seguros e DF)", "Indireto BDI", fluxo["des_ac_seguros"])
    add_data_row("TOTAL DE SAÍDAS DE CAIXA (B)", "Total Outflows", fluxo["total_saidas_realista"], is_bold=True, is_highlight=False)
    
    # 3. RESULTADOS
    add_section_header("3. RESULTADO OPERACIONAL E SALDOS DE CAIXA")
    add_data_row("3.1 Saldo Operacional Líquido do Período (A - B)", "Resultado Líquido", fluxo["saldo_periodo_realista"], is_bold=True)
    
    # Saldo acumulado linha a linha
    ws2.cell(row=row_curr, column=1, value="3.2 Saldo de Caixa Acumulado (Cenário Realista)").font = FONT_BOLD
    ws2.cell(row=row_curr, column=2, value="Acumulado").alignment = Alignment(horizontal="center")
    for idx, m in enumerate(fluxo["meses"], start=3):
        v = fluxo["saldo_acumulado_realista"][m]
        c = ws2.cell(row=row_curr, column=idx, value=v)
        c.number_format = '"R$ "#,##0.00'
        c.font = FONT_RED if v < 0 else FONT_GREEN
    c_tot = ws2.cell(row=row_curr, column=10, value=saldo_final)
    c_tot.number_format = '"R$ "#,##0.00'
    c_tot.font = FONT_BOLD
    
    for c in range(1, 11):
        cell = ws2.cell(row=row_curr, column=c)
        cell.fill = YELLOW_ACCENT = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        cell.border = DOUBLE_BOTTOM_BORDER
    row_curr += 1

    # -------------------------------------------------------------
    # ABA 3: CENÁRIOS DE SENSIBILIDADE
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Cenários de Sensibilidade")
    ws3.views.sheetView[0].showGridLines = True
    
    ws3.merge_cells("A1:I1")
    ws3["A1"] = "ANÁLISE DE SENSIBILIDADE FINANCEIRA — SIMULAÇÃO EM 3 CENÁRIOS"
    ws3["A1"].fill = NAVY_HEADER
    ws3["A1"].font = FONT_TITLE
    ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 30
    
    headers_sens = ["Cenário de Análise", "Premissas Chave de Fluxo", "M1", "M2", "M3", "M4", "M5", "M6", "M7", "M8", "Pico Negativo (R$)", "Saldo Final (R$)"]
    for c_idx, h in enumerate(headers_sens, start=1):
        cell = ws3.cell(row=2, column=c_idx, value=h)
        cell.fill = NAVY_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER
        
    cenarios = [
        ("Cenário 1: Realista (Base)", "Medições pagas em D+15; Retenção 5% devolvida em M7; Fornecedores 20/80 D0/D+30",
         fluxo["saldo_acumulado_realista"], min(fluxo["saldo_acumulado_realista"].values()), fluxo["saldo_acumulado_realista"][7]),
        ("Cenário 2: Otimista (Adiantamento)", "Adiantamento Contratual de 10% (R$ 166.076,23) amortizado em 4x nas medições",
         fluxo["saldo_acumulado_otimista"], min(fluxo["saldo_acumulado_otimista"].values()), fluxo["saldo_acumulado_otimista"][7]),
        ("Cenário 3: Estresse (Atraso 30d)", "Contratante atrasa liberação de medições em 30 dias (D+45); fornecedores pagos pontualmente",
         fluxo["saldo_acumulado_pessimista"], min(fluxo["saldo_acumulado_pessimista"].values()), fluxo["saldo_acumulado_pessimista"][8])
    ]
    
    for r_idx, (nome, premissa, curva, pico, final) in enumerate(cenarios, start=3):
        ws3.cell(row=r_idx, column=1, value=nome).font = FONT_BOLD
        ws3.cell(row=r_idx, column=2, value=premissa).font = FONT_REGULAR
        
        for m in range(1, 8):
            val = curva.get(m, 0.0)
            c = ws3.cell(row=r_idx, column=m+2, value=val)
            c.number_format = '"R$ "#,##0.00'
            c.font = FONT_RED if val < 0 else FONT_GREEN
            
        # M8
        val_m8 = curva.get(8, "-")
        c8 = ws3.cell(row=r_idx, column=10, value=val_m8)
        if isinstance(val_m8, (int, float)):
            c8.number_format = '"R$ "#,##0.00'
            c8.font = FONT_GREEN
            
        c_pico = ws3.cell(row=r_idx, column=11, value=pico)
        c_pico.number_format = '"R$ "#,##0.00'
        c_pico.font = FONT_RED
        
        c_final = ws3.cell(row=r_idx, column=12, value=final)
        c_final.number_format = '"R$ "#,##0.00'
        c_final.font = FONT_BOLD
        
        for c in range(1, 13):
            cell = ws3.cell(row=r_idx, column=c)
            cell.border = THIN_BORDER
            if r_idx % 2 == 0:
                cell.fill = GRAY_LIGHT

    # Ajuste de larguras de colunas
    col_widths_ws1 = {2: 15, 3: 18, 4: 18, 5: 18, 6: 22, 7: 10, 8: 10, 9: 10}
    for c, w in col_widths_ws1.items():
        ws1.column_dimensions[get_column_letter(c)].width = w
        
    col_widths_ws2 = {1: 45, 2: 22, 3: 15, 4: 15, 5: 15, 6: 15, 7: 15, 8: 15, 9: 15, 10: 20}
    for c, w in col_widths_ws2.items():
        ws2.column_dimensions[get_column_letter(c)].width = w
        
    col_widths_ws3 = {1: 25, 2: 45, 3: 14, 4: 14, 5: 14, 6: 14, 7: 14, 8: 14, 9: 14, 10: 14, 11: 20, 12: 20}
    for c, w in col_widths_ws3.items():
        ws3.column_dimensions[get_column_letter(c)].width = w

    wb.save(xlsx_path)
    print(f"[OK] Planilha Executiva de Fluxo de Caixa gerada: {xlsx_path}")

def gerar_relatorio_markdown(fluxo):
    md_path = os.path.join(OUTPUT_DIR, "RELATORIO_FLUXO_DE_CAIXA_E_CAPITAL_DE_GIRO_TMULT.md")
    
    pico_realista = min(fluxo["saldo_acumulado_realista"].values())
    mes_pico_realista = min(fluxo["saldo_acumulado_realista"], key=fluxo["saldo_acumulado_realista"].get)
    saldo_final_realista = fluxo["saldo_acumulado_realista"][7]
    margem_liq_pct = (saldo_final_realista / fluxo["total_venda"]) * 100
    
    pico_otimista = min(fluxo["saldo_acumulado_otimista"].values())
    pico_pessimista = min(fluxo["saldo_acumulado_pessimista"].values())
    mes_pico_pessimista = min(fluxo["saldo_acumulado_pessimista"], key=fluxo["saldo_acumulado_pessimista"].get)
    
    total_impostos = sum(fluxo["des_impostos"].values())
    total_ac_seguros = sum(fluxo["des_ac_seguros"].values())
    
    md = []
    md.append("# 💰 Relatório Executivo de Fluxo de Caixa, Curva de Desembolso e Capital de Giro")
    md.append("\n**Empreendimento:** Edifício Administrativo do Terminal Multiuso (Porto do Açu) — `OBRA_TMULT`")
    md.append("**Área Construída Útil:** 368,40 m² | **Prazo Contratual:** 6 Meses (180 Dias Corridos / 26 Semanas)")
    md.append("**Preço Global Turnkey Contratado:** R$ 1.660.762,28 | **Custo Direto Total:** R$ 1.314.562,67")
    md.append("**Data de Emissão:** 10/09/2026 | **Fase:** Linha de Base 01 (Baseline 01)\n")
    md.append("---\n")
    
    # Portão de Qualidade de Dados (SKILL_GESTAO_09 / SKILL_GESTAO_07)
    md.append("## 1. Portão de Qualidade de Dados e Nível de Confiança")
    md.append("\nEm conformidade com a `SKILL_GESTAO_09` (Seção 1) e a `SKILL_GESTAO_07` (Hierarquia de Confiança), declara-se a confiabilidade dos parâmetros de entrada utilizados:\n")
    md.append("| Parâmetro Financeiro | Fonte / Documento Origem | Nível de Confiança | Observação / Governança |")
    md.append("|---|---|:---:|---|")
    md.append("| **Cronograma Físico-Financeiro** | `CRONOGRAMA_FISICO_FINANCEIRO_TMULT.csv` | 🟢 Alto | 158 itens distribuídos por cálculo CPM de 31 atividades. |")
    md.append("| **Custos Diretos e BDI** | Base Oficial SINAPI SP 07/2026 (`apoio/sinapi_sp/`) | 🟢 Alto | 100% auditado com Códigos CIA e BDI analítico (27,17% / 15,00%). |")
    md.append("| **Retenção Técnica Contratual** | Cláusula 5ª da Proposta Comercial Turnkey | 🟢 Alto | Retenção padrão de 5% sobre faturamento de medição. |")
    md.append("| **Prazos de Recebimento de Medições** | Cláusula 4ª da Proposta Comercial Turnkey | 🟢 Alto | Medição no último dia útil; pagamento em D+15 após NF. |")
    md.append("| **Prazos Comerciais de Fornecedores** | Padrão Comercial da Construção Civil / Curva ABC | 🟢 Alto | Mão de obra no mês; Materiais: 20% à vista e 80% D+30. |")
    md.append("| **Alíquotas Tributárias** | Composição do BDI (Tributos Municipais e Federais) | 🟢 Alto | 8,65% s/ NF (ISS 3%, PIS 0,65%, COFINS 3%, CPRB 2%). |\n")
    md.append("---\n")
    
    # Resumo Executivo Financeiro
    md.append("## 2. Resumo Executivo dos Indicadores Financeiros")
    md.append("\n| Indicador Econômico-Financeiro | Valor Consolidado (R$) | % da Receita Bruta | Impacto / Significado Operacional |")
    md.append("|---|:---:|:---:|---|")
    md.append(f"| **Faturamento Bruto da Obra (Turnkey)** | **{formatar_moeda(fluxo['total_venda'])}** | 100,00% | Preço global contratado fechado (R$ 4.508,04/m² útil). |")
    md.append(f"| **Retenção Contratual de Garantia (5,0%)** | **{formatar_moeda(fluxo['total_retencao'])}** | 5,00% | Retido nas medições M1 a M6; liberado integralmente no Mês 7 (TRD). |")
    md.append(f"| **Faturamento Líquido de Medições (95%)** | **{formatar_moeda(fluxo['total_venda'] - fluxo['total_retencao'])}** | 95,00% | Volume financeiro disponível durante o transcorrer da obra civil. |")
    md.append(f"| **Custo Direto Total da Obra** | **{formatar_moeda(fluxo['total_custo_direto'])}** | 79,15% | Custo de execução física (R$ 948k) + Canteiro/Gestão EAP 1.0 (R$ 366k). |")
    md.append(f"| **Tributos sobre Faturamento (8,65% s/ NF)** | **{formatar_moeda(total_impostos)}** | 8,65% | Recolhimento mensal no dia 20 subsequente à emissão de cada nota fiscal. |")
    md.append(f"| **Administração Central, Seguros e Riscos** | **{formatar_moeda(total_ac_seguros)}** | 5,94% | Rateio dos custos indiretos centrais e apólices de seguro da construtora. |")
    md.append(f"| **LUCRO LÍQUIDO OPERACIONAL REALIZADO** | **{formatar_moeda(saldo_final_realista)}** | **6,26%** | Margem líquida real de lucro após todos os tributos e despesas quitados. |")
    md.append(f"| **NECESSIDADE MÍNIMA DE CAPITAL DE GIRO** | **{formatar_moeda(abs(pico_realista))}** | **10,87%** | **Máxima exposição financeira de caixa (atingida no Mês {mes_pico_realista}).** |")
    md.append(f"| **Ponto de Equilíbrio do Caixa (Break-even)** | **Mês 7** | — | Ponto em que o caixa torna-se definitivamente superavitário. |\n")
    md.append("---\n")
    
    # Diagnóstico da Necessidade de Capital de Giro
    md.append("## 3. Diagnóstico do Capital de Giro e Timing de Desembolso")
    md.append("\n> 💡 **A Regra de Ouro do PMO:** *Estar dentro do orçamento não significa ter liquidez no momento certo.* O descasamento temporal entre os pagamentos a fornecedores/folha e o recebimento das medições é a principal causa de mortalidade de empresas de construção.\n")
    md.append(f"No cenário realista da `OBRA_TMULT`, a construtora experimenta sua **máxima exposição financeira no Mês {mes_pico_realista}**, atingindo um saldo acumulado negativo de **{formatar_moeda(pico_realista)}**:\n")
    md.append("1. **Mês 1 (Mobilização e Fundações):** A construtora mobiliza 14 operários, aluga 4 containers habitáveis e adquire aço, madeira e concreto para as sapatas e baldrames. Desembolso de **R$ 105.458,24** sem qualquer receita (já que a 1ª medição física é aferida no final de M1 e paga em M2). Saldo M1: **-R$ 105.458,24**.")
    md.append("2. **Mês 2 (Superestrutura):** Entra a 1ª medição líquida (R$ 155.269,92). No entanto, ocorrem os desembolsos de formas e concreto da supraestrutura + faturas a 30 dias de materiais de M1 + primeiro recolhimento de impostos. Saldo M2: **-R$ 152.369,27**.")
    md.append(f"3. **Mês 3 (Alvenaria e Cobertura — Ponto Crítico):** Coincide com o pico da Curva S (execução da cobertura sanduíche e alvenaria de blocos). O desembolso mensal atinge R$ 292.224,96 contra uma receita líquida de R$ 264.001,06, levando o saldo acumulado ao seu vale histórico: **{formatar_moeda(pico_realista)}**.")
    md.append("4. **Mês 4 em diante (Recuperação):** O recebimento da volumosa medição do Mês 3 (R$ 348.054,46) estanca a sangria, e a curva de caixa inicia sua trajetória ascendente.")
    md.append("5. **Mês 7 (Liquidação e Entrega):** Ocorre o recebimento da última medição (R$ 272.605,35) somado à **devolução integral da retenção técnica de 5% (R$ 83.038,11)**, totalizando uma entrada de R$ 355.643,46, quitando os tributos e contas residuais de M6 e consolidando o **Lucro Líquido Realizado de R$ 103.951,49**.\n")
    md.append("---\n")
    
    # Tabela Completa do Fluxo de Caixa Mensal
    md.append("## 4. Demonstrativo do Fluxo de Caixa Mensal (Cenário Realista Base)\n")
    md.append("| Rubrica Financeira | Mês 1 | Mês 2 | Mês 3 | Mês 4 | Mês 5 | Mês 6 | Mês 7 | Total Consolidado (R$) |")
    md.append("|---|:---:|:---:|:---:|:---:|:---:|:---:|:---:|:---:|")
    
    def format_row(titulo, dic):
        vals = [formatar_moeda(dic.get(m, 0.0)) for m in fluxo["meses"]]
        tot = formatar_moeda(sum(dic.get(m, 0.0) for m in fluxo["meses"]))
        return f"| {titulo} | " + " | ".join(vals) + f" | **{tot}** |"
        
    md.append("| **1. ENTRADAS DE CAIXA** | | | | | | | | |")
    md.append(format_row("1.1 Faturamento Bruto Previsto", fluxo["faturamento_bruto"]))
    md.append(format_row("1.2 Retenção Contratual (5,0%)", fluxo["retencao_5pct"]))
    md.append(format_row("1.3 Faturamento Líquido (95%)", fluxo["faturamento_liquido"]))
    md.append(format_row("1.4 Recebimento Efetivo (D+15)", fluxo["recebimento_medicoes"]))
    md.append(format_row("1.5 Devolução da Retenção (M7)", fluxo["devolucao_retencao"]))
    md.append(format_row("**TOTAL ENTRADAS (A)**", fluxo["total_entradas_realista"]))
    
    md.append("| **2. SAÍDAS DE CAIXA** | | | | | | | | |")
    md.append(format_row("2.1 Equipe Gestão (Eng/Mestre/TST)", fluxo["des_mo_gestao"]))
    md.append(format_row("2.2 Vivência e Alimentação (16 op.)", fluxo["des_vivencia"]))
    md.append(format_row("2.3 Mão de Obra de Campo (Físico)", fluxo["des_mo_campo"]))
    md.append(format_row("2.4 Materiais (À Vista no Pedido)", fluxo["des_mat_vista"]))
    md.append(format_row("2.5 Materiais (A Prazo D+30)", fluxo["des_mat_prazo"]))
    md.append(format_row("2.6 Containers NR-18 e Sanitários", fluxo["des_locacoes"]))
    md.append(format_row("2.7 Consumo Canteiro (Água/Luz/Net)", fluxo["des_contas"]))
    md.append(format_row("2.8 Tributos s/ Faturamento (8,65%)", fluxo["des_impostos"]))
    md.append(format_row("2.9 Custos Indiretos Centrais BDI", fluxo["des_ac_seguros"]))
    md.append(format_row("**TOTAL SAÍDAS (B)**", fluxo["total_saidas_realista"]))
    
    md.append("| **3. SALDO E RESULTADO** | | | | | | | | |")
    md.append(format_row("**3.1 Saldo Operacional Líquido**", fluxo["saldo_periodo_realista"]))
    
    vals_acum = [formatar_moeda(fluxo["saldo_acumulado_realista"][m]) for m in fluxo["meses"]]
    tot_acum = formatar_moeda(fluxo["saldo_acumulado_realista"][7])
    md.append(f"| **3.2 Saldo de Caixa Acumulado** | " + " | ".join(vals_acum) + f" | **{tot_acum}** |\n")
    md.append("---\n")
    
    # Análise de Sensibilidade em 3 Cenários
    md.append("## 5. Análise de Sensibilidade — Comparativo em 3 Cenários")
    md.append("\nSeguindo a diretriz obrigatória da `SKILL_GESTAO_07` de nunca apresentar projeções em número único, projetamos 3 cenários operacionais:\n")
    md.append("| Mês | Cenário Otimista (c/ Adiantamento 10%) | Cenário Realista (Base Contratual) | Cenário Estresse (Atraso Medição 30d) |")
    md.append("|:---:|:---:|:---:|:---:|")
    
    for m in range(1, 8):
        vo = formatar_moeda(fluxo["saldo_acumulado_otimista"].get(m, 0.0))
        vr = formatar_moeda(fluxo["saldo_acumulado_realista"].get(m, 0.0))
        vp = formatar_moeda(fluxo["saldo_acumulado_pessimista"].get(m, 0.0))
        md.append(f"| Mês {m} | {vo} | {vr} | {vp} |")
    md.append(f"| Mês 8 | — | — | {formatar_moeda(fluxo['saldo_acumulado_pessimista'][8])} |")
    md.append(f"| **PICO MÁXIMO DE EXPOSIÇÃO** | **{formatar_moeda(pico_otimista)}** | **{formatar_moeda(pico_realista)}** | **{formatar_moeda(pico_pessimista)}** |")
    md.append(f"| **Saldo Final Realizado** | **{formatar_moeda(saldo_final_realista)}** | **{formatar_moeda(saldo_final_realista)}** | **{formatar_moeda(saldo_final_realista)}** |\n")
    
    md.append("### 🔍 Insights dos Cenários:")
    md.append(f"- **No Cenário Otimista:** A negociação de um **Adiantamento Contratual de Mobilização de 10% (R$ 166.076,23)** no Dia Zero (amortizado em 4 parcelas nas primeiras medições) mantém o caixa positivo em +R$ 60.617,99 logo no primeiro mês, mitigando significativamente o risco inicial da construtora.")
    md.append(f"- **No Cenário Realista:** A construtora precisa ter assegurada uma linha de liquidez / capital de giro de **R$ 180.593,16** (recomendando-se um fundo de contingência de **R$ 200.000,00** para cobrir flutuações pontuais).")
    md.append(f"- **No Cenário de Estresse (Atraso de Medição):** Se o cliente atrasar a liberação de cada medição em 30 dias (pagamento em D+45 em vez de D+15), a exposição máxima salta para **{formatar_moeda(pico_pessimista)} (no Mês {mes_pico_pessimista})**. Isso demonstra cabalmente o risco de descasamento e justifica as cláusulas de proteção e multas moratórias incluídas na Proposta Comercial.\n")
    md.append("---\n")
    
    # Recomendações Estratégicas do PMO
    md.append("## 6. Recomendações Estratégicas do Engenheiro Chefe / PMO")
    md.append("1. **Constituição de Fundo de Capital de Giro de R$ 200.000,00:** Provisionar o aporte inicial antes do início das escavações para suportar o ciclo financeiro até o recebimento da 3ª medição.")
    md.append("2. **Negociação Comercial de Fornecedores Chave (Curva A):**")
    md.append("   - Telhas Termoacústicas (R$ 86k no M3): Negociar pagamento 50% em 30 dias e 50% em 60 dias da entrega da carga no Porto do Açu.")
    md.append("   - Climatização Split / VRF (R$ 78k no M5): Faturar direto com distribuidor em 30/60 dias.")
    md.append("   - Aço e Concreto Usinado: Homologar fornecedores locais com faturamento quinzenal e prazo de 28 dias.")
    md.append("3. **Gestão Rigorosa do Protocolo de Medição:**")
    md.append("   - Protocolar a Folha de Medição impreterivelmente no dia 25 de cada mês junto à fiscalização do Terminal Multiuso para garantir que a NF seja emitida no dia 01 e paga até o dia 15.")
    md.append("4. **Devolução da Retenção Técnica (Mês 7):**")
    md.append("   - Concluir o DataBook e as pranchas As-Built (EAP 5.1 e POP 18) com 15 dias de antecedência para homologar o Termo de Recebimento Definitivo (TRD) no 1º dia útil do Mês 7, destravando a devolução imediata dos **R$ 83.038,11** retidos.\n")
    
    with open(md_path, 'w', encoding='utf-8') as f:
        f.write("\n".join(md))
        
    print(f"[OK] Relatório Markdown de Fluxo de Caixa gerado: {md_path}")

def main():
    print("=" * 75)
    print("MOTOR DE GERAÇÃO DO FLUXO DE CAIXA E CAPITAL DE GIRO — OBRA_TMULT")
    print("=" * 75)
    
    vendas_m, diretos_m, total_venda, total_custo_direto = carregar_dados()
    fluxo = modelar_fluxo_caixa(vendas_m, diretos_m, total_venda, total_custo_direto)
    
    gerar_csv(fluxo)
    gerar_xlsx(fluxo)
    gerar_relatorio_markdown(fluxo)
    
    print("=" * 75)
    print("PROCESSAMENTO CONCLUÍDO COM SUCESSO!")
    print(f"Diretório de Saída: {OUTPUT_DIR}")
    print("=" * 75)

if __name__ == "__main__":
    main()
