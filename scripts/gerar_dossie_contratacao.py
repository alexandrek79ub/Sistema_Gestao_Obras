#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Motor Universal de Geração do Dossiê Executivo de Contratação (v2.0 Lean):
1. Histograma de Mão de Obra (Headcount & Horas-Homem) -> 06_SST_E_RH
2. Histograma de Equipamentos e Instalações de Canteiro -> 04_PRODUCAO_E_AVANCO
3. Curva ABC Dupla (Composições de Serviços e Famílias de Insumos) -> 02_ORCAMENTO_BASE_E_CONTRATOS
4. Dossiê Executivo Consolidado em Markdown -> 02_ORCAMENTO_BASE_E_CONTRATOS

Uso:
    python scripts/gerar_dossie_contratacao.py --obra OBRA_TMULT
    python scripts/gerar_dossie_contratacao.py --obra RESIDENCIAL_ALPHA
    python scripts/gerar_dossie_contratacao.py --dir projetos/OBRA_TMULT
"""

import os
import sys
import json
import argparse
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Adiciona raiz do projeto ao path para importação da infraestrutura compartilhada
ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

from scripts.common.obra_io import resolver_obra_dir, carregar_config_obra
from scripts.common.excel_theme import (
    NAVY, BLUE_DARK, GOLD_ACCENT, GRAY_LIGHT, GREEN_FILL as GREEN_LIGHT, BLUE_LIGHT,
    FONT_TITLE, FONT_HEADER, FONT_BOLD, FONT_REGULAR,
    THIN_BORDER, DOUBLE_BOTTOM_BORDER
)

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

NAVY_HEADER = NAVY
DOUBLE_BOTTOM = DOUBLE_BOTTOM_BORDER


def parse_args():
    parser = argparse.ArgumentParser(description="Motor Universal de Dossiê Executivo de Contratação (v2.0 Lean).")
    parser.add_argument("--obra", type=str, default="OBRA_TMULT", help="Nome da pasta da obra em projetos/")
    parser.add_argument("--dir", type=str, default=None, help="Caminho completo da pasta da obra")
    return parser.parse_args()


def parse_val(v):
    if isinstance(v, str):
        v = v.replace('R$', '').strip().replace('.', '').replace(',', '.')
    return float(v)


# ==============================================================================
# 1. HISTOGRAMA DE MÃO DE OBRA
# ==============================================================================
def gerar_histograma_mao_de_obra(dir_rh, dados_mo, sigla, titulo_obra, prazo_meses=6):
    print("-> Gerando Histograma de Mão de Obra...")
    cols = ["Grupo", "Função / Cargo", "Categoria", "Custo Base Ref (R$/mês)"] + [f"Mês {m}" for m in range(1, prazo_meses + 1)]

    dados_mo_ajustados = []
    for row in dados_mo:
        prefixo = row[:4]
        meses_vals = row[4:]
        if len(meses_vals) >= prazo_meses:
            ajustado = prefixo + meses_vals[:prazo_meses]
        else:
            ajustado = prefixo + meses_vals + [meses_vals[-1] if meses_vals else 0] * (prazo_meses - len(meses_vals))
        dados_mo_ajustados.append(ajustado)

    df_mo = pd.DataFrame(dados_mo_ajustados, columns=cols)
    totais_headcount = [int(df_mo[f"Mês {m}"].sum()) for m in range(1, prazo_meses + 1)]
    totais_hh = [hc * 220 for hc in totais_headcount]

    csv_path = os.path.join(dir_rh, f"HISTOGRAMA_MAO_DE_OBRA_{sigla}.csv")
    df_mo.to_csv(csv_path, sep=';', index=False, encoding='utf-8-sig')

    xlsx_path = os.path.join(dir_rh, f"HISTOGRAMA_MAO_DE_OBRA_{sigla}.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Histograma de Mão de Obra"
    ws.views.sheetView[0].showGridLines = True

    last_ltr = get_column_letter(len(cols))
    ws.merge_cells(f"A1:{last_ltr}2")
    ws["A1"] = f"{sigla} — HISTOGRAMA DE MÃO DE OBRA & HEADCOUNT MENSAL"
    ws["A1"].fill = NAVY_HEADER
    ws["A1"].font = FONT_TITLE
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(f"A3:{last_ltr}3")
    ws["A3"] = f"Planejamento Físico de Efetivo | {prazo_meses} Meses ({prazo_meses * 30} Dias) | Carga Horária Padrão: 220 Horas-Homem (HH) / mês"
    ws["A3"].fill = BLUE_DARK
    ws["A3"].font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

    row_hdr = 5
    for col_idx, h in enumerate(cols, start=1):
        cell = ws.cell(row=row_hdr, column=col_idx, value=h)
        cell.fill = NAVY_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    for r_idx, row_data in enumerate(dados_mo_ajustados, start=6):
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = FONT_REGULAR
            cell.border = THIN_BORDER
            if c_idx == 4:
                cell.number_format = '"R$ "#,##0.00'
                cell.alignment = Alignment(horizontal="right")
            elif c_idx >= 5:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="left" if c_idx == 2 else "center")
            if r_idx % 2 == 1:
                cell.fill = GRAY_LIGHT

    r_hc = len(dados_mo_ajustados) + 6
    ws.cell(row=r_hc, column=1, value="TOTAL").alignment = Alignment(horizontal="center")
    ws.cell(row=r_hc, column=2, value="HEADCOUNT TOTAL DE CAMPO (Operários + Gestão)").alignment = Alignment(horizontal="left")
    for m_idx, hc in enumerate(totais_headcount, start=5):
        ws.cell(row=r_hc, column=m_idx, value=hc).number_format = '#,##0'
        ws.cell(row=r_hc, column=m_idx).alignment = Alignment(horizontal="center")

    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=r_hc, column=c)
        cell.fill = GOLD_ACCENT
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.border = THIN_BORDER

    r_hh = r_hc + 1
    ws.cell(row=r_hh, column=1, value="TOTAL HH").alignment = Alignment(horizontal="center")
    ws.cell(row=r_hh, column=2, value="TOTAL DE HORAS-HOMEM PREVISTAS (220h / profissional)").alignment = Alignment(horizontal="left")
    for m_idx, hh in enumerate(totais_hh, start=5):
        ws.cell(row=r_hh, column=m_idx, value=hh).number_format = '#,##0'
        ws.cell(row=r_hh, column=m_idx).alignment = Alignment(horizontal="center")

    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=r_hh, column=c)
        cell.fill = GREEN_LIGHT
        cell.font = FONT_BOLD
        cell.border = DOUBLE_BOTTOM

    col_widths = {1: 14, 2: 44, 3: 18, 4: 22}
    for c_i in range(5, len(cols) + 1):
        col_widths[c_i] = 12
    for c_idx, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = w

    wb.save(xlsx_path)
    print(f"   -> Salvo: {csv_path}")
    print(f"   -> Salvo: {xlsx_path}")
    return df_mo, totais_headcount, totais_hh


# ==============================================================================
# 2. HISTOGRAMA DE EQUIPAMENTOS
# ==============================================================================
def gerar_histograma_equipamentos(dir_prod, dados_eq, sigla, titulo_obra, prazo_meses=6):
    print("-> Gerando Histograma de Equipamentos...")
    cols = ["Equipamento / Instalação Provisória", "Especificação / Modelo", "Unid"] + [f"Mês {m}" for m in range(1, prazo_meses + 1)] + ["Finalidade Operacional"]

    dados_eq_ajustados = []
    for row in dados_eq:
        prefixo = row[:3]
        meses_vals = row[3:-1]
        sufixo = [row[-1]]
        if len(meses_vals) >= prazo_meses:
            ajustado = prefixo + meses_vals[:prazo_meses] + sufixo
        else:
            ajustado = prefixo + meses_vals + [meses_vals[-1] if meses_vals else 0] * (prazo_meses - len(meses_vals)) + sufixo
        dados_eq_ajustados.append(ajustado)

    df_eq = pd.DataFrame(dados_eq_ajustados, columns=cols)

    csv_path = os.path.join(dir_prod, f"HISTOGRAMA_EQUIPAMENTOS_{sigla}.csv")
    df_eq.to_csv(csv_path, sep=';', index=False, encoding='utf-8-sig')

    xlsx_path = os.path.join(dir_prod, f"HISTOGRAMA_EQUIPAMENTOS_{sigla}.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Histograma de Equipamentos"
    ws.views.sheetView[0].showGridLines = True

    last_ltr = get_column_letter(len(cols))
    ws.merge_cells(f"A1:{last_ltr}2")
    ws["A1"] = f"{sigla} — HISTOGRAMA DE EQUIPAMENTOS & INFRAESTRUTURA DE CANTEIRO"
    ws["A1"].fill = NAVY_HEADER
    ws["A1"].font = FONT_TITLE
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(f"A3:{last_ltr}3")
    ws["A3"] = f"Dimensionamento de Maquinário, Andaimes e Módulos Habitáveis ao longo dos {prazo_meses} Meses de Execução"
    ws["A3"].fill = BLUE_DARK
    ws["A3"].font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

    row_hdr = 5
    for col_idx, h in enumerate(cols, start=1):
        cell = ws.cell(row=row_hdr, column=col_idx, value=h)
        cell.fill = NAVY_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    for r_idx, row_data in enumerate(dados_eq_ajustados, start=6):
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = FONT_REGULAR
            cell.border = THIN_BORDER
            if 4 <= c_idx <= 3 + prazo_meses:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="center")
            elif c_idx == 3:
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="left")
            if r_idx % 2 == 1:
                cell.fill = GRAY_LIGHT

    col_widths = {1: 40, 2: 38, 3: 8}
    for c_i in range(4, 4 + prazo_meses):
        col_widths[c_i] = 10
    col_widths[4 + prazo_meses] = 35
    for c_idx, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = w

    wb.save(xlsx_path)
    print(f"   -> Salvo: {csv_path}")
    print(f"   -> Salvo: {xlsx_path}")
    return df_eq


# ==============================================================================
# 3. CURVA ABC DE COMPOSIÇÕES E INSUMOS
# ==============================================================================
def gerar_curva_abc_completa(dir_orcamento, orc_path, dados_insumos, sigla, titulo_obra):
    print("-> Gerando Curva ABC de Serviços e Insumos...")

    df_orc = pd.read_csv(orc_path, sep=';', encoding='utf-8')
    eap_col = [c for c in df_orc.columns if 'digo EAP' in c or 'EAP' in c][0]
    desc_col = [c for c in df_orc.columns if 'Descricao' in c or 'Item' in c][0]

    df_orc['Preco_Total'] = df_orc['Custo Total (R$)'].apply(parse_val)
    df_orc['Preco_Unit'] = df_orc['Preço Unitário (R$)'].apply(parse_val)

    df_serv = df_orc.sort_values(by='Preco_Total', ascending=False).reset_index(drop=True)
    total_geral = df_serv['Preco_Total'].sum()

    df_serv['Pct_Individual'] = (df_serv['Preco_Total'] / total_geral) * 100
    df_serv['Pct_Acumulado'] = df_serv['Pct_Individual'].cumsum()

    def classificar_abc(acum):
        if acum <= 80.0001:
            return 'A'
        elif acum <= 95.0001:
            return 'B'
        else:
            return 'C'

    df_serv['Classe_ABC'] = df_serv['Pct_Acumulado'].apply(classificar_abc)

    cols_ins = ["Família de Insumo / Recurso Chave", "Natureza de Custo", "Custo Direto Total (R$)", "Qtd Macro", "Unid", "Referência / Aplicação na Obra"]
    df_ins = pd.DataFrame(dados_insumos, columns=cols_ins)
    tot_ins = df_ins["Custo Direto Total (R$)"].sum()
    df_ins["Pct_Individual"] = (df_ins["Custo Direto Total (R$)"] / tot_ins) * 100
    df_ins["Pct_Acumulado"] = df_ins["Pct_Individual"].cumsum()
    df_ins["Classe_ABC"] = df_ins["Pct_Acumulado"].apply(classificar_abc)

    csv_serv_path = os.path.join(dir_orcamento, f"CURVA_ABC_SERVICOS_{sigla}.csv")
    df_serv.to_csv(csv_serv_path, sep=';', index=False, encoding='utf-8-sig')

    csv_ins_path = os.path.join(dir_orcamento, f"CURVA_ABC_INSUMOS_{sigla}.csv")
    df_ins.to_csv(csv_ins_path, sep=';', index=False, encoding='utf-8-sig')

    xlsx_path = os.path.join(dir_orcamento, f"CURVA_ABC_SERVICOS_E_INSUMOS_{sigla}.xlsx")
    wb = openpyxl.Workbook()

    # Aba 1: Insumos
    ws1 = wb.active
    ws1.title = "Curva ABC Insumos"
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:I2")
    ws1["A1"] = f"CURVA ABC DE FAMÍLIAS DE INSUMOS & RECURSOS — {sigla}"
    ws1["A1"].fill = NAVY_HEADER
    ws1["A1"].font = FONT_TITLE
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws1.merge_cells("A3:I3")
    ws1["A3"] = f"Abertura Analítica do Custo Direto Total: R$ {tot_ins:,.2f}"
    ws1["A3"].fill = BLUE_DARK
    ws1["A3"].font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
    ws1["A3"].alignment = Alignment(horizontal="center", vertical="center")

    headers_ins = ["Item", "Família de Insumo / Recurso Chave", "Natureza", "Custo Direto (R$)", "Qtd", "Unid", "% Indiv.", "% Acum.", "Classe ABC"]
    row_hdr = 5
    for c_idx, h in enumerate(headers_ins, start=1):
        cell = ws1.cell(row=row_hdr, column=c_idx, value=h)
        cell.fill = NAVY_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for idx, r in df_ins.iterrows():
        r_num = idx + 6
        ws1.cell(row=r_num, column=1, value=idx+1).alignment = Alignment(horizontal="center")
        ws1.cell(row=r_num, column=2, value=r["Família de Insumo / Recurso Chave"]).alignment = Alignment(horizontal="left")
        ws1.cell(row=r_num, column=3, value=r["Natureza de Custo"]).alignment = Alignment(horizontal="center")
        ws1.cell(row=r_num, column=4, value=r["Custo Direto Total (R$)"]).number_format = '"R$ "#,##0.00'
        ws1.cell(row=r_num, column=5, value=r["Qtd Macro"]).number_format = '#,##0.00'
        ws1.cell(row=r_num, column=6, value=r["Unid"]).alignment = Alignment(horizontal="center")
        ws1.cell(row=r_num, column=7, value=r["Pct_Individual"] / 100.0).number_format = '0.00%'
        ws1.cell(row=r_num, column=8, value=r["Pct_Acumulado"] / 100.0).number_format = '0.00%'

        cell_abc = ws1.cell(row=r_num, column=9, value=r["Classe_ABC"])
        cell_abc.alignment = Alignment(horizontal="center")
        cell_abc.font = Font(name="Calibri", size=11, bold=True)
        if r["Classe_ABC"] == 'A':
            cell_abc.fill = PatternFill(start_color="FFD1D1", end_color="FFD1D1", fill_type="solid")
        elif r["Classe_ABC"] == 'B':
            cell_abc.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        else:
            cell_abc.fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")

        for c in range(1, 9):
            ws1.cell(row=r_num, column=c).font = FONT_REGULAR
            ws1.cell(row=r_num, column=c).border = THIN_BORDER
            if idx % 2 == 1:
                ws1.cell(row=r_num, column=c).fill = GRAY_LIGHT

    r_tot_ins = len(df_ins) + 6
    ws1.cell(row=r_tot_ins, column=1, value="TOTAL").alignment = Alignment(horizontal="center")
    ws1.cell(row=r_tot_ins, column=2, value="TOTAL CUSTO DIRETO DA OBRA").alignment = Alignment(horizontal="left")
    ws1.cell(row=r_tot_ins, column=4, value=tot_ins).number_format = '"R$ "#,##0.00'
    ws1.cell(row=r_tot_ins, column=7, value=1.0).number_format = '0.00%'
    ws1.cell(row=r_tot_ins, column=8, value=1.0).number_format = '0.00%'
    for c in range(1, 10):
        cell = ws1.cell(row=r_tot_ins, column=c)
        cell.fill = GREEN_LIGHT
        cell.font = FONT_BOLD
        cell.border = DOUBLE_BOTTOM

    col_w_ins = {1: 8, 2: 45, 3: 16, 4: 22, 5: 12, 6: 8, 7: 12, 8: 12, 9: 14}
    for c_idx, w in col_w_ins.items():
        ws1.column_dimensions[get_column_letter(c_idx)].width = w

    # Aba 2: Serviços
    ws2 = wb.create_sheet(title=f"Curva ABC Serviços ({len(df_serv)}I)")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:J2")
    ws2["A1"] = f"CURVA ABC DE COMPOSIÇÕES E PACOTES DE SERVIÇOS ({len(df_serv)} ITENS DA EAP)"
    ws2["A1"].fill = NAVY_HEADER
    ws2["A1"].font = FONT_TITLE
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers_serv = ["Rank", "Código EAP", "Descrição do Pacote / Serviço", "Disciplina", "Qtd Proj", "Unid", "Preço Total Turnkey (R$)", "% Indiv.", "% Acum.", "Classe ABC"]
    for c_idx, h in enumerate(headers_serv, start=1):
        cell = ws2.cell(row=4, column=c_idx, value=h)
        cell.fill = NAVY_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for idx, r in df_serv.iterrows():
        r_num = idx + 5
        ws2.cell(row=r_num, column=1, value=idx+1).alignment = Alignment(horizontal="center")
        ws2.cell(row=r_num, column=2, value=str(r[eap_col]).strip()).alignment = Alignment(horizontal="center")
        ws2.cell(row=r_num, column=3, value=str(r[desc_col]).strip()).alignment = Alignment(horizontal="left")
        ws2.cell(row=r_num, column=4, value=str(r.get("Disciplina", "")).strip()).alignment = Alignment(horizontal="center")
        ws2.cell(row=r_num, column=5, value=r["Qtd Projeto"]).number_format = '#,##0.00'
        ws2.cell(row=r_num, column=6, value=str(r.get("Unidade Proj", "")).strip()).alignment = Alignment(horizontal="center")
        ws2.cell(row=r_num, column=7, value=r["Preco_Total"]).number_format = '"R$ "#,##0.00'
        ws2.cell(row=r_num, column=8, value=r["Pct_Individual"] / 100.0).number_format = '0.00%'
        ws2.cell(row=r_num, column=9, value=r["Pct_Acumulado"] / 100.0).number_format = '0.00%'

        cell_abc = ws2.cell(row=r_num, column=10, value=r["Classe_ABC"])
        cell_abc.alignment = Alignment(horizontal="center")
        cell_abc.font = Font(name="Calibri", size=11, bold=True)
        if r["Classe_ABC"] == 'A':
            cell_abc.fill = PatternFill(start_color="FFD1D1", end_color="FFD1D1", fill_type="solid")
        elif r["Classe_ABC"] == 'B':
            cell_abc.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        else:
            cell_abc.fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")

        for c in range(1, 10):
            ws2.cell(row=r_num, column=c).font = FONT_REGULAR
            ws2.cell(row=r_num, column=c).border = THIN_BORDER
            if idx % 2 == 1:
                ws2.cell(row=r_num, column=c).fill = GRAY_LIGHT

    col_w_serv = {1: 8, 2: 14, 3: 45, 4: 20, 5: 12, 6: 8, 7: 22, 8: 12, 9: 12, 10: 14}
    for c_idx, w in col_w_serv.items():
        ws2.column_dimensions[get_column_letter(c_idx)].width = w

    wb.save(xlsx_path)
    print(f"   -> Salvo: {csv_serv_path}")
    print(f"   -> Salvo: {csv_ins_path}")
    print(f"   -> Salvo: {xlsx_path}")
    return df_serv, df_ins


# ==============================================================================
# 4. GERAÇÃO DO RELATÓRIO DO DOSSIÊ EXECUTIVO EM MARKDOWN
# ==============================================================================
def gerar_relatorio_dossie_markdown(dir_orcamento, sigla, titulo, prazo_meses, totais_hc, dados_eq, tot_ins):
    template_path = os.path.join(ROOT_DIR, "scripts", "templates", "contratacao", "template_dossie_executivo.md")
    if not os.path.exists(template_path):
        return

    with open(template_path, "r", encoding="utf-8") as f:
        tpl = f.read()

    hc_medio = sum(totais_hc) / len(totais_hc) if totais_hc else 0
    hc_pico = max(totais_hc) if totais_hc else 0
    mes_pico_idx = (totais_hc.index(hc_pico) + 1) if totais_hc else 1
    mes_pico = f"Mês {mes_pico_idx}"

    # Tabela síntese MO
    tab_mo_linhas = [
        "| Mês | " + " | ".join([f"Mês {m}" for m in range(1, prazo_meses + 1)]) + " | **Total HH** |",
        "| :---: | " + " | ".join([":---:" for _ in range(prazo_meses)]) + " | :---: |",
        "| **Headcount (Op)** | " + " | ".join([str(hc) for hc in totais_hc]) + f" | **{sum(totais_hc)*220:,.0f} HH** |"
    ]
    tab_mo_str = "\n".join(tab_mo_linhas)

    # Tabela síntese Equipamentos
    tab_eq_linhas = [
        "| Equipamento / Instalação | Qtd | Un | Finalidade Operacional |",
        "| :--- | :---: | :---: | :--- |"
    ]
    for row in dados_eq[:6]:
        tab_eq_linhas.append(f"| **{row[0]}** | {row[3]} | {row[2]} | {row[-1]} |")
    if len(dados_eq) > 6:
        tab_eq_linhas.append(f"| *... e mais {len(dados_eq)-6} itens menores/módulos NR-18* | — | — | *Ver HISTOGRAMA_EQUIPAMENTOS.xlsx* |")
    tab_eq_str = "\n".join(tab_eq_linhas)

    conteudo = (
        tpl
        .replace("{{NOME_OBRA_UPPER}}", titulo.upper())
        .replace("{{NOME_OBRA}}", titulo)
        .replace("{{SIGLA_OBRA}}", sigla)
        .replace("{{PRAZO_MESES}}", str(prazo_meses))
        .replace("{{DIAS_CORRIDOS}}", str(prazo_meses * 30))
        .replace("{{DATA_ATUAL}}", datetime.now().strftime("%d/%m/%Y"))
        .replace("{{HEADCOUNT_MEDIO}}", f"{hc_medio:.1f}".replace(".", ","))
        .replace("{{HEADCOUNT_PICO}}", str(hc_pico))
        .replace("{{MES_PICO}}", mes_pico)
        .replace("{{TABELA_RESUMO_MO}}", tab_mo_str)
        .replace("{{TABELA_RESUMO_EQ}}", tab_eq_str)
        .replace("{{VALOR_TOTAL_CD}}", f"R$ {tot_ins:,.2f}".replace(",", "v").replace(".", ",").replace("v", "."))
    )

    out_md = os.path.join(dir_orcamento, f"DOSSIE_EXECUTIVO_CONTRATACAO_{sigla}.md")
    with open(out_md, "w", encoding="utf-8") as f:
        f.write(conteudo)
    print(f"   -> Salvo: {out_md}")


def main():
    args = parse_args()
    project_dir = resolver_obra_dir(args)

    if not os.path.exists(project_dir):
        print(f"[ERRO] Diretório não encontrado: {project_dir}")
        sys.exit(1)

    config = carregar_config_obra(project_dir)
    dados_obra = config.get("dados_obra", {})
    sigla = dados_obra.get("sigla", config.get("sigla_obra", os.path.basename(project_dir).replace("OBRA_", "")))
    titulo = dados_obra.get("nome_obra", config.get("nome_obra", f"Obra {sigla}"))
    prazo_meses = int(dados_obra.get("prazo_meses", config.get("prazo_meses", 6)))

    dir_orcamento = os.path.join(project_dir, "02_ORCAMENTO_BASE_E_CONTRATOS")
    dir_producao = os.path.join(project_dir, "04_PRODUCAO_E_AVANCO")
    dir_rh = os.path.join(project_dir, "06_SST_E_RH")

    os.makedirs(dir_orcamento, exist_ok=True)
    os.makedirs(dir_producao, exist_ok=True)
    os.makedirs(dir_rh, exist_ok=True)

    mo_json = os.path.join(dir_rh, "dados_histograma_mo.json")
    eq_json = os.path.join(dir_producao, "dados_histograma_equipamentos.json")
    ins_json = os.path.join(dir_orcamento, "dados_familias_insumos.json")
    orc_path = os.path.join(dir_orcamento, "ORCAMENTO_BASE_CONSOLIDADO.csv")

    if not os.path.exists(orc_path):
        print(f"[ERRO] Orçamento base não encontrado: {orc_path}")
        sys.exit(1)

    with open(mo_json, "r", encoding="utf-8") as f:
        dados_mo = json.load(f)
    with open(eq_json, "r", encoding="utf-8") as f:
        dados_eq = json.load(f)
    with open(ins_json, "r", encoding="utf-8") as f:
        dados_insumos = json.load(f)

    print(f"=== MOTOR UNIVERSAL DE DOSSIÊ DE CONTRATAÇÃO: {titulo} ({sigla}) ===")
    _, totais_hc, _ = gerar_histograma_mao_de_obra(dir_rh, dados_mo, sigla, titulo, prazo_meses=prazo_meses)
    gerar_histograma_equipamentos(dir_producao, dados_eq, sigla, titulo, prazo_meses=prazo_meses)
    _, df_ins = gerar_curva_abc_completa(dir_orcamento, orc_path, dados_insumos, sigla, titulo)

    tot_ins = df_ins["Custo Direto Total (R$)"].sum()
    gerar_relatorio_dossie_markdown(dir_orcamento, sigla, titulo, prazo_meses, totais_hc, dados_eq, tot_ins)
    print("=== DOSSIÊ DE CONTRATAÇÃO GERADO COM SUCESSO! ===")


if __name__ == "__main__":
    main()
