#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Motor Universal de Geração de Contratos de Locação de Equipamentos e Planilha de Gestão.

Uso:
    python scripts/gerar_contratos_locacao.py --obra OBRA_TMULT
    python scripts/gerar_contratos_locacao.py --obra RESIDENCIAL_ALPHA
    python scripts/gerar_contratos_locacao.py --dir projetos/OBRA_TMULT
"""

import os
import sys
import json
import argparse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

def parse_args():
    parser = argparse.ArgumentParser(description="Motor Universal de Locação de Equipamentos.")
    parser.add_argument("--obra", type=str, default="OBRA_TMULT", help="Nome da pasta da obra em projetos/")
    parser.add_argument("--dir", type=str, default=None, help="Caminho completo da pasta da obra")
    return parser.parse_args()

def gerar_contratos_markdown(pacotes, output_dir, titulo_obra, sigla, area_m2=368.4):
    tot = sum(p.get('valor_total', 0.0) for p in pacotes)
    
    # 1. ÍNDICE MESTRE
    indice_md = os.path.join(output_dir, "INDICE_MESTRE_CONTRATOS_LOCACAO_EQUIPAMENTOS.md")
    with open(indice_md, "w", encoding="utf-8") as f:
        f.write(f"""# 🚜 ÍNDICE MESTRE: CONTRATOS DE LOCAÇÃO DE EQUIPAMENTOS & MÁQUINAS

> **Empreendimento:** {titulo_obra} ({area_m2:.2f} m²) — `{sigla}`  
> **Volume de Locação:** {len(pacotes)} Macro-Pacotes Contratuais  
> **Valor Total Consolidado das Locações:** **R$ {tot:,.2f}**  
> **Governança:** POP 04 (Equipamentos), POP 06 (Recebimento), NR-12, NR-18 e NR-35  
> **Planilha de Gestão e Controle:** [`PLANILHA_GESTAO_LOCACAO_EQUIPAMENTOS.xlsx`](file:///{output_dir.replace('\\', '/')}/PLANILHA_GESTAO_LOCACAO_EQUIPAMENTOS.xlsx)

---

## 1. Quadro Geral de Contratos de Locação Pré-Criados

| Contrato | Objeto da Locação de Equipamentos | Famílias Atendidas | Fornecedor Homologado | Prazo / Vigência | Centro Custo | Valor Total (R$) |
| :---: | :--- | :---: | :--- | :---: | :---: | :---: |
""")
        for p in pacotes:
            f.write(f"| **[{p['cod']}]({p['nome_arquivo']})** | {p['titulo'][:55]}... | `{p.get('itens_equip', '')[:30]}...` | {p.get('locador', '').split('/')[0].strip()} | {p.get('periodo', '')[:22]} | `{p.get('cc', '').split('(')[0].strip()}` | **R$ {p.get('valor_total', 0.0):,.2f}** |\n")
            
        f.write(f"\n**VALOR TOTAL CONSOLIDADO DE LOCAÇÃO DE MÁQUINAS E INSTALAÇÕES:** **R$ {tot:,.2f}**\n\n---\n\n")
        f.write(f"""## 2. Protocolo de Entrada e Segurança de Equipamentos no Canteiro (POP 04)

Antes de qualquer equipamento iniciar operação no canteiro da {sigla}, o Engenheiro Residente e o Técnico de Segurança do Trabalho (TST) devem exigir:

1. **Checklist de Conformidade NR-12 / NR-18:**
   - Botões de parada de emergência funcionais e acessíveis;
   - Proteção física rígida de correias, polias e cremalheiras;
   - Aterramento elétrico temporário do chassi ligado à malha de terra da obra;
   - Alarme sonoro e giroflex acionados automaticamente na ré para máquinas móveis.
2. **Documentação de Operadores Terceirizados (POP 17):**
   - Atestado de Saúde Ocupacional (ASO) apto para a função;
   - Certificado de treinamento de operador de máquinas pesadas / plataformas (NR-11, NR-12, NR-18 e NR-35);
   - Ficha de entrega de EPIs com CAs vigentes.
3. **ART de Responsabilidade Técnica:**
   - Obrigatória para o Cimbramento Metálico, Andaimes Fachadeiros, Instalações de Containers e Grupo Gerador.

---

## 3. Matriz de Manutenção e Substituição Preventiva (SLA)

| Pacote | Manutenção Programada | Tempo Máximo de Atendimento (SLA) | Penalidade por Indisponibilidade |
| :--- | :--- | :---: | :--- |
| **Módulos / Canteiro** | Sucção e limpeza sanitária 2x por semana | Até 24 horas úteis | Glosa diária de 5% sobre a locação |
| **Terraplenagem** | Revisão diária de graxa e nível de óleo pelo operador | Até 12 horas úteis | Desconto das horas paralisadas + reposição de maquinário |
| **Grupo Gerador** | Troca de filtros e óleo a cada 250 horas | Até 4 horas (Plantão 24/7) | Multa por risco de paralisação de concretagem |
| **Cimbramento** | Vistoria técnica antes da concretagem | Até 24 horas antes do lançamento | Bloqueio do Portão de Qualidade da Laje |
| **Andaimes / Plataformas** | Carga de baterias e verificação mensal | Até 24 horas úteis | Substituição imediata por plataforma reserva |
| **Equipamentos Produção** | Limpeza diária e aferição de manômetros | Até 12 horas úteis | Troca imediata da bomba de argamassa |
""")
    print(f"-> Índice Mestre de Locações gerado: {indice_md}")

    # 2. CONTRATOS INDIVIDUAIS
    for p in pacotes:
        caminho_c = os.path.join(output_dir, p["nome_arquivo"])
        with open(caminho_c, "w", encoding="utf-8") as f:
            f.write(f"""# 🚜 INSTRUMENTO PARTICULAR DE CONTRATO DE LOCAÇÃO DE EQUIPAMENTOS: {p['cod']}
### {p['titulo']}

> **Empreendimento:** {titulo_obra} ({area_m2:.2f} m²) — `{sigla}`  
> **Locatária:** CONSTRUTORA EXECUTIVA DO PORTO LTDA  
> **Locadora:** {p.get('locador', '')} | **CNPJ:** {p.get('cnpj', '')}  
> **Valor Total do Contrato:** R$ {p.get('valor_total', 0.0):,.2f} | **Centro de Custo:** {p.get('cc', '')}  
> **Prazo e Vigência:** {p.get('periodo', '')} | **EAP:** {p.get('eap', '')}  
> **Governança:** POP 04 (Equipamentos) / NR-12 / NR-18 / NR-35

---

## CLÁUSULA PRIMEIRA — DO OBJETO E EQUIPAMENTOS LOCADOS
1.1. O presente instrumento tem por objeto a locação das máquinas, equipamentos e instalações provisórias descritos abaixo, correspondentes ao pacote **{p['cod']}** da Linha de Base 01 da {sigla}:
""")
            for esc in p.get("escopo", []):
                f.write(f"- {esc};\n")
                
            f.write(f"""
1.2. Os equipamentos destinam-se exclusivamente ao atendimento das obras de construção civil do empreendimento {titulo_obra}, sendo vedada a sublocação ou desvio de finalidade sem anuência prévia da LOCATÁRIA.

## CLÁUSULA SEGUNDA — DOS VALORES E FORMA DE PAGAMENTO
2.1. Pela locação dos bens, a LOCATÁRIA pagará à LOCADORA o valor global de **R$ {p.get('valor_total', 0.0):,.2f}**, em parcelas mensais vinculadas ao relatório de medição e efetiva disponibilidade operacional no canteiro.
2.2. **Condição de Pagamento:** Faturamento em D+30 dias corridos após a medição quinzenal aprovada pelo Engenheiro Residente e apresentação da Nota Fiscal com o Centro de Custo `{p.get('cc', '').split('(')[0].strip()}`.
2.3. Não incidirá pagamento sobre dias ou horas em que o equipamento permanecer paralisado por quebra mecânica, falta de operador da LOCADORA ou não conformidade técnica com as normas de segurança (NR-12 / NR-18).

## CLÁUSULA TERCEIRA — DOS PRAZOS DE MOBILIZAÇÃO E DESMOBILIZAÇÃO
3.1. O prazo de entrega e início da locação no canteiro é impreterivelmente **{p.get('periodo', '').split('a')[0].strip()}**, sob pena de multa diária de 1% sobre o valor da locação.
3.2. A desmobilização ocorrerá em **{p.get('periodo', '').split('a')[1].strip() if 'a' in p.get('periodo', '') else p.get('periodo', '')}**, devendo a LOCATÁRIA comunicar com 5 dias úteis de antecedência para vistoria conjunta de encerramento.

## CLÁUSULA QUARTA — DAS OBRIGAÇÕES DA LOCADORA
""")
            for ob_l in p.get("obrigacoes_locador", []):
                f.write(f"- {ob_l};\n")
                
            f.write(f"""
## CLÁUSULA QUINTA — DAS OBRIGAÇÕES DA LOCATÁRIA
""")
            for ob_t in p.get("obrigacoes_locatario", []):
                f.write(f"- {ob_t};\n")
                
            f.write(f"""
## CLÁUSULA SEXTA — DA SEGURANÇA DO TRABALHO E PORTÃO DE ENTRADA (POP 04 / NR-18)
6.1. Todos os equipamentos deverão ser submetidos à vistoria de entrada no canteiro, com preenchimento da **Ficha de Inspeção de Equipamento (FIE)** pelo TST da obra.
6.2. Nenhum operador da LOCADORA poderá adentrar o canteiro sem portar o crachá de identificação, ASO com aptidão específica, certificados NR-11/12/18 e EPIs obrigatórios com CA válido.

---

## ANEXO I: CHECKLIST DE RECEBIMENTO TÉCNICO E CONFORMIDADE (POP 04)

| Item Verificado | Critério Normativo | Status Admissão | Responsável Vistoria |
| :--- | :--- | :---: | :---: |
| **Pintura e Estrutura** | Ausência de corrosão acentuada, trincas ou deformações nas chapas | [ ] Aprovado | Almoxarife / Eng. Residente |
| **Proteções Móveis (NR-12)**| Carenagens fixas com parafusos e travas de segurança operacionais | [ ] Aprovado | TST da Obra |
| **Sistema Elétrico (NR-10)**| Cabos íntegros sem emendas, tomadas industriais steck e aterramento | [ ] Aprovado | Eletricista de Manutenção |
| **Dispositivos de Emergência**| Botão tipo cogumelo com retenção e rearme manual testado in-loco | [ ] Aprovado | TST da Obra |
| **ART / Laudo do Fabricante**| ART recolhida por Responsável Técnico habilitado e registrada no CREA | [ ] Aprovado | Engenheiro Fiscal |

---

*Contrato pré-criado e auditado conforme a Linha de Base 01 da {sigla}.*
""")
        print(f"-> Contrato de locação gerado: {caminho_c}")

def gerar_planilha_gestao_locacoes(pacotes, output_dir, sigla):
    caminho = os.path.join(output_dir, "PLANILHA_GESTAO_LOCACAO_EQUIPAMENTOS.xlsx")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    ws = wb.create_sheet(title="Controle de Locações")
    ws.views.sheetView[0].showGridLines = True
    
    NAVY_HEADER = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    BLUE_HEADER = PatternFill(start_color="2B5B84", end_color="2B5B84", fill_type="solid")
    GOLD_HEADER = PatternFill(start_color="D99B26", end_color="D99B26", fill_type="solid")
    GRAY_LIGHT = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
    BLUE_LIGHT = PatternFill(start_color="E8EEF5", end_color="E8EEF5", fill_type="solid")
    
    FONT_TITLE = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    FONT_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    FONT_BOLD = Font(name="Calibri", size=10, bold=True, color="1B365D")
    FONT_REGULAR = Font(name="Calibri", size=10, color="333333")
    
    BORDER_THIN = Border(
        left=Side(style='thin', color='DDDDDD'), right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'), bottom=Side(style='thin', color='DDDDDD')
    )
    BORDER_TOTAL = Border(
        top=Side(style='thin', color='1B365D'), bottom=Side(style='double', color='1B365D'),
        left=Side(style='thin', color='DDDDDD'), right=Side(style='thin', color='DDDDDD')
    )

    tot = sum(p.get("valor_total", 0.0) for p in pacotes)
    ws.merge_cells("A1:K1")
    ws["A1"] = f"PAINEL DE GESTÃO DE CONTRATOS DE LOCAÇÃO DE EQUIPAMENTOS ({sigla} - R$ {tot:,.2f})"
    ws["A1"].font = FONT_TITLE
    ws["A1"].fill = NAVY_HEADER
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26
    
    headers = [
        "Contrato", "Objeto da Locação", "Equipamentos Atendidos", "Locador Homologado",
        "Centro Custo", "EAP", "Mobilização", "Desmobilização", "Prazo", "Valor Mensal (R$)", "Valor Total (R$)"
    ]
    ws.row_dimensions[3].height = 24
    for c_i, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=c_i, value=h)
        cell.font = FONT_HEADER
        cell.fill = BLUE_HEADER if c_i <= 8 else GOLD_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_THIN
        
    for idx, p in enumerate(pacotes, start=4):
        ws.row_dimensions[idx].height = 20
        ws.cell(row=idx, column=1, value=p["cod"]).alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=2, value=p.get("titulo", "").replace("CONTRATO DE LOCAÇÃO DE ", "")[:40]).alignment = Alignment(horizontal="left")
        ws.cell(row=idx, column=3, value=p.get("itens_equip", "")[:35]).alignment = Alignment(horizontal="left")
        ws.cell(row=idx, column=4, value=p.get("locador", "").split('/')[0].strip()).alignment = Alignment(horizontal="left")
        ws.cell(row=idx, column=5, value=p.get("cc", "").split('(')[0].strip()).alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=6, value=p.get("eap", "")).alignment = Alignment(horizontal="center")
        
        per = p.get("periodo", "")
        mob = per.split('a')[0].strip().split('(')[0].strip() if 'a' in per else per
        desmob = per.split('a')[1].strip().split('(')[0].strip() if 'a' in per else ""
        prz = per.split('(')[1].split('/')[0].strip() if '(' in per else per
        
        ws.cell(row=idx, column=7, value=mob).alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=8, value=desmob).alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=9, value=prz).alignment = Alignment(horizontal="center")
        
        c_vm = ws.cell(row=idx, column=10, value=p.get("valor_mensal", 0.0))
        c_vm.number_format = '"R$ "#,##0.00'
        c_vm.alignment = Alignment(horizontal="right")
        
        c_vt = ws.cell(row=idx, column=11, value=p.get("valor_total", 0.0))
        c_vt.number_format = '"R$ "#,##0.00'
        c_vt.font = FONT_BOLD
        c_vt.alignment = Alignment(horizontal="right")
        
        for c_i in range(1, 12):
            cell = ws.cell(row=idx, column=c_i)
            cell.border = BORDER_THIN
            if idx % 2 == 0:
                cell.fill = GRAY_LIGHT

    r_tot = 4 + len(pacotes)
    ws.row_dimensions[r_tot].height = 24
    ws.merge_cells(start_row=r_tot, start_column=1, end_row=r_tot, end_column=9)
    ws.cell(row=r_tot, column=1, value="TOTAL CONSOLIDADO DE LOCAÇÃO DE EQUIPAMENTOS:").font = FONT_BOLD
    ws.cell(row=r_tot, column=1).alignment = Alignment(horizontal="right", vertical="center")
    
    c_tot_vm = ws.cell(row=r_tot, column=10, value=f"=SUM(J4:J{r_tot-1})")
    c_tot_vm.number_format = '"R$ "#,##0.00'
    c_tot_vm.font = FONT_BOLD
    c_tot_vm.alignment = Alignment(horizontal="right")
    
    c_tot_vt = ws.cell(row=r_tot, column=11, value=f"=SUM(K4:K{r_tot-1})")
    c_tot_vt.number_format = '"R$ "#,##0.00'
    c_tot_vt.font = FONT_BOLD
    c_tot_vt.alignment = Alignment(horizontal="right")
    
    for c_i in range(1, 12):
        cell = ws.cell(row=r_tot, column=c_i)
        cell.fill = BLUE_LIGHT
        cell.border = BORDER_TOTAL

    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 14
    ws.column_dimensions["H"].width = 14
    ws.column_dimensions["I"].width = 12
    ws.column_dimensions["J"].width = 18
    ws.column_dimensions["K"].width = 18

    wb.save(caminho)
    print(f"-> Planilha de Locações gerada: {caminho}")

def main():
    args = parse_args()
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    
    if args.dir:
        project_dir = os.path.abspath(args.dir)
    else:
        project_dir = os.path.join(base_dir, "projetos", args.obra)
        
    if not os.path.exists(project_dir):
        print(f"[ERRO] Diretório não encontrado: {project_dir}")
        sys.exit(1)
        
    config_file = os.path.join(project_dir, "config_obra.json")
    config = {}
    if os.path.exists(config_file):
        with open(config_file, "r", encoding="utf-8") as f:
            config = json.load(f)
            
    dados_obra = config.get("dados_obra", {})
    sigla = dados_obra.get("sigla", config.get("sigla_obra", os.path.basename(project_dir).replace("OBRA_", "")))
    titulo = dados_obra.get("nome_obra", config.get("nome_obra", f"Obra {sigla}"))
    area_m2 = float(dados_obra.get("area_construida_m2", config.get("area_construida_m2", 368.4)))
    
    loc_json = os.path.join(project_dir, "02_ORCAMENTO_BASE_E_CONTRATOS", "contratos_locacao.json")
    if not os.path.exists(loc_json):
        print(f"[ERRO] Arquivo de locações não encontrado: {loc_json}")
        sys.exit(1)
        
    with open(loc_json, "r", encoding="utf-8") as f:
        pacotes = json.load(f)
        
    output_dir = os.path.join(project_dir, "02_ORCAMENTO_BASE_E_CONTRATOS", "CONTRATOS_EQUIPAMENTOS")
    os.makedirs(output_dir, exist_ok=True)
    
    print(f"=== MOTOR UNIVERSAL DE LOCAÇÃO DE EQUIPAMENTOS: {titulo} ({sigla}) ===")
    gerar_contratos_markdown(pacotes, output_dir, titulo, sigla, area_m2=area_m2)
    gerar_planilha_gestao_locacoes(pacotes, output_dir, sigla)
    print("=== CONTRATOS DE LOCAÇÃO GERADOS COM SUCESSO! ===")

if __name__ == "__main__":
    main()
