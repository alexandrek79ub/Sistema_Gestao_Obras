#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gerador Completo da Planilha de Medição Evolutiva de Empreiteiros (12 Quinzenas)
Atende rigorosamente ao formato especificado pelo usuário:
- Coluna 1: Item EAP
- Coluna 2: Descrição do Serviço
- Coluna 3: Unidade
- Coluna 4: Quantidade Contratada
- Coluna 5: Preço Unitário (R$)
- Coluna 6: Total Contrato (R$)
- Coluna 7: Medição 1
- Coluna 8: Medição 2
- Coluna 9: Medição 3
...
- Coluna 18: Medição 12
- Colunas Finais: Total Acumulado, Saldo a Medir, % Concluído
- Rodapé de cada coluna de medição:
  1. Total Medido no Mês / Quinzena (R$)
  2. Total Acumulado até a Medição (R$)
  3. Saldo Remanescente do Contrato (R$)
  4. % Avanço Físico Acumulado (%)
  5. (-) Retenção Técnica Contratual (5,0%)
  6. (=) Valor Líquido a Pagar na NF-e (R$)
"""

import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_DIR = os.path.join(BASE_DIR, "projetos", "OBRA_TMULT", "02_ORCAMENTO_BASE_E_CONTRATOS", "CONTRATOS_EMPREITEIROS")
os.makedirs(OUTPUT_DIR, exist_ok=True)

# -------------------------------------------------------------
# DADOS CONTRATUAIS DOS 8 PACOTES DE EMPREITEIROS DA OBRA TMULT
# -------------------------------------------------------------
PACOTES = [
    {
        "cod": "SUB-01",
        "tab_name": "SUB-01 Estrutura",
        "titulo": "EMPREITADA DE FUNDAÇÕES E ESTRUTURA DE CONCRETO ARMADO",
        "empreiteiro": "Empreiteira de Estruturas Homologada Ltda",
        "cc": "CC-200 / CC-300",
        "vigencia": "Semanas S01 a S09 (60 dias corridos)",
        "total_contrato": 195400.00,
        "servicos": [
            {"eap": "1.1.1", "desc": "Escavação manual e ajuste fino de fundo de cavas das sapatas", "und": "m³", "qtd": 51.27, "pu": 14.28},
            {"eap": "1.1.2", "desc": "Apiloamento e compactação do fundo de cavas e valas de fundação", "und": "m²", "qtd": 60.74, "pu": 8.65},
            {"eap": "1.1.5", "desc": "Montagem e desforma de painéis de fôrma compensado 17mm para 32 sapatas", "und": "m²", "qtd": 35.88, "pu": 55.00},
            {"eap": "1.1.7", "desc": "Montagem e desforma de fôrmas compensadas 17mm para arranques de pilares", "und": "m²", "qtd": 25.76, "pu": 65.00},
            {"eap": "1.1.9", "desc": "Montagem e desforma de fôrmas compensadas para vigas baldrames (140,44m)", "und": "m²", "qtd": 112.35, "pu": 55.00},
            {"eap": "1.1.10", "desc": "Armação e amarração de aço CA-50 Ø8,0mm nas 32 sapatas isoladas", "und": "kg", "qtd": 256.00, "pu": 4.50},
            {"eap": "1.1.11", "desc": "Armação e amarração de arranques dos pilares P1-P24 c/ gancho 30cm", "und": "kg", "qtd": 145.70, "pu": 4.50},
            {"eap": "1.1.12", "desc": "Armação e amarração de armaduras positivas/estribos vigas baldrames", "und": "kg", "qtd": 1077.30, "pu": 4.50},
            {"eap": "1.1.13", "desc": "Aplicação de emulsão asfáltica impermeabilizante em sapatas e baldrames", "und": "m²", "qtd": 183.36, "pu": 18.00},
            {"eap": "1.2.1", "desc": "Lançamento, adensamento vibrado e acabamento de concreto usinado pilares", "und": "m³", "qtd": 6.44, "pu": 120.00},
            {"eap": "1.2.2", "desc": "Montagem, prumo, travamento e desforma de fôrmas de pilares P1-P24 (4 faces)", "und": "m²", "qtd": 85.82, "pu": 70.00},
            {"eap": "1.2.3", "desc": "Lançamento e vibração de concreto fck 30 MPa vigas superiores V101-V115", "und": "m³", "qtd": 17.38, "pu": 110.00},
            {"eap": "1.2.4", "desc": "Montagem, nivelamento e desforma de fôrmas vigas superiores (fundo e laterais)", "und": "m²", "qtd": 208.62, "pu": 65.00},
            {"eap": "1.2.5", "desc": "Lançamento e sarrafeamento de concreto usinado na capa de laje e=5cm", "und": "m³", "qtd": 13.17, "pu": 110.00},
            {"eap": "1.2.6", "desc": "Montagem de escoramento metálico e assoalho compensado para laje treliçada", "und": "m²", "qtd": 263.48, "pu": 45.00},
            {"eap": "1.2.8", "desc": "Montagem de vigotas treliçadas TR 16745, lajotas EPS H12 e armadura capeamento", "und": "m²", "qtd": 298.25, "pu": 35.00},
            {"eap": "1.2.10", "desc": "Mão de obra global de cura úmida (7 dias) e desforma gradual das estruturas", "und": "vb", "qtd": 1.00, "pu": 128363.08}
        ],
        "simulacao": {
            "1.1.1": [51.27, 0, 0, 0], "1.1.2": [60.74, 0, 0, 0], "1.1.5": [35.88, 0, 0, 0],
            "1.1.7": [25.76, 0, 0, 0], "1.1.9": [56.17, 56.18, 0, 0], "1.1.10": [256.0, 0, 0, 0],
            "1.1.11": [145.70, 0, 0, 0], "1.1.12": [538.65, 538.65, 0, 0], "1.1.13": [0, 91.68, 91.68, 0],
            "1.2.1": [0, 0, 6.44, 0], "1.2.2": [0, 0, 85.82, 0], "1.2.3": [0, 0, 0, 17.38],
            "1.2.4": [0, 0, 104.31, 104.31], "1.2.5": [0, 0, 0, 13.17], "1.2.6": [0, 0, 131.74, 131.74],
            "1.2.8": [0, 0, 0, 298.25], "1.2.10": [0.20, 0.25, 0.25, 0.30]
        }
    },
    {
        "cod": "SUB-02",
        "tab_name": "SUB-02 Alvenaria",
        "titulo": "EMPREITADA DE ALVENARIA DE VEDAÇÃO E REVESTIMENTOS",
        "empreiteiro": "Empreiteira de Alvenaria e Fachadas Ltda",
        "cc": "CC-400 (Alvenaria e Revestimentos)",
        "vigencia": "Semanas S10 a S17 (60 dias corridos)",
        "total_contrato": 138600.00,
        "servicos": [
            {"eap": "1.3.1", "desc": "Alvenaria de vedação em blocos de concreto 14x19x39cm e=14cm", "und": "m²", "qtd": 450.00, "pu": 42.00},
            {"eap": "1.3.2", "desc": "Execução de vergas e contravergas pré-moldadas em vãos de portas e janelas", "und": "m", "qtd": 68.00, "pu": 25.00},
            {"eap": "1.3.3", "desc": "Encunhamento superior elástico de alvenaria com poliuretano expandido", "und": "m", "qtd": 140.00, "pu": 15.00},
            {"eap": "1.5.1", "desc": "Chapisco rolado industrializado com aditivo polimérico Bianco", "und": "m²", "qtd": 920.00, "pu": 8.00},
            {"eap": "1.5.2", "desc": "Emboço paulista mecanizado com projeção contínua sarrafeado e desempenado", "und": "m²", "qtd": 920.00, "pu": 28.00},
            {"eap": "1.5.3", "desc": "Mão de obra especializada de operação de andaimes, recortes e arremates", "und": "vb", "qtd": 1.00, "pu": 82780.00}
        ],
        "simulacao": {
            "1.3.1": [0, 0, 0, 0, 225.00, 225.00, 0, 0],
            "1.3.2": [0, 0, 0, 0, 34.00, 34.00, 0, 0],
            "1.3.3": [0, 0, 0, 0, 0, 140.00, 0, 0],
            "1.5.1": [0, 0, 0, 0, 0, 460.00, 460.00, 0],
            "1.5.2": [0, 0, 0, 0, 0, 0, 460.00, 460.00],
            "1.5.3": [0, 0, 0, 0, 0.25, 0.25, 0.25, 0.25]
        }
    },
    {
        "cod": "SUB-03",
        "tab_name": "SUB-03 Cobertura",
        "titulo": "EMPREITADA DE COBERTURA TERMOACÚSTICA E CALHAS",
        "empreiteiro": "Empreiteira de Coberturas Industriais Ltda",
        "cc": "CC-500 (Cobertura Metálica)",
        "vigencia": "Semanas S10 a S15 (45 dias corridos)",
        "total_contrato": 42500.00,
        "servicos": [
            {"eap": "1.4.1", "desc": "Montagem e fixação de terças metálicas em perfil U enrijecido galvanizado", "und": "kg", "qtd": 1850.00, "pu": 8.00},
            {"eap": "1.4.2", "desc": "Montagem e fixação de telhas sandwich termoacústicas PIR 30mm trapézio", "und": "m²", "qtd": 381.29, "pu": 38.00},
            {"eap": "1.4.3", "desc": "Instalação de calhas metálicas galvanizadas corte 50 e condutores", "und": "m", "qtd": 64.00, "pu": 45.00},
            {"eap": "1.4.4", "desc": "Instalação de rufos de platibanda, cumeeiras seladas e arremates estanques", "und": "m", "qtd": 82.00, "pu": 35.00},
            {"eap": "1.4.5", "desc": "Mão de obra global de montagem com linha de vida e teste de mangueira", "und": "vb", "qtd": 1.00, "pu": 7460.98}
        ],
        "simulacao": {
            "1.4.1": [0, 0, 0, 0, 1850.00, 0, 0],
            "1.4.2": [0, 0, 0, 0, 0, 381.29, 0],
            "1.4.3": [0, 0, 0, 0, 0, 64.00, 0],
            "1.4.4": [0, 0, 0, 0, 0, 0, 82.00],
            "1.4.5": [0, 0, 0, 0, 0.40, 0.40, 0.20]
        }
    },
    {
        "cod": "SUB-04",
        "tab_name": "SUB-04 Pisos",
        "titulo": "EMPREITADA DE PISOS, CONTRAPISO E REVESTIMENTOS CERÂMICOS",
        "empreiteiro": "Empreiteira de Pisos e Acabamentos Ltda",
        "cc": "CC-704 (Pisos e Pavimentações)",
        "vigencia": "Semanas S18 a S21 (30 dias corridos)",
        "total_contrato": 36800.00,
        "servicos": [
            {"eap": "1.7.1", "desc": "Execução de contrapiso de regularização farofa e=3cm com impermeabilizante", "und": "m²", "qtd": 368.40, "pu": 22.00},
            {"eap": "1.7.2", "desc": "Assentamento de porcelanato retificado 60x60cm dupla colagem com AC-III", "und": "m²", "qtd": 368.40, "pu": 48.00},
            {"eap": "1.7.3", "desc": "Instalação de rodapé em porcelanato h=10cm cortado com meia esquadria", "und": "m", "qtd": 180.00, "pu": 15.00},
            {"eap": "1.7.4", "desc": "Aplicação de rejunte flexível resinado e epóxi anti-fungo em áreas molhadas", "und": "m²", "qtd": 368.40, "pu": 8.00},
            {"eap": "1.7.5", "desc": "Proteção superficial de pisos com papelão ondulado e fita crepe", "und": "vb", "qtd": 1.00, "pu": 5364.80}
        ],
        "simulacao": {
            "1.7.1": [0, 0, 0, 0, 0, 0, 0, 0, 368.40, 0],
            "1.7.2": [0, 0, 0, 0, 0, 0, 0, 0, 0, 368.40],
            "1.7.3": [0, 0, 0, 0, 0, 0, 0, 0, 0, 180.00],
            "1.7.4": [0, 0, 0, 0, 0, 0, 0, 0, 0, 368.40],
            "1.7.5": [0, 0, 0, 0, 0, 0, 0, 0, 0, 1.00]
        }
    },
    {
        "cod": "SUB-05",
        "tab_name": "SUB-05 Eletrica",
        "titulo": "EMPREITADA DE INSTALAÇÕES ELÉTRICAS, CABEAMENTO E TELECOM",
        "empreiteiro": "Empreiteira de Engenharia Elétrica Ltda",
        "cc": "CC-604 / CC-605 / CC-606",
        "vigencia": "Semanas S12 a S21 (60 dias corridos)",
        "total_contrato": 48200.00,
        "servicos": [
            {"eap": "2.2.1", "desc": "Chumbamento de caixas 4x2 e eletrodutos de PVC rígido roscável embutidos", "und": "m", "qtd": 650.00, "pu": 12.00},
            {"eap": "2.2.3", "desc": "Enfiamento e puxamento de condutores de cobre flexível 750V de 1,5 a 50 mm²", "und": "m", "qtd": 3800.00, "pu": 4.50},
            {"eap": "2.2.5", "desc": "Montagem, barramentos de cobre e fiação interna dos quadros QDG e QDF", "und": "cj", "qtd": 2.00, "pu": 3200.00},
            {"eap": "2.2.6", "desc": "Instalação de 85 conjuntos de interruptores, tomadas 10A/20A e tampas 4x2", "und": "un", "qtd": 85.00, "pu": 25.00},
            {"eap": "2.2.7", "desc": "Instalação de 48 luminárias de embutir LED 60x60cm e spots decorativos", "und": "un", "qtd": 48.00, "pu": 35.00},
            {"eap": "2.2.8", "desc": "Passagem e conectorização RJ-45 de rede de dados e telecom Cat6 (Rack 12U)", "und": "pt", "qtd": 32.00, "pu": 45.00},
            {"eap": "2.2.9", "desc": "Mão de obra de testes, ensaios de isolamento, etiquetagem e ART elétrica", "und": "vb", "qtd": 1.00, "pu": 11655.00}
        ],
        "simulacao": {
            "2.2.1": [0, 0, 0, 0, 0, 325.00, 325.00, 0, 0, 0],
            "2.2.3": [0, 0, 0, 0, 0, 0, 1900.00, 1900.00, 0, 0],
            "2.2.5": [0, 0, 0, 0, 0, 0, 0, 1.00, 1.00, 0],
            "2.2.6": [0, 0, 0, 0, 0, 0, 0, 0, 42.00, 43.00],
            "2.2.7": [0, 0, 0, 0, 0, 0, 0, 0, 24.00, 24.00],
            "2.2.8": [0, 0, 0, 0, 0, 0, 0, 0, 16.00, 16.00],
            "2.2.9": [0, 0, 0, 0, 0, 0, 0.20, 0.30, 0.25, 0.25]
        }
    },
    {
        "cod": "SUB-06",
        "tab_name": "SUB-06 Hidraulica",
        "titulo": "EMPREITADA DE INSTALAÇÕES HIDROSSANITÁRIAS E DRENAGEM",
        "empreiteiro": "Empreiteira de Instalações Hidráulicas Ltda",
        "cc": "CC-601 / CC-602 / CC-603",
        "vigencia": "Semanas S12 a S20 (60 dias corridos)",
        "total_contrato": 38900.00,
        "servicos": [
            {"eap": "2.1.1", "desc": "Tubulações de água fria em PVC soldável de 20 a 50mm embutidas", "und": "m", "qtd": 240.00, "pu": 18.00},
            {"eap": "2.1.2", "desc": "Tubulações de esgoto sanitário e ventilação em PVC série normal e reforçada", "und": "m", "qtd": 180.00, "pu": 24.00},
            {"eap": "2.1.3", "desc": "Instalação de caixas sifonadas, ralos secos e caixas de inspeção/gordura", "und": "un", "qtd": 16.00, "pu": 85.00},
            {"eap": "2.1.4", "desc": "Instalação de louças sanitárias Deca (bacias acopladas e lavatórios de coluna)", "und": "un", "qtd": 12.00, "pu": 120.00},
            {"eap": "2.1.5", "desc": "Instalação de metais Docol (torneiras temporizadas e registros de gaveta)", "und": "un", "qtd": 18.00, "pu": 65.00},
            {"eap": "2.1.6", "desc": "Mão de obra especializada de testes hidrostáticos, manômetro e laudo técnico", "und": "vb", "qtd": 1.00, "pu": 26290.00}
        ],
        "simulacao": {
            "2.1.1": [0, 0, 0, 0, 0, 120.00, 120.00, 0, 0],
            "2.1.2": [0, 0, 0, 0, 0, 90.00, 90.00, 0, 0],
            "2.1.3": [0, 0, 0, 0, 0, 8.00, 8.00, 0, 0],
            "2.1.4": [0, 0, 0, 0, 0, 0, 0, 6.00, 6.00],
            "2.1.5": [0, 0, 0, 0, 0, 0, 0, 9.00, 9.00],
            "2.1.6": [0, 0, 0, 0, 0, 0.25, 0.25, 0.25, 0.25]
        }
    },
    {
        "cod": "SUB-07",
        "tab_name": "SUB-07 Climatizacao",
        "titulo": "EMPREITADA DE CLIMATIZAÇÃO E SISTEMA VRF/SPLIT",
        "empreiteiro": "Empreiteira de Refrigeração e HVAC Ltda",
        "cc": "CC-608 (Climatização e Ventilação)",
        "vigencia": "Semanas S20 a S24 (30 dias corridos)",
        "total_contrato": 24800.00,
        "servicos": [
            {"eap": "2.3.1", "desc": "Instalação de tubulação frigorígena de cobre sem costura isolada c/ elastomérico", "und": "m", "qtd": 160.00, "pu": 35.00},
            {"eap": "2.3.2", "desc": "Execução de rede de drenagem de condensado em PVC rígido soldável Ø25mm", "und": "m", "qtd": 80.00, "pu": 22.00},
            {"eap": "2.3.3", "desc": "Fixação de 8 evaporadoras e condensadoras com suportes metálicos c/ coxins", "und": "un", "qtd": 8.00, "pu": 450.00},
            {"eap": "2.3.4", "desc": "Serviço de pressurização com N2, vácuo profundo <500 microns e recolhimento", "und": "cj", "qtd": 8.00, "pu": 300.00},
            {"eap": "2.3.5", "desc": "Start-up, medição de corrente/tensão, teste de vazão de ar e emissão de ART", "und": "vb", "qtd": 1.00, "pu": 11440.00}
        ],
        "simulacao": {
            "2.3.1": [0, 0, 0, 0, 0, 0, 0, 0, 0, 80.00, 80.00, 0],
            "2.3.2": [0, 0, 0, 0, 0, 0, 0, 0, 0, 40.00, 40.00, 0],
            "2.3.3": [0, 0, 0, 0, 0, 0, 0, 0, 0, 4.00, 4.00, 0],
            "2.3.4": [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 8.00, 0],
            "2.3.5": [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0.50, 0.50]
        }
    },
    {
        "cod": "SUB-08",
        "tab_name": "SUB-08 Pintura",
        "titulo": "EMPREITADA DE PINTURA PREDIAL, MASSA CORRIDA E TEXTURAS",
        "empreiteiro": "Empreiteira de Pintura e Tratamento de Superfícies Ltda",
        "cc": "CC-706 (Pintura Predial)",
        "vigencia": "Semanas S20 a S25 (40 dias corridos)",
        "total_contrato": 39500.00,
        "servicos": [
            {"eap": "1.8.1", "desc": "Lixamento de superfícies de emboço e aplicação de fundo preparador/selador", "und": "m²", "qtd": 920.00, "pu": 6.00},
            {"eap": "1.8.2", "desc": "Aplicação de 2 demãos de massa corrida (PVA interna / Acrílica externa e WCs)", "und": "m²", "qtd": 920.00, "pu": 18.00},
            {"eap": "1.8.3", "desc": "Lixamento fino com lixadeira orbital c/ coletor de pó e limpeza de poeira", "und": "m²", "qtd": 920.00, "pu": 5.00},
            {"eap": "1.8.4", "desc": "Pintura em 2 demãos de látex acrílico premium acabamento fosco lavável", "und": "m²", "qtd": 920.00, "pu": 12.00},
            {"eap": "1.8.5", "desc": "Proteção geral de vidros e caixilhos com fita crepe azul e papel kraft", "und": "vb", "qtd": 1.00, "pu": 1780.00}
        ],
        "simulacao": {
            "1.8.1": [0, 0, 0, 0, 0, 0, 0, 0, 0, 460.00, 460.00, 0],
            "1.8.2": [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 460.00, 460.00],
            "1.8.3": [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 460.00, 460.00],
            "1.8.4": [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 920.00],
            "1.8.5": [0, 0, 0, 0, 0, 0, 0, 0, 0, 0.50, 0.25, 0.25]
        }
    }
]

NUM_MEDICOES = 12

def construir_planilha_medicao_completa():
    caminho = os.path.join(OUTPUT_DIR, "PLANILHA_MEDICAO_QUINZENAL_EMPREITEIROS.xlsx")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove aba em branco padrão

    # Estilos Visuais
    NAVY_HEADER = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    BLUE_HEADER = PatternFill(start_color="2B5B84", end_color="2B5B84", fill_type="solid")
    GOLD_HEADER = PatternFill(start_color="D99B26", end_color="D99B26", fill_type="solid")
    TEAL_HEADER = PatternFill(start_color="0E6655", end_color="0E6655", fill_type="solid")
    
    GRAY_LIGHT = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
    BLUE_LIGHT = PatternFill(start_color="E8EEF5", end_color="E8EEF5", fill_type="solid")
    GREEN_LIGHT = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
    YELLOW_LIGHT = PatternFill(start_color="FEF7E0", end_color="FEF7E0", fill_type="solid")
    
    FONT_TITLE = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    FONT_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    FONT_BOLD = Font(name="Calibri", size=10, bold=True, color="1B365D")
    FONT_REGULAR = Font(name="Calibri", size=10, color="333333")
    FONT_GREEN = Font(name="Calibri", size=10, bold=True, color="137333")
    FONT_GOLD = Font(name="Calibri", size=10, bold=True, color="B06000")
    
    BORDER_THIN = Border(
        left=Side(style='thin', color='DDDDDD'), right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'), bottom=Side(style='thin', color='DDDDDD')
    )
    BORDER_TOTAL = Border(
        top=Side(style='thin', color='1B365D'), bottom=Side(style='double', color='1B365D'),
        left=Side(style='thin', color='DDDDDD'), right=Side(style='thin', color='DDDDDD')
    )

    # =========================================================================
    # 1. ABA RESUMO GERAL CONSOLIDADO (DASHBOARD FINANCEIRO DA OBRA TMULT)
    # =========================================================================
    ws_dash = wb.create_sheet(title="Painel Geral 12 Medições")
    ws_dash.views.sheetView[0].showGridLines = True
    
    # Cabeçalho Principal do Painel
    ws_dash.merge_cells("A1:U1")
    ws_dash["A1"] = "PAINEL MASTER DE MEDIÇÕES EVOLUTIVAS DE EMPREITEIROS (12 QUINZENAS - R$ 564.700,00)"
    ws_dash["A1"].font = FONT_TITLE
    ws_dash["A1"].fill = NAVY_HEADER
    ws_dash["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_dash.row_dimensions[1].height = 28
    
    # Headers das colunas
    headers_dash = [
        "Cód", "Pacote de Empreitada", "Empreiteiro Homologado", "Centro Custo", "Total Contrato (R$)"
    ] + [f"Med {m}" for m in range(1, NUM_MEDICOES + 1)] + ["Total Acumulado (R$)", "Saldo a Pagar (R$)", "% Executado", "Retenção 5% Acum."]
    
    ws_dash.row_dimensions[3].height = 24
    for c_i, h_n in enumerate(headers_dash, start=1):
        cell = ws_dash.cell(row=3, column=c_i, value=h_n)
        cell.font = FONT_HEADER
        cell.fill = BLUE_HEADER if 6 <= c_i <= 17 else (GOLD_HEADER if c_i > 17 else NAVY_HEADER)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_THIN
        
    for p_idx, pac in enumerate(PACOTES, start=4):
        ws_dash.row_dimensions[p_idx].height = 20
        ws_dash.cell(row=p_idx, column=1, value=pac["cod"]).alignment = Alignment(horizontal="center")
        ws_dash.cell(row=p_idx, column=2, value=pac["titulo"]).alignment = Alignment(horizontal="left")
        ws_dash.cell(row=p_idx, column=3, value=pac["empreiteiro"]).alignment = Alignment(horizontal="left")
        ws_dash.cell(row=p_idx, column=4, value=pac["cc"]).alignment = Alignment(horizontal="center")
        
        # Total Contrato puxado da aba do empreiteiro
        tab = pac["tab_name"]
        num_items = len(pac["servicos"])
        r_tot = 6 + num_items  # linha do total bruto na aba do empreiteiro (start_row=6 + num_items)
        
        c_tc = ws_dash.cell(row=p_idx, column=5, value=f"='{tab}'!F{r_tot}")
        c_tc.number_format = '"R$ "#,##0.00'
        c_tc.font = FONT_BOLD
        c_tc.alignment = Alignment(horizontal="right")
        
        # Medições 1 a 12 (Colunas 6 a 17 no Dashboard, que correspondem a Col G até R na aba do pacote)
        for m in range(NUM_MEDICOES):
            col_dash = 6 + m
            col_pacote_letter = get_column_letter(7 + m)
            c_m = ws_dash.cell(row=p_idx, column=col_dash, value=f"='{tab}'!{col_pacote_letter}{r_tot}")
            c_m.number_format = '"R$ "#,##0.00'
            c_m.alignment = Alignment(horizontal="right")
            
        # Total Acumulado = soma de Med 1 a 12
        col_acum = 6 + NUM_MEDICOES
        c_ac = ws_dash.cell(row=p_idx, column=col_acum, value=f"=SUM(F{p_idx}:{get_column_letter(col_acum-1)}{p_idx})")
        c_ac.number_format = '"R$ "#,##0.00'
        c_ac.font = FONT_BOLD
        c_ac.alignment = Alignment(horizontal="right")
        
        # Saldo a Pagar = Total Contrato (E) - Total Acumulado
        col_sal = col_acum + 1
        c_sal = ws_dash.cell(row=p_idx, column=col_sal, value=f"=E{p_idx}-{get_column_letter(col_acum)}{p_idx}")
        c_sal.number_format = '"R$ "#,##0.00'
        c_sal.font = FONT_REGULAR
        c_sal.alignment = Alignment(horizontal="right")
        
        # % Executado = Acumulado / Total Contrato
        col_pct = col_sal + 1
        c_pct = ws_dash.cell(row=p_idx, column=col_pct, value=f"={get_column_letter(col_acum)}{p_idx}/E{p_idx}")
        c_pct.number_format = '0.00%'
        c_pct.font = FONT_BOLD
        c_pct.alignment = Alignment(horizontal="center")
        
        # Retenção 5% Acumulada
        col_ret = col_pct + 1
        c_ret = ws_dash.cell(row=p_idx, column=col_ret, value=f"={get_column_letter(col_acum)}{p_idx}*0.05")
        c_ret.number_format = '"R$ "#,##0.00'
        c_ret.font = FONT_GOLD
        c_ret.alignment = Alignment(horizontal="right")
        
        for ci in range(1, col_ret + 1):
            c_cell = ws_dash.cell(row=p_idx, column=ci)
            c_cell.border = BORDER_THIN
            if p_idx % 2 == 0:
                c_cell.fill = GRAY_LIGHT

    # Linha Totalizadora da Obra no Painel
    r_total_dash = 4 + len(PACOTES)
    ws_dash.row_dimensions[r_total_dash].height = 24
    ws_dash.merge_cells(start_row=r_total_dash, start_column=1, end_row=r_total_dash, end_column=4)
    ws_dash.cell(row=r_total_dash, column=1, value="TOTAL CONSOLIDADO DA OBRA TMULT (R$):").font = FONT_BOLD
    ws_dash.cell(row=r_total_dash, column=1).alignment = Alignment(horizontal="right", vertical="center")
    
    for c_i in range(5, col_ret + 1):
        letter = get_column_letter(c_i)
        if c_i == col_pct:
            cell = ws_dash.cell(row=r_total_dash, column=c_i, value=f"={get_column_letter(col_acum)}{r_total_dash}/E{r_total_dash}")
            cell.number_format = '0.00%'
            cell.alignment = Alignment(horizontal="center")
        else:
            cell = ws_dash.cell(row=r_total_dash, column=c_i, value=f"=SUM({letter}4:{letter}{r_total_dash-1})")
            cell.number_format = '"R$ "#,##0.00'
            cell.alignment = Alignment(horizontal="right")
        cell.font = FONT_BOLD
        cell.fill = BLUE_LIGHT
        cell.border = BORDER_TOTAL

    # Ajuste colunas Dashboard
    ws_dash.column_dimensions["A"].width = 10
    ws_dash.column_dimensions["B"].width = 38
    ws_dash.column_dimensions["C"].width = 32
    ws_dash.column_dimensions["D"].width = 18
    ws_dash.column_dimensions["E"].width = 18
    for m in range(NUM_MEDICOES):
        ws_dash.column_dimensions[get_column_letter(6 + m)].width = 15
    ws_dash.column_dimensions[get_column_letter(col_acum)].width = 18
    ws_dash.column_dimensions[get_column_letter(col_sal)].width = 18
    ws_dash.column_dimensions[get_column_letter(col_pct)].width = 14
    ws_dash.column_dimensions[get_column_letter(col_ret)].width = 18

    # =========================================================================
    # 2. ABAS INDIVIDUAIS DOS 8 EMPREITEIROS (FORMATO EXATO PEDIDO PELO USUÁRIO)
    # =========================================================================
    for pac in PACOTES:
        ws = wb.create_sheet(title=pac["tab_name"])
        ws.views.sheetView[0].showGridLines = True
        
        # 2.1 Título do Cabeçalho
        ws.merge_cells("A1:U1")
        ws["A1"] = f"PLANILHA DE MEDIÇÃO EVOLUTIVA: {pac['cod']} — {pac['titulo']}"
        ws["A1"].font = FONT_TITLE
        ws["A1"].fill = NAVY_HEADER
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 26
        
        # Metadados do Contrato
        meta = [
            ("Empreendimento:", "Edifício Administrativo TMULT (368,40 m²)", "Empreiteiro:", pac["empreiteiro"], "Contrato:", f"{pac['cod']} (R$ {pac['total_contrato']:,.2f})"),
            ("Localização:", "Terminal Multiuso - Porto do Açu (SJB/RJ)", "Centro Custo:", pac["cc"], "Prazo Vigência:", pac["vigencia"]),
            ("Engenheiro Fiscal:", "Eng. Alexandre (CREA-RJ)", "Regra Medição:", "POP 09 (Regra da Trena in-loco)", "Retenção Técnica:", "5,0% sobre valor bruto medido")
        ]
        for r_idx, (k1, v1, k2, v2, k3, v3) in enumerate(meta, start=2):
            ws.row_dimensions[r_idx].height = 18
            ws.cell(row=r_idx, column=1, value=k1).font = FONT_BOLD
            ws.cell(row=r_idx, column=2, value=v1).font = FONT_REGULAR
            ws.cell(row=r_idx, column=5, value=k2).font = FONT_BOLD
            ws.cell(row=r_idx, column=6, value=v2).font = FONT_REGULAR
            ws.cell(row=r_idx, column=10, value=k3).font = FONT_BOLD
            ws.cell(row=r_idx, column=11, value=v3).font = FONT_REGULAR

        # 2.2 Cabeçalhos das Colunas
        # Coluna 1: Item | Coluna 2: Descrição | Coluna 3: Unidade | Coluna 4: Quantidade | Coluna 5: Preço Unit | Coluna 6: Total Contrato
        # Coluna 7: Medição 1 | Coluna 8: Medição 2 | ... | Coluna 18: Medição 12
        # Coluna 19: Total Medido Acumulado | Coluna 20: Saldo a Medir | Coluna 21: % Concluído
        row_h = 5
        ws.row_dimensions[row_h].height = 28
        
        headers_base = [
            (1, "Item EAP"),
            (2, "Descrição do Serviço Contratado"),
            (3, "Unid"),
            (4, "Quantidade Contratada"),
            (5, "Preço Unitário (R$)"),
            (6, "Total Contrato (R$)")
        ]
        for c_idx, nome in headers_base:
            c = ws.cell(row=row_h, column=c_idx, value=nome)
            c.font = FONT_HEADER
            c.fill = NAVY_HEADER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER_THIN
            
        # Colunas 7 a 18: Medição 1 a 12
        for m in range(1, NUM_MEDICOES + 1):
            col_idx = 6 + m
            c = ws.cell(row=row_h, column=col_idx, value=f"Medição {m}")
            c.font = FONT_HEADER
            c.fill = BLUE_HEADER if m <= 4 else PatternFill(start_color="3A5A78", end_color="3A5A78", fill_type="solid")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER_THIN
            
        # Colunas 19 a 21: Total Acumulado, Saldo e %
        col_acum = 6 + NUM_MEDICOES + 1  # 19
        col_saldo = col_acum + 1          # 20
        col_pct = col_saldo + 1           # 21
        
        headers_pos = [
            (col_acum, "Total Medido Acumulado (Qtd)"),
            (col_saldo, "Saldo a Medir (Qtd)"),
            (col_pct, "% Avanço Concluído")
        ]
        for c_idx, nome in headers_pos:
            c = ws.cell(row=row_h, column=c_idx, value=nome)
            c.font = FONT_HEADER
            c.fill = GOLD_HEADER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER_THIN

        # 2.3 Linhas de Dados dos Serviços Contratados
        start_row = 6
        num_servicos = len(pac["servicos"])
        end_data_row = start_row + num_servicos - 1
        
        for s_idx, s in enumerate(pac["servicos"]):
            r = start_row + s_idx
            ws.row_dimensions[r].height = 20
            
            # Col 1: Item
            ws.cell(row=r, column=1, value=s["eap"]).alignment = Alignment(horizontal="center", vertical="center")
            # Col 2: Descrição
            ws.cell(row=r, column=2, value=s["desc"]).alignment = Alignment(horizontal="left", vertical="center")
            # Col 3: Unid
            ws.cell(row=r, column=3, value=s["und"]).alignment = Alignment(horizontal="center", vertical="center")
            # Col 4: Qtd
            c_qtd = ws.cell(row=r, column=4, value=s["qtd"])
            c_qtd.number_format = '#,##0.00'
            c_qtd.alignment = Alignment(horizontal="right", vertical="center")
            # Col 5: Preço Unit.
            c_pu = ws.cell(row=r, column=5, value=s["pu"])
            c_pu.number_format = '"R$ "#,##0.00'
            c_pu.alignment = Alignment(horizontal="right", vertical="center")
            # Col 6: Total Contrato = Qtd * PU (=D{r}*E{r})
            c_tot = ws.cell(row=r, column=6, value=f"=D{r}*E{r}")
            c_tot.number_format = '"R$ "#,##0.00'
            c_tot.font = FONT_BOLD
            c_tot.alignment = Alignment(horizontal="right", vertical="center")
            
            # Col 7 a 18: Medições 1 a 12 (Qtd medida na quinzena)
            for m in range(1, NUM_MEDICOES + 1):
                col_m = 6 + m
                sim_val = 0.0
                if "simulacao" in pac and s["eap"] in pac["simulacao"]:
                    sim_arr = pac["simulacao"][s["eap"]]
                    if m - 1 < len(sim_arr):
                        sim_val = sim_arr[m - 1]
                        
                c_cell = ws.cell(row=r, column=col_m, value=sim_val)
                c_cell.number_format = '#,##0.00'
                c_cell.alignment = Alignment(horizontal="right", vertical="center")
                if sim_val > 0:
                    c_cell.fill = YELLOW_LIGHT
                    c_cell.font = FONT_BOLD
                    
            # Col 19: Total Medido Acumulado = SOMA(G{r}:R{r})
            col_med_ini = get_column_letter(7)
            col_med_fim = get_column_letter(6 + NUM_MEDICOES)
            c_acum = ws.cell(row=r, column=col_acum, value=f"=SUM({col_med_ini}{r}:{col_med_fim}{r})")
            c_acum.number_format = '#,##0.00'
            c_acum.font = FONT_BOLD
            c_acum.alignment = Alignment(horizontal="right", vertical="center")
            
            # Col 20: Saldo a Medir = Qtd Contrato (D{r}) - Total Acumulado (S{r})
            letter_acum = get_column_letter(col_acum)
            c_saldo = ws.cell(row=r, column=col_saldo, value=f"=D{r}-{letter_acum}{r}")
            c_saldo.number_format = '#,##0.00'
            c_saldo.alignment = Alignment(horizontal="right", vertical="center")
            
            # Col 21: % Avanço = Total Acumulado / Qtd Contratada
            c_pct = ws.cell(row=r, column=col_pct, value=f"={letter_acum}{r}/D{r}")
            c_pct.number_format = '0.00%'
            c_pct.font = FONT_BOLD
            c_pct.alignment = Alignment(horizontal="center", vertical="center")
            
            # Bordas e zebra
            for ci in range(1, col_pct + 1):
                cell_item = ws.cell(row=r, column=ci)
                cell_item.border = BORDER_THIN
                if s_idx % 2 == 0 and cell_item.fill.start_color.index != 'FEF7E0':
                    cell_item.fill = GRAY_LIGHT

        # 2.4 LINHAS DE RODAPÉ (TOTALIZADORES E FECHAMENTOS POR MEDIÇÃO)
        # -------------------------------------------------------------
        # Linha 1: TOTAL MEDIDO NO MÊS / PERÍODO (R$)
        r_bruto = end_data_row + 1
        ws.row_dimensions[r_bruto].height = 24
        ws.merge_cells(start_row=r_bruto, start_column=1, end_row=r_bruto, end_column=5)
        ws.cell(row=r_bruto, column=1, value="TOTAL BRUTO MEDIDO NO PERÍODO (R$):").font = FONT_BOLD
        ws.cell(row=r_bruto, column=1).alignment = Alignment(horizontal="right", vertical="center")
        
        # Total Contrato somado
        ws.cell(row=r_bruto, column=6, value=f"=SUM(F{start_row}:F{end_data_row})").number_format = '"R$ "#,##0.00'
        ws.cell(row=r_bruto, column=6).font = FONT_BOLD
        ws.cell(row=r_bruto, column=6).alignment = Alignment(horizontal="right", vertical="center")
        
        # Para cada coluna de Medição (7 a 18), calcula o Total R$ Medido via SOMARPRODUTO (Preço Unit * Qtd Medida)
        for m in range(1, NUM_MEDICOES + 1):
            col_m = 6 + m
            col_letter = get_column_letter(col_m)
            # FÓRMULA OFICIAL EXCEL: =SUMPRODUCT($E$6:$E$end, G6:Gend)
            c_bruto = ws.cell(
                row=r_bruto, column=col_m,
                value=f"=SUMPRODUCT($E${start_row}:$E${end_data_row},{col_letter}${start_row}:{col_letter}${end_data_row})"
            )
            c_bruto.number_format = '"R$ "#,##0.00'
            c_bruto.font = FONT_BOLD
            c_bruto.alignment = Alignment(horizontal="right", vertical="center")
            
        # Colunas finais da linha de Total Bruto
        # Total Acumulado em R$
        c_tot_acum_r = ws.cell(
            row=r_bruto, column=col_acum,
            value=f"=SUM({get_column_letter(7)}{r_bruto}:{get_column_letter(6+NUM_MEDICOES)}{r_bruto})"
        )
        c_tot_acum_r.number_format = '"R$ "#,##0.00'
        c_tot_acum_r.font = FONT_BOLD
        c_tot_acum_r.alignment = Alignment(horizontal="right", vertical="center")
        
        # Saldo do Contrato em R$
        c_tot_saldo_r = ws.cell(
            row=r_bruto, column=col_saldo,
            value=f"=F{r_bruto}-{get_column_letter(col_acum)}{r_bruto}"
        )
        c_tot_saldo_r.number_format = '"R$ "#,##0.00'
        c_tot_saldo_r.font = FONT_BOLD
        c_tot_saldo_r.alignment = Alignment(horizontal="right", vertical="center")
        
        # % Geral Executado
        c_tot_pct = ws.cell(
            row=r_bruto, column=col_pct,
            value=f"={get_column_letter(col_acum)}{r_bruto}/F{r_bruto}"
        )
        c_tot_pct.number_format = '0.00%'
        c_tot_pct.font = FONT_BOLD
        c_tot_pct.alignment = Alignment(horizontal="center", vertical="center")
        
        for ci in range(1, col_pct + 1):
            cell_b = ws.cell(row=r_bruto, column=ci)
            cell_b.fill = BLUE_LIGHT
            cell_b.border = BORDER_TOTAL

        # Linha 2: TOTAL ACUMULADO ATÉ A MEDIÇÃO (R$)
        r_acum = r_bruto + 1
        ws.row_dimensions[r_acum].height = 20
        ws.merge_cells(start_row=r_acum, start_column=1, end_row=r_acum, end_column=6)
        ws.cell(row=r_acum, column=1, value="TOTAL ACUMULADO ATÉ A MEDIÇÃO (R$):").font = FONT_BOLD
        ws.cell(row=r_acum, column=1).alignment = Alignment(horizontal="right", vertical="center")
        
        for m in range(1, NUM_MEDICOES + 1):
            col_m = 6 + m
            col_letter = get_column_letter(col_m)
            col_ini_letter = get_column_letter(7)
            # Soma do total bruto de Med 1 até Med M
            c_ac = ws.cell(row=r_acum, column=col_m, value=f"=SUM({col_ini_letter}${r_bruto}:{col_letter}${r_bruto})")
            c_ac.number_format = '"R$ "#,##0.00'
            c_ac.font = FONT_BOLD
            c_ac.alignment = Alignment(horizontal="right", vertical="center")

        # Linha 3: SALDO REMANESCENTE DO CONTRATO (R$)
        r_saldo = r_acum + 1
        ws.row_dimensions[r_saldo].height = 20
        ws.merge_cells(start_row=r_saldo, start_column=1, end_row=r_saldo, end_column=6)
        ws.cell(row=r_saldo, column=1, value="SALDO DO CONTRATO APÓS MEDIÇÃO (R$):").font = FONT_BOLD
        ws.cell(row=r_saldo, column=1).alignment = Alignment(horizontal="right", vertical="center")
        
        for m in range(1, NUM_MEDICOES + 1):
            col_m = 6 + m
            col_letter = get_column_letter(col_m)
            # Saldo = Total Contrato ($F$r_bruto) - Acumulado nesta medição
            c_sd = ws.cell(row=r_saldo, column=col_m, value=f"=$F${r_bruto}-{col_letter}{r_acum}")
            c_sd.number_format = '"R$ "#,##0.00'
            c_sd.font = FONT_REGULAR
            c_sd.alignment = Alignment(horizontal="right", vertical="center")

        # Linha 4: % AVANÇO FÍSICO ACUMULADO (%)
        r_pct_acum = r_saldo + 1
        ws.row_dimensions[r_pct_acum].height = 20
        ws.merge_cells(start_row=r_pct_acum, start_column=1, end_row=r_pct_acum, end_column=6)
        ws.cell(row=r_pct_acum, column=1, value="% AVANÇO ACUMULADO NO CONTRATO:").font = FONT_BOLD
        ws.cell(row=r_pct_acum, column=1).alignment = Alignment(horizontal="right", vertical="center")
        
        for m in range(1, NUM_MEDICOES + 1):
            col_m = 6 + m
            col_letter = get_column_letter(col_m)
            # % Acumulado = Acumulado / Total Contrato
            c_pc = ws.cell(row=r_pct_acum, column=col_m, value=f"={col_letter}{r_acum}/$F${r_bruto}")
            c_pc.number_format = '0.00%'
            c_pc.font = FONT_BOLD
            c_pc.alignment = Alignment(horizontal="center", vertical="center")

        # Linha 5: RETENÇÃO TÉCNICA 5% (CAUÇÃO CONTRATUAL)
        r_ret = r_pct_acum + 1
        ws.row_dimensions[r_ret].height = 20
        ws.merge_cells(start_row=r_ret, start_column=1, end_row=r_ret, end_column=6)
        ws.cell(row=r_ret, column=1, value="(-) RETENÇÃO TÉCNICA DE GARANTIA (5,0%):").font = FONT_BOLD
        ws.cell(row=r_ret, column=1).alignment = Alignment(horizontal="right", vertical="center")
        
        for m in range(1, NUM_MEDICOES + 1):
            col_m = 6 + m
            col_letter = get_column_letter(col_m)
            # Retenção = Total Bruto Medido no Mês * 5%
            c_rt = ws.cell(row=r_ret, column=col_m, value=f"={col_letter}{r_bruto}*0.05")
            c_rt.number_format = '"R$ "#,##0.00'
            c_rt.font = FONT_GOLD
            c_rt.alignment = Alignment(horizontal="right", vertical="center")

        # Linha 6: VALOR LÍQUIDO A PAGAR NA QUINZENA (NF-e)
        r_liq = r_ret + 1
        ws.row_dimensions[r_liq].height = 24
        ws.merge_cells(start_row=r_liq, start_column=1, end_row=r_liq, end_column=6)
        ws.cell(row=r_liq, column=1, value="(=) VALOR LÍQUIDO A LIBERAR NA NF-e:").font = FONT_BOLD
        ws.cell(row=r_liq, column=1).alignment = Alignment(horizontal="right", vertical="center")
        
        for m in range(1, NUM_MEDICOES + 1):
            col_m = 6 + m
            col_letter = get_column_letter(col_m)
            # Líquido = Total Bruto Medido no Mês - Retenção
            c_lq = ws.cell(row=r_liq, column=col_m, value=f"={col_letter}{r_bruto}-{col_letter}{r_ret}")
            c_lq.number_format = '"R$ "#,##0.00'
            c_lq.font = FONT_GREEN
            c_lq.fill = GREEN_LIGHT
            c_lq.alignment = Alignment(horizontal="right", vertical="center")

        for r_foot in [r_acum, r_saldo, r_pct_acum, r_ret, r_liq]:
            for ci in range(1, col_pct + 1):
                ws.cell(row=r_foot, column=ci).border = BORDER_THIN

        # Congelamento de Painéis: Trava colunas 1 a 6 (Item até Total Contrato) e linhas 1 a 5 de cabeçalho
        ws.freeze_panes = "G6"

        # Ajustes de Largura das Colunas
        ws.column_dimensions["A"].width = 11
        ws.column_dimensions["B"].width = 46
        ws.column_dimensions["C"].width = 8
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 16
        ws.column_dimensions["F"].width = 18
        for m in range(1, NUM_MEDICOES + 1):
            ws.column_dimensions[get_column_letter(6 + m)].width = 15
        ws.column_dimensions[get_column_letter(col_acum)].width = 18
        ws.column_dimensions[get_column_letter(col_saldo)].width = 16
        ws.column_dimensions[get_column_letter(col_pct)].width = 14

    # Salva Pasta de Trabalho Completa
    wb.save(caminho)
    print(f"Planilha Master gerada com sucesso em: {caminho}")

if __name__ == "__main__":
    construir_planilha_medicao_completa()
