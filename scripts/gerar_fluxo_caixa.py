#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Motor Universal de Geração de Fluxo de Caixa, Curva de Desembolso e Capital de Giro.
Arquitetura Modular v2.0 Lean (Manual de Boas Práticas §2).

Uso:
    python scripts/gerar_fluxo_caixa.py --obra OBRA_TMULT
    python scripts/gerar_fluxo_caixa.py --dir projetos/OBRA_TMULT
"""

import os
import sys
import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference
import pandas as pd

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from scripts.common.obra_io import (
    parse_obra_args, resolver_obra_dir, carregar_config_obra, salvar_csv_utf8_sig, salvar_markdown
)
from scripts.common.excel_theme import (
    NAVY, GOLD_ACCENT, GRAY_LIGHT, GREEN_FILL, RED_FILL, BLUE_LIGHT, YELLOW_ACCENT,
    FONT_TITLE, FONT_HEADER, FONT_BOLD, FONT_REGULAR, FONT_RED, FONT_GREEN,
    THIN_BORDER, DOUBLE_BOTTOM_BORDER, ALIGN_CENTER, ALIGN_LEFT, ALIGN_RIGHT, formatar_moeda
)

RUBRICAS_ENTRADAS = [
    ("1.1 Faturamento Bruto de Medições", "Receita Bruta", "faturamento_bruto", "1.1 Faturamento Bruto Previsto"),
    ("1.2 Retenção Técnica Contratual (5%)", "Dedução", "retencao_5pct", "1.2 Retenção Contratual (5,0%)"),
    ("1.3 Faturamento Líquido de Medições", "Receita Líquida", "faturamento_liquido", "1.3 Faturamento Líquido (95%)"),
    ("1.4 Recebimento Efetivo de Medições (D+15)", "Inflow", "recebimento_medicoes", "1.4 Recebimento Efetivo (D+15)"),
    ("1.5 Devolução da Retenção Técnica (TRD M7)", "Inflow", "devolucao_retencao", "1.5 Devolução da Retenção (M{m_fim})"),
]

RUBRICAS_SAIDAS = [
    ("2.1 Equipe Técnica de Gestão (Eng/Mestre/TST)", "MO Gestão", "des_mo_gestao", "2.1 Equipe Gestão (Eng/Mestre/TST)"),
    ("2.2 Vivência, Alimentação (16 op.) e Transporte", "Vivência", "des_vivencia", "2.2 Vivência e Alimentação"),
    ("2.3 Mão de Obra de Campo (Serviços Físicos)", "MO Campo", "des_mo_campo", "2.3 Mão de Obra de Campo (Físico)"),
    ("2.4 Materiais e Insumos Civis (Sinal à Vista)", "Material", "des_mat_vista", "2.4 Materiais (À Vista no Pedido)"),
    ("2.5 Materiais e Insumos Civis (Prazo D+30)", "Material", "des_mat_prazo", "2.5 Materiais (A Prazo D+30)"),
    ("2.6 Locação de Containers NR-18 e Sanitários", "Equipamento", "des_locacoes", "2.6 Containers NR-18 e Sanitários"),
    ("2.7 Contas de Consumo (Água, Luz, Internet)", "Canteiro", "des_contas", "2.7 Consumo Canteiro (Água/Luz/Net)"),
    ("2.8 Tributos sobre Faturamento (8,65% NF)", "Tributos", "des_impostos", "2.8 Tributos s/ Faturamento (8,65%)"),
    ("2.9 Custos Indiretos Centrais BDI (AC/Seguros)", "Indireto", "des_ac_seguros", "2.9 Custos Indiretos Centrais BDI"),
]


def carregar_dados_cronograma(cronograma_csv, prazo_meses=6, bdi_servico=0.2717, bdi_equip=0.15, config=None):
    """Lê cronograma físico-financeiro e calcula vendas e custos diretos por mês."""
    df = pd.read_csv(cronograma_csv, sep=';', encoding='utf-8')
    eap_col = [c for c in df.columns if 'EAP' in c][0]
    _conv = lambda v: float(v.replace('R$', '').replace('.', '').replace(',', '.').strip()) if isinstance(v, str) else float(v)

    vendas_m = {m: round(df[f'R$ Mês {m}'].apply(_conv).sum(), 2) if f'R$ Mês {m}' in df.columns else 0.0 for m in range(1, prazo_meses + 1)}
    total_venda = sum(vendas_m.values())

    diretos_m = {}
    for m in range(1, prazo_meses + 1):
        col_m = f'R$ Mês {m}'
        diretos_m[m] = round(sum(
            _conv(r[col_m]) / (1.0 + (bdi_equip if str(r[eap_col]).strip().startswith('3.3') else bdi_servico))
            for _, r in df.iterrows()
        ), 2) if col_m in df.columns else 0.0
    total_custo_direto = sum(diretos_m.values())

    params = config.get("parametros_fluxo_caixa", {}) if config else {}
    target_cd = params.get("total_custo_direto_alvo")
    if target_cd is not None and total_custo_direto > 0:
        target_cd = float(target_cd)
        fator = target_cd / total_custo_direto
        diretos_m = {m: round(diretos_m[m] * fator, 2) for m in range(1, prazo_meses + 1)}
        diretos_m[prazo_meses] = round(diretos_m[prazo_meses] + (target_cd - sum(diretos_m.values())), 2)
        total_custo_direto = target_cd

    return vendas_m, diretos_m, total_venda, total_custo_direto


def modelar_fluxo_caixa(vendas_m, diretos_m, total_venda, total_custo_direto, config, prazo_meses=6):
    """Simula o fluxo de caixa nos cenários Realista, Otimista e Estresse."""
    meses = list(range(1, prazo_meses + 2))
    params = config.get("parametros_fluxo_caixa", {})
    retencao_pct = float(params.get("retencao_pct", 0.05))
    adiant_pct = float(params.get("adiantamento_pct", 0.10))
    aliq_imp = float(params.get("aliquota_impostos_pct", 0.0865))
    particao_mo = float(params.get("particao_mo_pct", 0.35))
    particao_mat = float(params.get("particao_mat_pct", 0.65))

    adm = config.get("administracao_local", [])
    c_g = sum(float(i.get("custo_unitario", 0)) for i in adm if i.get("eap") in ["1.0.1", "1.0.2"])
    c_v = sum(float(i.get("custo_unitario", 0)) for i in adm if i.get("eap") in ["1.0.5", "1.0.6"])
    c_l = sum(float(i.get("custo_unitario", 0)) for i in adm if i.get("eap") == "1.0.3")
    c_c = sum(float(i.get("custo_unitario", 0)) for i in adm if i.get("eap") == "1.0.4")
    c_cant_m = c_g + c_v + c_l + c_c
    bdi_ind_m = round((total_venda * params.get("bdi_indireto_pct", 0.0594)) / prazo_meses, 2)

    # 1. ENTRADAS REALISTA
    faturamento_bruto = {m: vendas_m.get(m, 0.0) for m in meses}
    retencao = {m: round(vendas_m.get(m, 0.0) * retencao_pct, 2) for m in meses}
    fat_liq = {m: round(vendas_m.get(m, 0.0) * (1.0 - retencao_pct), 2) for m in meses}
    tot_ret = sum(retencao.values())

    rec_med = {m: 0.0 for m in meses}
    for m in range(1, prazo_meses + 1):
        rec_med[m + 1] = fat_liq[m]
    dev_ret = {m: 0.0 for m in meses}
    dev_ret[prazo_meses + 1] = tot_ret
    ent_real = {m: round(rec_med[m] + dev_ret[m], 2) for m in meses}

    # 2. SAÍDAS REALISTA
    des_g = {m: 0.0 for m in meses}
    des_v = {m: 0.0 for m in meses}
    des_mc = {m: 0.0 for m in meses}
    des_mv = {m: 0.0 for m in meses}
    des_mp = {m: 0.0 for m in meses}
    des_loc = {m: 0.0 for m in meses}
    des_con = {m: 0.0 for m in meses}
    des_imp = {m: 0.0 for m in meses}
    des_ac = {m: 0.0 for m in meses}

    for m in range(1, prazo_meses + 1):
        des_g[m] += round(c_g, 2)
        des_v[m] += round(c_v, 2)
        des_loc[m] += round(c_l * 0.50, 2)
        des_loc[m + 1] += round(c_l * 0.50, 2)
        des_con[m] += round(c_c * 0.50, 2)
        des_con[m + 1] += round(c_c * 0.50, 2)

        custo_fis = max(0.0, diretos_m[m] - c_cant_m)
        mo_c = round(custo_fis * particao_mo, 2)
        mat_c = round(custo_fis * particao_mat, 2)

        des_mc[m] += mo_c
        des_mv[m] += round(mat_c * 0.20, 2)
        des_mp[m + 1] += round(mat_c * 0.80, 2)

        des_imp[m + 1] += round(vendas_m[m] * aliq_imp, 2)
        des_ac[m] += bdi_ind_m

    sai_real = {
        m: round(des_g[m] + des_v[m] + des_mc[m] + des_mv[m] + des_mp[m] + des_loc[m] + des_con[m] + des_imp[m] + des_ac[m], 2)
        for m in meses
    }
    saldo_per, saldo_acum_real, acum = {}, {}, 0.0
    for m in meses:
        liq = round(ent_real[m] - sai_real[m], 2)
        acum = round(acum + liq, 2)
        saldo_per[m], saldo_acum_real[m] = liq, acum

    # 3. CENÁRIO OTIMISTA (Adiantamento)
    adiant = round(total_venda * adiant_pct, 2)
    amort_m = round(adiant / 4.0, 2)
    ent_oti = {m: 0.0 for m in meses}
    ent_oti[1] += adiant
    for m in range(1, prazo_meses + 1):
        amort = amort_m if m <= 4 else 0.0
        ent_oti[m + 1] += round((vendas_m[m] * (1.0 - retencao_pct)) - amort, 2)
    ent_oti[prazo_meses + 1] += tot_ret

    saldo_acum_oti, ac_o = {}, 0.0
    for m in meses:
        ac_o = round(ac_o + ent_oti[m] - sai_real[m], 2)
        saldo_acum_oti[m] = ac_o

    # 4. CENÁRIO ESTRESSE (Atraso 30d)
    meses_pess = list(range(1, prazo_meses + 3))
    ent_pess = {m: 0.0 for m in meses_pess}
    for m in range(1, prazo_meses + 1):
        ent_pess[m + 2] += fat_liq[m]
    ent_pess[prazo_meses + 2] += tot_ret
    sai_pess = {m: sai_real.get(m, 0.0) for m in meses_pess}
    saldo_acum_pess, ac_p = {}, 0.0
    for m in meses_pess:
        ac_p = round(ac_p + ent_pess[m] - sai_pess[m], 2)
        saldo_acum_pess[m] = ac_p

    return {
        "meses": meses, "prazo_meses": prazo_meses, "vendas_m": vendas_m, "diretos_m": diretos_m,
        "faturamento_bruto": faturamento_bruto, "retencao_5pct": retencao, "faturamento_liquido": fat_liq,
        "total_retencao": tot_ret, "recebimento_medicoes": rec_med, "devolucao_retencao": dev_ret,
        "total_entradas_realista": ent_real, "des_mo_gestao": des_g, "des_vivencia": des_v,
        "des_mo_campo": des_mc, "des_mat_vista": des_mv, "des_mat_prazo": des_mp,
        "des_locacoes": des_loc, "des_contas": des_con, "des_impostos": des_imp, "des_ac_seguros": des_ac,
        "total_saidas_realista": sai_real, "saldo_periodo_realista": saldo_per, "saldo_acumulado_realista": saldo_acum_real,
        "saldo_acumulado_otimista": saldo_acum_oti, "meses_pess": meses_pess, "total_entradas_pessimista": ent_pess,
        "total_saidas_pessimista": sai_pess, "saldo_acumulado_pessimista": saldo_acum_pess,
        "total_venda": total_venda, "total_custo_direto": total_custo_direto
    }


def gerar_csv(fluxo, output_dir, sigla):
    """Gera o arquivo CSV oficial de Fluxo de Caixa no padrão utf-8-sig."""
    meses = fluxo["meses"]
    pad = [""] * (len(meses) + 1)
    header = ["Conta / Rubrica Financeira", "Natureza"] + [f"Mês {m}" for m in meses] + ["Total Consolidado (R$)"]
    _row = lambda t, n, k: [t, n] + [f"{fluxo[k].get(m, 0.0):.2f}" for m in meses] + [f"{sum(fluxo[k].get(m, 0.0) for m in meses):.2f}"]

    linhas = [
        ["1. ENTRADAS DE CAIXA (RECEBIMENTOS)", "GRUPO"] + pad,
        *[_row(t, n, k) for t, n, k, _ in RUBRICAS_ENTRADAS],
        ["TOTAL DE ENTRADAS DE CAIXA (A)", "Total Inflow"] + [f"{fluxo['total_entradas_realista'].get(m, 0.0):.2f}" for m in meses] + [f"{sum(fluxo['total_entradas_realista'].values()):.2f}"],
        [""] * (len(meses) + 3),
        ["2. SAÍDAS DE CAIXA (DESEMBOLSOS)", "GRUPO"] + pad,
        *[_row(t, n, k) for t, n, k, _ in RUBRICAS_SAIDAS],
        ["TOTAL DE SAÍDAS DE CAIXA (B)", "Total Outflow"] + [f"{fluxo['total_saidas_realista'].get(m, 0.0):.2f}" for m in meses] + [f"{sum(fluxo['total_saidas_realista'].values()):.2f}"],
        [""] * (len(meses) + 3),
        ["3. BALANÇO E SALDOS DE CAIXA", "GRUPO"] + pad,
        ["3.1 Saldo Operacional Líquido do Mês (A - B)", "Resultado"] + [f"{fluxo['saldo_periodo_realista'].get(m, 0.0):.2f}" for m in meses] + [f"{sum(fluxo['saldo_periodo_realista'].values()):.2f}"],
        ["3.2 Saldo de Caixa Acumulado (Cenário Realista)", "Acumulado"] + [f"{fluxo['saldo_acumulado_realista'][m]:.2f}" for m in meses] + [f"{fluxo['saldo_acumulado_realista'][meses[-1]]:.2f}"],
        ["3.3 Saldo Acumulado (Cenário Otimista c/ Adiantamento)", "Sensibilidade"] + [f"{fluxo['saldo_acumulado_otimista'][m]:.2f}" for m in meses] + [f"{fluxo['saldo_acumulado_otimista'][meses[-1]]:.2f}"],
        ["3.4 Saldo Acumulado (Cenário Estresse Atraso 30d)", "Sensibilidade"] + [f"{fluxo['saldo_acumulado_pessimista'][m]:.2f}" for m in meses] + [f"{fluxo['saldo_acumulado_pessimista'][meses[-1]]:.2f}"]
    ]
    salvar_csv_utf8_sig(os.path.join(output_dir, f"FLUXO_DE_CAIXA_{sigla}.csv"), header, linhas)


def _sc(ws, r, c, val, font=FONT_REGULAR, fill=None, align=ALIGN_LEFT, border=THIN_BORDER, num_fmt=None):
    cell = ws.cell(row=r, column=c, value=val)
    if font: cell.font = font
    if fill: cell.fill = fill
    if align: cell.alignment = align
    if border: cell.border = border
    if num_fmt: cell.number_format = num_fmt
    return cell


def gerar_xlsx(fluxo, output_dir, sigla, titulo_obra):
    """Gera a planilha executiva formatada com 3 abas, KPIs e gráficos analíticos."""
    xlsx_path = os.path.join(output_dir, f"FLUXO_DE_CAIXA_{sigla}.xlsx")
    wb = openpyxl.Workbook()
    prazo_m, m_fim = fluxo["prazo_meses"], fluxo["prazo_meses"] + 1

    # ABA 1: RESUMO EXECUTIVO
    ws1 = wb.active
    ws1.title, ws1.views.sheetView[0].showGridLines = "Resumo Executivo", True
    ws1.merge_cells("B2:I2")
    _sc(ws1, 2, 2, "RELATÓRIO FINANCEIRO EXECUTIVO — FLUXO DE CAIXA E CAPITAL DE GIRO", FONT_TITLE, NAVY, ALIGN_CENTER)
    ws1.row_dimensions[2].height, ws1.row_dimensions[3].height = 35, 22
    ws1.merge_cells("B3:I3")
    _sc(ws1, 3, 2, f"Empreendimento: {titulo_obra} | Prazo: {prazo_m} Meses ({prazo_m * 30} Dias) | Baseline 01",
        Font(name="Calibri", size=11, bold=True, color="FFFFFF"), GOLD_ACCENT, ALIGN_CENTER)

    pico_r = min(fluxo["saldo_acumulado_realista"].values())
    m_pico = min(fluxo["saldo_acumulado_realista"], key=fluxo["saldo_acumulado_realista"].get)
    s_fim = fluxo["saldo_acumulado_realista"][m_fim]
    m_liq = (s_fim / fluxo["total_venda"]) * 100 if fluxo["total_venda"] > 0 else 0.0

    kpis = [
        ("Preço Global Turnkey da Obra", fluxo["total_venda"], '"R$ "#,##0.00'),
        ("Custo Direto Total da Obra", fluxo["total_custo_direto"], '"R$ "#,##0.00'),
        ("Total de Impostos s/ Faturamento (8,65%)", sum(fluxo["des_impostos"].values()), '"R$ "#,##0.00'),
        ("Despesas Administrativas Centrais BDI", sum(fluxo["des_ac_seguros"].values()), '"R$ "#,##0.00'),
        ("Lucro Líquido Operacional Realizado", s_fim, '"R$ "#,##0.00'),
        ("Margem Líquida Realizada sobre Venda", m_liq / 100.0, '0.00%'),
        ("Retenção Contratual Técnica (5,0%)", fluxo["total_retencao"], '"R$ "#,##0.00'),
        ("NECESSIDADE MÍNIMA DE CAPITAL DE GIRO", abs(pico_r), '"R$ "#,##0.00'),
        ("Mês de Máxima Exposição de Caixa", f"Mês {m_pico} (Pico da Envoltória)", "@"),
        ("Ponto de Equilíbrio do Caixa (Break-even)", f"Mês {m_fim} (Liquidação do TRD)", "@")
    ]
    ws1.merge_cells("B5:E5")
    _sc(ws1, 5, 2, "INDICADORES FINANCEIROS E PARÂMETROS DE LIQUIDEZ", FONT_HEADER, NAVY, ALIGN_CENTER)
    for i, (kpi, val, fmt) in enumerate(kpis, start=1):
        r, fill = 5 + i, RED_FILL if "CAPITAL DE GIRO" in kpi else (GREEN_FILL if "Lucro" in kpi else (GRAY_LIGHT if i % 2 == 0 else None))
        f_lbl = FONT_BOLD if ("CAPITAL DE GIRO" in kpi or "Lucro" in kpi) else FONT_REGULAR
        f_val = FONT_RED if "CAPITAL DE GIRO" in kpi else (FONT_GREEN if "Lucro" in kpi else FONT_REGULAR)
        ws1.merge_cells(f"B{r}:D{r}")
        _sc(ws1, r, 2, kpi, font=f_lbl, fill=fill, align=ALIGN_LEFT)
        for c in (3, 4): _sc(ws1, r, c, None, fill=fill)
        _sc(ws1, r, 5, val, font=f_val, fill=fill, align=ALIGN_RIGHT, num_fmt=fmt)

    for c_i, h in enumerate(["Período", "Entradas (R$)", "Saídas (R$)", "Saldo Mês (R$)", "Saldo Acumulado (R$)"], start=2):
        _sc(ws1, 17, c_i, h, FONT_HEADER, NAVY, ALIGN_CENTER)
    for idx, m in enumerate(fluxo["meses"], start=1):
        r, fill = 17 + idx, GRAY_LIGHT if idx % 2 == 0 else None
        _sc(ws1, r, 2, f"Mês {m}", FONT_REGULAR, fill, ALIGN_CENTER)
        _sc(ws1, r, 3, fluxo["total_entradas_realista"][m], FONT_REGULAR, fill, ALIGN_RIGHT, num_fmt='"R$ "#,##0.00')
        _sc(ws1, r, 4, fluxo["total_saidas_realista"][m], FONT_REGULAR, fill, ALIGN_RIGHT, num_fmt='"R$ "#,##0.00')
        _sc(ws1, r, 5, fluxo["saldo_periodo_realista"][m], FONT_REGULAR, fill, ALIGN_RIGHT, num_fmt='"R$ "#,##0.00')
        f_ac = FONT_RED if fluxo["saldo_acumulado_realista"][m] < 0 else FONT_GREEN
        _sc(ws1, r, 6, fluxo["saldo_acumulado_realista"][m], f_ac, fill, ALIGN_RIGHT, num_fmt='"R$ "#,##0.00')

    r_tot = 18 + len(fluxo["meses"])
    _sc(ws1, r_tot, 2, "TOTAL GERAL", FONT_BOLD, GREEN_FILL, ALIGN_CENTER, DOUBLE_BOTTOM_BORDER)
    for c_i, val in enumerate([sum(fluxo["total_entradas_realista"].values()), sum(fluxo["total_saidas_realista"].values()), s_fim, s_fim], start=3):
        _sc(ws1, r_tot, c_i, val, FONT_BOLD, GREEN_FILL, ALIGN_RIGHT, DOUBLE_BOTTOM_BORDER, '"R$ "#,##0.00')

    cl = LineChart()
    cl.title, cl.style, cl.width, cl.height = "Curva de Balanço Financeiro (Saldo de Caixa Acumulado - Cenário Realista)", 13, 18, 11
    cl.y_axis.title, cl.x_axis.title = "Saldo Acumulado (R$)", "Cronograma (Meses)"
    cl.add_data(Reference(ws1, min_col=6, min_row=17, max_col=6, max_row=17 + len(fluxo["meses"])), titles_from_data=True)
    cl.set_categories(Reference(ws1, min_col=2, min_row=18, max_row=17 + len(fluxo["meses"])))
    ws1.add_chart(cl, "H5")

    cb = BarChart()
    cb.type, cb.style, cb.width, cb.height = "col", 10, 18, 11
    cb.title, cb.y_axis.title, cb.x_axis.title = "Entradas vs Saídas de Caixa por Mês (Inflows vs Outflows)", "Valor Mensal (R$)", "Mês"
    cb.add_data(Reference(ws1, min_col=3, min_row=17, max_col=4, max_row=17 + len(fluxo["meses"])), titles_from_data=True)
    cb.set_categories(Reference(ws1, min_col=2, min_row=18, max_row=17 + len(fluxo["meses"])))
    ws1.add_chart(cb, "H19")

    # ABA 2: FLUXO ANALÍTICO
    ws2 = wb.create_sheet(title="Fluxo Analítico")
    ws2.views.sheetView[0].showGridLines = True
    tot_col = len(fluxo["meses"]) + 3
    ws2.merge_cells(f"A1:{get_column_letter(tot_col)}1")
    _sc(ws2, 1, 1, "FLUXO DE CAIXA ANALÍTICO MENSAL — ENTRADAS, DESEMBOLSOS E BALANÇO (R$)", FONT_TITLE, NAVY, ALIGN_CENTER)
    ws2.row_dimensions[1].height, ws2.row_dimensions[2].height = 30, 24
    for c_idx, h in enumerate(["Conta / Rubrica Operacional", "Natureza Contábil"] + [f"Mês {m}" for m in fluxo["meses"]] + ["Total Global (R$)"], start=1):
        _sc(ws2, 2, c_idx, h, FONT_HEADER, NAVY, ALIGN_CENTER)

    r_curr = 3
    def _sec(t):
        nonlocal r_curr
        ws2.merge_cells(start_row=r_curr, start_column=1, end_row=r_curr, end_column=tot_col)
        for c in range(1, tot_col + 1): _sc(ws2, r_curr, c, t if c == 1 else None, FONT_BOLD, BLUE_LIGHT, ALIGN_LEFT)
        r_curr += 1

    def _drow(t, n, d, b=False, hi=False):
        nonlocal r_curr
        f, fill = FONT_BOLD if b else FONT_REGULAR, GREEN_FILL if hi else (GRAY_LIGHT if r_curr % 2 == 0 else None)
        _sc(ws2, r_curr, 1, t, f, fill, ALIGN_LEFT); _sc(ws2, r_curr, 2, n, f, fill, ALIGN_CENTER)
        vals = [d.get(m, 0.0) for m in fluxo["meses"]]
        for idx, v in enumerate(vals, start=3): _sc(ws2, r_curr, idx, v, f, fill, ALIGN_RIGHT, num_fmt='"R$ "#,##0.00')
        _sc(ws2, r_curr, tot_col, sum(vals), FONT_BOLD, fill, ALIGN_RIGHT, num_fmt='"R$ "#,##0.00')
        r_curr += 1

    _sec("1. ENTRADAS DE CAIXA (RECEBIMENTOS DE MEDIÇÕES)")
    for t, n, k, _ in RUBRICAS_ENTRADAS: _drow(t, n, fluxo[k])
    _drow("TOTAL DE ENTRADAS DE CAIXA (A)", "Total Inflows", fluxo["total_entradas_realista"], b=True, hi=True)

    _sec("2. SAÍDAS DE CAIXA (DESEMBOLSOS DE PRODUÇÃO E GESTÃO)")
    for t, n, k, _ in RUBRICAS_SAIDAS: _drow(t, n, fluxo[k])
    _drow("TOTAL DE SAÍDAS DE CAIXA (B)", "Total Outflows", fluxo["total_saidas_realista"], b=True)

    _sec("3. RESULTADO OPERACIONAL E SALDOS DE CAIXA")
    _drow("3.1 Saldo Operacional Líquido do Período (A - B)", "Resultado Líquido", fluxo["saldo_periodo_realista"], b=True)
    _sc(ws2, r_curr, 1, "3.2 Saldo de Caixa Acumulado (Cenário Realista)", FONT_BOLD, YELLOW_ACCENT, ALIGN_LEFT, DOUBLE_BOTTOM_BORDER)
    _sc(ws2, r_curr, 2, "Acumulado", FONT_BOLD, YELLOW_ACCENT, ALIGN_CENTER, DOUBLE_BOTTOM_BORDER)
    for idx, m in enumerate(fluxo["meses"], start=3):
        v = fluxo["saldo_acumulado_realista"][m]
        _sc(ws2, r_curr, idx, v, FONT_RED if v < 0 else FONT_GREEN, YELLOW_ACCENT, ALIGN_RIGHT, DOUBLE_BOTTOM_BORDER, '"R$ "#,##0.00')
    _sc(ws2, r_curr, tot_col, s_fim, FONT_BOLD, YELLOW_ACCENT, ALIGN_RIGHT, DOUBLE_BOTTOM_BORDER, '"R$ "#,##0.00')

    # ABA 3: SENSIBILIDADE
    ws3 = wb.create_sheet(title="Cenários de Sensibilidade")
    ws3.views.sheetView[0].showGridLines = True
    ws3.merge_cells("A1:L1")
    _sc(ws3, 1, 1, "ANÁLISE DE SENSIBILIDADE FINANCEIRA — SIMULAÇÃO EM 3 CENÁRIOS", FONT_TITLE, NAVY, ALIGN_CENTER)
    ws3.row_dimensions[1].height = 30
    for c_idx, h in enumerate(["Cenário de Análise", "Premissas Chave de Fluxo"] + [f"M{m}" for m in range(1, prazo_m + 3)] + ["Pico Negativo (R$)", "Saldo Final (R$)"], start=1):
        _sc(ws3, 2, c_idx, h, FONT_HEADER, NAVY, ALIGN_CENTER)

    cenarios = [
        ("Cenário 1: Realista (Base)", "Medições pagas em D+15; Retenção 5% devolvida em M7; Fornecedores 20/80 D0/D+30",
         fluxo["saldo_acumulado_realista"], min(fluxo["saldo_acumulado_realista"].values()), fluxo["saldo_acumulado_realista"][m_fim]),
        ("Cenário 2: Otimista (Adiantamento)", "Adiantamento Contratual de 10% amortizado em 4x nas medições",
         fluxo["saldo_acumulado_otimista"], min(fluxo["saldo_acumulado_otimista"].values()), fluxo["saldo_acumulado_otimista"][m_fim]),
        ("Cenário 3: Estresse (Atraso 30d)", "Contratante atrasa liberação de medições em 30 dias (D+45); fornecedores pagos pontualmente",
         fluxo["saldo_acumulado_pessimista"], min(fluxo["saldo_acumulado_pessimista"].values()), fluxo["saldo_acumulado_pessimista"][prazo_m + 2])
    ]
    for r_idx, (nome, premissa, curva, pico, final) in enumerate(cenarios, start=3):
        fill = GRAY_LIGHT if r_idx % 2 == 0 else None
        _sc(ws3, r_idx, 1, nome, FONT_BOLD, fill, ALIGN_LEFT)
        _sc(ws3, r_idx, 2, premissa, FONT_REGULAR, fill, ALIGN_LEFT)
        for m in range(1, prazo_m + 2):
            val = curva.get(m, 0.0)
            _sc(ws3, r_idx, m + 2, val, FONT_RED if val < 0 else FONT_GREEN, fill, ALIGN_RIGHT, num_fmt='"R$ "#,##0.00')
        vm_end = curva.get(prazo_m + 2, "-")
        f_end = FONT_GREEN if isinstance(vm_end, (int, float)) else FONT_REGULAR
        fmt_end = '"R$ "#,##0.00' if isinstance(vm_end, (int, float)) else None
        _sc(ws3, r_idx, prazo_m + 4, vm_end, f_end, fill, ALIGN_RIGHT, num_fmt=fmt_end)
        _sc(ws3, r_idx, prazo_m + 5, pico, FONT_RED, fill, ALIGN_RIGHT, num_fmt='"R$ "#,##0.00')
        _sc(ws3, r_idx, prazo_m + 6, final, FONT_BOLD, fill, ALIGN_RIGHT, num_fmt='"R$ "#,##0.00')

    for c, w in {2: 15, 3: 18, 4: 18, 5: 18, 6: 22, 7: 10, 8: 10, 9: 10}.items(): ws1.column_dimensions[get_column_letter(c)].width = w
    for c in range(1, tot_col + 1): ws2.column_dimensions[get_column_letter(c)].width = 45 if c == 1 else (22 if c == 2 else 16)
    for c in range(1, prazo_m + 7): ws3.column_dimensions[get_column_letter(c)].width = 25 if c == 1 else (45 if c == 2 else 15)

    wb.save(xlsx_path)
    print(f"-> Excel Executivo de Fluxo de Caixa gerado: {xlsx_path}")


def gerar_relatorio_markdown(fluxo, output_dir, sigla, titulo_obra, area_m2=368.4, prazo_meses=6):
    """Gera o relatório Markdown executivo limpo (sem blocos KaTeX, conforme Manual §5.1)."""
    m_fim = prazo_meses + 1
    pico_r = min(fluxo["saldo_acumulado_realista"].values())
    m_pico = min(fluxo["saldo_acumulado_realista"], key=fluxo["saldo_acumulado_realista"].get)
    s_fim = fluxo["saldo_acumulado_realista"][m_fim]
    m_liq = (s_fim / fluxo["total_venda"]) * 100 if fluxo["total_venda"] > 0 else 0.0
    pico_o, pico_p = min(fluxo["saldo_acumulado_otimista"].values()), min(fluxo["saldo_acumulado_pessimista"].values())

    cols_hdr = [f"Mês {m}" for m in fluxo["meses"]]
    _fr = lambda t, d: f"| {t} | " + " | ".join(formatar_moeda(d.get(m, 0.0)) for m in fluxo["meses"]) + f" | **{formatar_moeda(sum(d.get(m, 0.0) for m in fluxo['meses']))}** |"

    md = [
        f"# 💰 Relatório Executivo de Fluxo de Caixa, Curva de Desembolso e Capital de Giro — {sigla}\n",
        f"**Empreendimento:** {titulo_obra} (`{sigla}`)",
        f"**Área Construída Útil:** {area_m2:.2f} m² | **Prazo Contratual:** {prazo_meses} Meses ({prazo_meses * 30} Dias Corridos)",
        f"**Preço Global Turnkey Contratado:** {formatar_moeda(fluxo['total_venda'])} | **Custo Direto Total:** {formatar_moeda(fluxo['total_custo_direto'])}",
        f"**Fase:** Linha de Base 01 (Baseline 01)\n\n---\n",
        "## 1. Portão de Qualidade de Dados e Nível de Confiança\n",
        "| Parâmetro Financeiro | Fonte / Documento Origem | Nível de Confiança | Observação / Governança |",
        "|---|---|:---:|---|",
        f"| **Cronograma Físico-Financeiro** | `CRONOGRAMA_FISICO_FINANCEIRO_{sigla}.csv` | 🟢 Alto | Itens da EAP distribuídos no cronograma executivo. |",
        "| **Custos Diretos e BDI** | Base Oficial SINAPI / Orçamento Mestre | 🟢 Alto | Códigos e BDI aplicados analiticamente. |",
        "| **Retenção Técnica Contratual** | Cláusula Contratual Turnkey | 🟢 Alto | Retenção padrão de 5% sobre faturamento de medição. |",
        "| **Prazos de Recebimento de Medições** | Contrato de Empreitada Turnkey | 🟢 Alto | Medição no fim do período; pagamento em D+15 após NF. |",
        "| **Prazos Comerciais de Fornecedores** | Padrão Comercial da Construção Civil | 🟢 Alto | Mão de obra no mês; Materiais: 20% à vista e 80% D+30. |",
        "| **Alíquotas Tributárias** | Composição do BDI | 🟢 Alto | 8,65% s/ NF (ISS, PIS, COFINS, CPRB). |\n\n---\n",
        "## 2. Resumo Executivo dos Indicadores Financeiros\n",
        "| Indicador Econômico-Financeiro | Valor Consolidado (R$) | % da Receita Bruta | Impacto / Significado Operacional |",
        "|---|:---:|:---:|---|",
        f"| **Faturamento Bruto da Obra (Turnkey)** | **{formatar_moeda(fluxo['total_venda'])}** | 100,00% | Preço global contratado fechado. |",
        f"| **Retenção Contratual de Garantia (5,0%)** | **{formatar_moeda(fluxo['total_retencao'])}** | 5,00% | Retido nas medições; liberado integralmente no encerramento (TRD). |",
        f"| **Faturamento Líquido de Medições (95%)** | **{formatar_moeda(fluxo['total_venda'] - fluxo['total_retencao'])}** | 95,00% | Volume financeiro disponível durante o transcorrer da obra civil. |",
        f"| **Custo Direto Total da Obra** | **{formatar_moeda(fluxo['total_custo_direto'])}** | {(fluxo['total_custo_direto']/fluxo['total_venda']*100 if fluxo['total_venda']>0 else 0):.2f}% | Custo de execução física + Canteiro/Gestão EAP 1.0. |",
        f"| **Tributos sobre Faturamento (8,65% s/ NF)** | **{formatar_moeda(sum(fluxo['des_impostos'].values()))}** | 8,65% | Recolhimento mensal subsequente à emissão de cada nota fiscal. |",
        f"| **Administração Central, Seguros e Riscos** | **{formatar_moeda(sum(fluxo['des_ac_seguros'].values()))}** | {(sum(fluxo['des_ac_seguros'].values())/fluxo['total_venda']*100 if fluxo['total_venda']>0 else 0):.2f}% | Rateio dos custos indiretos centrais e seguros da construtora. |",
        f"| **LUCRO LÍQUIDO OPERACIONAL REALIZADO** | **{formatar_moeda(s_fim)}** | **{m_liq:.2f}%** | Margem líquida real de lucro após todos os tributos e despesas quitados. |",
        f"| **NECESSIDADE MÍNIMA DE CAPITAL DE GIRO** | **{formatar_moeda(abs(pico_r))}** | **{(abs(pico_r)/fluxo['total_venda']*100 if fluxo['total_venda']>0 else 0):.2f}%** | **Máxima exposição financeira de caixa (atingida no Mês {m_pico}).** |",
        f"| **Ponto de Equilíbrio do Caixa (Break-even)** | **Mês {m_fim}** | — | Ponto em que o caixa torna-se definitivamente superavitário. |\n\n---\n",
        "## 3. Diagnóstico do Capital de Giro e Timing de Desembolso\n",
        f"No cenário realista da obra `{sigla}`, a construtora experimenta sua **máxima exposição financeira no Mês {m_pico}**, atingindo um saldo acumulado negativo de **{formatar_moeda(pico_r)}**.",
        f"A liquidação definitiva do caixa ocorre no Mês {m_fim} com a devolução da retenção técnica de **{formatar_moeda(fluxo['total_retencao'])}**, consolidando um saldo superavitário de **{formatar_moeda(s_fim)}**.\n\n---\n",
        "## 4. Demonstrativo do Fluxo de Caixa Mensal (Cenário Realista Base)\n",
        "| Rubrica Financeira | " + " | ".join(cols_hdr) + " | Total Consolidado (R$) |",
        "|---|" + ":---:|"*len(cols_hdr) + ":---:|",
        "| **1. ENTRADAS DE CAIXA** |" + " |"*len(cols_hdr) + " |",
        *[_fr(lbl.replace("{m_fim}", str(m_fim)), fluxo[k]) for _, _, k, lbl in RUBRICAS_ENTRADAS],
        _fr("**TOTAL ENTRADAS (A)**", fluxo["total_entradas_realista"]),
        "| **2. SAÍDAS DE CAIXA** |" + " |"*len(cols_hdr) + " |",
        *[_fr(lbl, fluxo[k]) for _, _, k, lbl in RUBRICAS_SAIDAS],
        _fr("**TOTAL SAÍDAS (B)**", fluxo["total_saidas_realista"]),
        "| **3. SALDO E RESULTADO** |" + " |"*len(cols_hdr) + " |",
        _fr("**3.1 Saldo Operacional Líquido**", fluxo["saldo_periodo_realista"]),
        f"| **3.2 Saldo de Caixa Acumulado** | " + " | ".join(formatar_moeda(fluxo['saldo_acumulado_realista'][m]) for m in fluxo['meses']) + f" | **{formatar_moeda(s_fim)}** |\n\n---\n",
        "## 5. Análise de Sensibilidade — Comparativo em 3 Cenários\n",
        "| Mês | Cenário Otimista (c/ Adiantamento 10%) | Cenário Realista (Base Contratual) | Cenário Estresse (Atraso Medição 30d) |",
        "|:---:|:---:|:---:|:---:|",
        *[f"| Mês {m} | {formatar_moeda(fluxo['saldo_acumulado_otimista'].get(m, 0.0))} | {formatar_moeda(fluxo['saldo_acumulado_realista'].get(m, 0.0))} | {formatar_moeda(fluxo['saldo_acumulado_pessimista'].get(m, 0.0))} |" for m in range(1, prazo_meses + 2)],
        f"| Mês {prazo_meses + 2} | — | — | {formatar_moeda(fluxo['saldo_acumulado_pessimista'].get(prazo_meses + 2, 0.0))} |",
        f"| **PICO MÁXIMO DE EXPOSIÇÃO** | **{formatar_moeda(pico_o)}** | **{formatar_moeda(pico_r)}** | **{formatar_moeda(pico_p)}** |",
        f"| **Saldo Final Realizado** | **{formatar_moeda(s_fim)}** | **{formatar_moeda(s_fim)}** | **{formatar_moeda(s_fim)}** |\n"
    ]
    salvar_markdown(os.path.join(output_dir, f"RELATORIO_FLUXO_DE_CAIXA_E_CAPITAL_DE_GIRO_{sigla}.md"), "\n".join(md))


def main():
    args = parse_obra_args("Motor Universal de Geração de Fluxo de Caixa, Curva de Desembolso e Capital de Giro.")
    obra_dir = resolver_obra_dir(args)
    config = carregar_config_obra(obra_dir)

    dados_obra = config.get("dados_obra", {})
    sigla = dados_obra.get("sigla", config.get("sigla_obra", os.path.basename(obra_dir).replace("OBRA_", "")))
    titulo = dados_obra.get("nome_obra", config.get("nome_obra", f"Obra {sigla}"))
    area_m2 = float(dados_obra.get("area_construida_m2", config.get("area_construida_m2", 368.4)))
    prazo_meses = int(dados_obra.get("prazo_meses", config.get("prazo_meses", 6)))
    bdi_servico = float(dados_obra.get("bdi_servico_pct", config.get("bdi_servico_pct", 27.17))) / 100.0
    bdi_equip = float(dados_obra.get("bdi_equipamento_pct", config.get("bdi_equipamento_pct", 15.0))) / 100.0

    cronograma_dir = os.path.join(obra_dir, "03_PLANEJAMENTO_E_CRONOGRAMA")
    cronograma_csv = os.path.join(cronograma_dir, f"CRONOGRAMA_FISICO_FINANCEIRO_{sigla}.csv")
    if not os.path.exists(cronograma_csv):
        cand = [os.path.join(cronograma_dir, f) for f in os.listdir(cronograma_dir) if f.startswith("CRONOGRAMA_FISICO_FINANCEIRO") and f.endswith(".csv")] if os.path.exists(cronograma_dir) else []
        if cand: cronograma_csv = cand[0]

    if not os.path.exists(cronograma_csv):
        print(f"[ERRO] Cronograma Físico-Financeiro CSV não encontrado: {cronograma_csv}")
        print("Execute primeiro: python scripts/gerar_cronograma.py --obra <NOME>")
        sys.exit(1)

    output_dir = os.path.join(obra_dir, "05_SUPRIMENTOS_E_FINANCEIRO")
    os.makedirs(output_dir, exist_ok=True)

    print(f"=== MOTOR UNIVERSAL DE FLUXO DE CAIXA: {titulo} ({sigla}) ===")
    vendas_m, diretos_m, total_venda, total_custo_direto = carregar_dados_cronograma(
        cronograma_csv, prazo_meses=prazo_meses, bdi_servico=bdi_servico, bdi_equip=bdi_equip, config=config
    )
    fluxo = modelar_fluxo_caixa(vendas_m, diretos_m, total_venda, total_custo_direto, config, prazo_meses=prazo_meses)

    gerar_csv(fluxo, output_dir, sigla)
    gerar_xlsx(fluxo, output_dir, sigla, titulo)
    gerar_relatorio_markdown(fluxo, output_dir, sigla, titulo, area_m2=area_m2, prazo_meses=prazo_meses)
    print("=== FLUXO DE CAIXA E CAPITAL DE GIRO GERADO COM SUCESSO! ===")


if __name__ == "__main__":
    main()
