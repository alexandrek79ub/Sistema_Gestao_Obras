#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Motor Universal de Geração da Planilha de Medição Evolutiva de Empreiteiros (12 Quinzenas).
Arquitetura Modular v2.0 Lean (Manual de Boas Práticas §2).

Uso:
    python scripts/gerar_planilha_medicao.py --obra OBRA_TMULT
    python scripts/gerar_planilha_medicao.py --dir projetos/OBRA_TMULT
"""

import os
import sys
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from scripts.common.obra_io import parse_obra_args, resolver_obra_dir, carregar_config_obra
from scripts.common.excel_theme import (
    NAVY, BLUE_DARK, BLUE_LIGHT, GREEN_DARK, GREEN_FILL, YELLOW_LIGHT, GRAY_LIGHT,
    FONT_TITLE, FONT_HEADER, FONT_BOLD, FONT_REGULAR, FONT_GREEN, FONT_GOLD,
    THIN_BORDER, DOUBLE_BOTTOM_BORDER, ALIGN_CENTER, ALIGN_RIGHT, ALIGN_LEFT
)

FONT_INFO = Font(name="Calibri", size=10, italic=True, color="FFFFFF")


def _sc(ws, r, c, val=None, font=FONT_REGULAR, fill=None, align=ALIGN_LEFT, border=THIN_BORDER, num_fmt=None):
    cell = ws.cell(row=r, column=c, value=val)
    if font: cell.font = font
    if fill: cell.fill = fill
    if align: cell.alignment = align
    if border: cell.border = border
    if num_fmt: cell.number_format = num_fmt
    return cell


def construir_planilha_medicao_completa(pacotes, output_dir, titulo_obra, sigla, num_medicoes=12):
    """Gera o arquivo Excel completo com as abas de medição quinzenal de cada pacote."""
    caminho = os.path.join(output_dir, "PLANILHA_MEDICAO_QUINZENAL_EMPREITEIROS.xlsx")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    col_med_ini = get_column_letter(7)
    col_med_fim = get_column_letter(6 + num_medicoes)
    col_acum, col_saldo, col_pct = 6 + num_medicoes + 1, 6 + num_medicoes + 2, 6 + num_medicoes + 3
    let_acum, let_saldo, let_pct = get_column_letter(col_acum), get_column_letter(col_saldo), get_column_letter(col_pct)

    for pac in pacotes:
        ws = wb.create_sheet(title=pac.get("tab_name", pac["cod"]))
        ws.views.sheetView[0].showGridLines = True

        # 1. Cabeçalho Principal e Subcabeçalho
        ws.merge_cells(f"A1:{let_pct}1")
        _sc(ws, 1, 1, f"{sigla} — BOLETIM EVOLUTIVO DE MEDIÇÃO QUINZENAL: {pac['cod']} ({pac['titulo'].upper()})", FONT_TITLE, NAVY, ALIGN_CENTER)
        ws.row_dimensions[1].height = 28

        ws.merge_cells(f"A2:{let_pct}2")
        _sc(ws, 2, 1, f"Empreiteiro: {pac.get('empreiteiro', 'Homologado')} | Centro de Custo: {pac.get('cc', '')} | Vigência: {pac.get('vigencia', '')} | Retenção Técnica: 5,0% | Governança: POP 09 e POP 17", FONT_INFO, BLUE_DARK, ALIGN_CENTER)
        ws.row_dimensions[2].height, ws.row_dimensions[4].height, ws.row_dimensions[5].height = 20, 18, 24

        # 2. Cabeçalhos da Tabela (Linhas 4 e 5)
        cols_fixas = [("A", "Código\nEAP"), ("B", "Descrição Pormenorizada do Serviço Contratado"), ("C", "Unid"),
                      ("D", "Quantidade\nContratada"), ("E", "Preço Unitário\n(R$)"), ("F", "Total Contrato\n(R$)")]
        for c_let, tit in cols_fixas:
            ws.merge_cells(f"{c_let}4:{c_let}5")
            _sc(ws, 4, openpyxl.utils.column_index_from_string(c_let), tit, FONT_HEADER, NAVY, Alignment(horizontal="center", vertical="center", wrap_text=True))

        ws.merge_cells(f"{col_med_ini}4:{col_med_fim}4")
        _sc(ws, 4, 7, f"AVANÇO FÍSICO QUINZENAL DE PRODUÇÃO ({num_medicoes} MEDIÇÕES CONTRATUAIS - QTD MEDIDA)", FONT_HEADER, NAVY, ALIGN_CENTER)
        for m in range(1, num_medicoes + 1):
            _sc(ws, 5, 6 + m, f"Medição {m}\n(Qtd)", FONT_HEADER, NAVY, Alignment(horizontal="center", vertical="center", wrap_text=True))

        ws.merge_cells(f"{let_acum}4:{let_pct}4")
        _sc(ws, 4, col_acum, "SALDOS E FECHAMENTO ACUMULADO", FONT_HEADER, GREEN_DARK, ALIGN_CENTER)
        for c_i, tit in [(col_acum, "Total Medido\nAcumulado"), (col_saldo, "Saldo a Medir\n(Quantidade)"), (col_pct, "% Avanço\nFísico")]:
            _sc(ws, 5, c_i, tit, FONT_HEADER, GREEN_DARK, Alignment(horizontal="center", vertical="center", wrap_text=True))

        # 3. Linhas de Dados de Serviços
        servicos = pac.get("servicos", [])
        start_row, end_data_row = 6, 6 + len(servicos) - 1

        for s_idx, s in enumerate(servicos, start=start_row):
            r = s_idx
            ws.row_dimensions[r].height = 20
            _sc(ws, r, 1, s["eap"], FONT_REGULAR, align=ALIGN_CENTER)
            _sc(ws, r, 2, s["desc"], FONT_REGULAR, align=ALIGN_LEFT)
            _sc(ws, r, 3, s["und"], FONT_REGULAR, align=ALIGN_CENTER)
            _sc(ws, r, 4, s["qtd"], FONT_REGULAR, align=ALIGN_RIGHT, num_fmt='#,##0.00')
            _sc(ws, r, 5, s["pu"], FONT_REGULAR, align=ALIGN_RIGHT, num_fmt='"R$ "#,##0.00')
            _sc(ws, r, 6, f"=D{r}*E{r}", FONT_BOLD, align=ALIGN_RIGHT, num_fmt='"R$ "#,##0.00')

            for m in range(1, num_medicoes + 1):
                sim_val = pac.get("simulacao", {}).get(s["eap"], [])[m - 1] if (m - 1) < len(pac.get("simulacao", {}).get(s["eap"], [])) else 0.0
                f_s = FONT_BOLD if sim_val > 0 else FONT_REGULAR
                fill_s = YELLOW_LIGHT if sim_val > 0 else (GRAY_LIGHT if s_idx % 2 == 0 else None)
                _sc(ws, r, 6 + m, sim_val, f_s, fill=fill_s, align=ALIGN_RIGHT, num_fmt='#,##0.00')

            _sc(ws, r, col_acum, f"=SUM({col_med_ini}{r}:{col_med_fim}{r})", FONT_BOLD, align=ALIGN_RIGHT, num_fmt='#,##0.00')
            _sc(ws, r, col_saldo, f"=D{r}-{let_acum}{r}", FONT_REGULAR, align=ALIGN_RIGHT, num_fmt='#,##0.00')
            _sc(ws, r, col_pct, f"={let_acum}{r}/D{r}", FONT_BOLD, align=ALIGN_CENTER, num_fmt='0.00%')

            if s_idx % 2 == 0:
                for ci in [1, 2, 3, 4, 5, 6, col_acum, col_saldo, col_pct]:
                    if ws.cell(row=r, column=ci).fill.start_color.index != 'FEF7E0':
                        ws.cell(row=r, column=ci).fill = GRAY_LIGHT

        # 4. Linhas de Rodapé / Fechamento Financeiro
        r_bruto = end_data_row + 1
        ws.row_dimensions[r_bruto].height = 24
        ws.merge_cells(start_row=r_bruto, start_column=1, end_row=r_bruto, end_column=5)
        _sc(ws, r_bruto, 1, "TOTAL BRUTO MEDIDO NO PERÍODO (R$):", FONT_BOLD, BLUE_LIGHT, ALIGN_RIGHT, DOUBLE_BOTTOM_BORDER)
        _sc(ws, r_bruto, 6, f"=SUM(F{start_row}:F{end_data_row})", FONT_BOLD, BLUE_LIGHT, ALIGN_RIGHT, DOUBLE_BOTTOM_BORDER, '"R$ "#,##0.00')

        for m in range(1, num_medicoes + 1):
            cl = get_column_letter(6 + m)
            _sc(ws, r_bruto, 6 + m, f"=SUMPRODUCT($E${start_row}:$E${end_data_row},{cl}${start_row}:{cl}${end_data_row})",
                FONT_BOLD, BLUE_LIGHT, ALIGN_RIGHT, DOUBLE_BOTTOM_BORDER, '"R$ "#,##0.00')

        _sc(ws, r_bruto, col_acum, f"=SUM({get_column_letter(7)}{r_bruto}:{get_column_letter(6+num_medicoes)}{r_bruto})",
            FONT_BOLD, BLUE_LIGHT, ALIGN_RIGHT, DOUBLE_BOTTOM_BORDER, '"R$ "#,##0.00')
        _sc(ws, r_bruto, col_saldo, f"=F{r_bruto}-{let_acum}{r_bruto}", FONT_BOLD, BLUE_LIGHT, ALIGN_RIGHT, DOUBLE_BOTTOM_BORDER, '"R$ "#,##0.00')
        _sc(ws, r_bruto, col_pct, f"={let_acum}{r_bruto}/F{r_bruto}", FONT_BOLD, BLUE_LIGHT, ALIGN_CENTER, DOUBLE_BOTTOM_BORDER, '0.00%')

        # Demais linhas de rodapé
        footers = [
            (r_bruto + 1, "TOTAL ACUMULADO ATÉ A MEDIÇÃO (R$):", lambda cl: f"=SUM({get_column_letter(7)}${r_bruto}:{cl}${r_bruto})", FONT_BOLD, '"R$ "#,##0.00', None),
            (r_bruto + 2, "SALDO DO CONTRATO APÓS MEDIÇÃO (R$):", lambda cl: f"=$F${r_bruto}-{cl}{r_bruto + 1}", FONT_REGULAR, '"R$ "#,##0.00', None),
            (r_bruto + 3, "% AVANÇO ACUMULADO NO CONTRATO:", lambda cl: f"={cl}{r_bruto + 1}/$F${r_bruto}", FONT_BOLD, '0.00%', None),
            (r_bruto + 4, "(-) RETENÇÃO TÉCNICA DE GARANTIA (5,0%):", lambda cl: f"={cl}{r_bruto}*0.05", FONT_GOLD, '"R$ "#,##0.00', None),
            (r_bruto + 5, "(=) VALOR LÍQUIDO A LIBERAR NA NF-e:", lambda cl: f"={cl}{r_bruto}-{cl}{r_bruto + 4}", FONT_GREEN, '"R$ "#,##0.00', GREEN_FILL),
        ]
        for r_f, label, formula_fn, font_f, fmt_f, fill_f in footers:
            ws.row_dimensions[r_f].height = 24 if "(=)" in label else 20
            ws.merge_cells(start_row=r_f, start_column=1, end_row=r_f, end_column=6)
            _sc(ws, r_f, 1, label, FONT_BOLD, align=ALIGN_RIGHT)
            for m in range(1, num_medicoes + 1):
                cl = get_column_letter(6 + m)
                _sc(ws, r_f, 6 + m, formula_fn(cl), font_f, fill=fill_f, align=ALIGN_CENTER if fmt_f == '0.00%' else ALIGN_RIGHT, num_fmt=fmt_f)
            for ci in range(1, col_pct + 1):
                ws.cell(row=r_f, column=ci).border = THIN_BORDER

        # 5. Geometria de Colunas e Congelamento
        ws.freeze_panes = "G6"
        dim_map = {"A": 11, "B": 46, "C": 8, "D": 16, "E": 16, "F": 18, let_acum: 18, let_saldo: 16, let_pct: 14}
        for c_l, w in dim_map.items(): ws.column_dimensions[c_l].width = w
        for m in range(1, num_medicoes + 1): ws.column_dimensions[get_column_letter(6 + m)].width = 15

    wb.save(caminho)
    print(f"-> Planilha Master de Medição gerada: {caminho}")


def main():
    args = parse_obra_args("Motor Universal de Medição de Empreiteiros.")
    obra_dir = resolver_obra_dir(args)
    config = carregar_config_obra(obra_dir)

    dados_obra = config.get("dados_obra", {})
    sigla = dados_obra.get("sigla", config.get("sigla_obra", os.path.basename(obra_dir).replace("OBRA_", "")))
    titulo = dados_obra.get("nome_obra", config.get("nome_obra", f"Obra {sigla}"))
    prazo_meses = int(dados_obra.get("prazo_meses", config.get("prazo_meses", 6)))
    num_medicoes = prazo_meses * 2

    med_json = os.path.join(obra_dir, "02_ORCAMENTO_BASE_E_CONTRATOS", "dados_medicoes_empreiteiros.json")
    if not os.path.exists(med_json):
        print(f"[ERRO] Arquivo de medições não encontrado: {med_json}")
        sys.exit(1)

    with open(med_json, "r", encoding="utf-8") as f:
        pacotes = json.load(f)

    output_dir = os.path.join(obra_dir, "02_ORCAMENTO_BASE_E_CONTRATOS", "CONTRATOS_EMPREITEIROS")
    os.makedirs(output_dir, exist_ok=True)

    print(f"=== MOTOR UNIVERSAL DE PLANILHA DE MEDIÇÃO: {titulo} ({sigla}) ===")
    construir_planilha_medicao_completa(pacotes, output_dir, titulo, sigla, num_medicoes=num_medicoes)
    print("=== PLANILHA DE MEDIÇÃO GERADA COM SUCESSO! ===")


if __name__ == "__main__":
    main()
