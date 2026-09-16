from typing import Any
from openpyxl.styles import Alignment, Font, PatternFill
from motor_quantitativos.exportadores.excel.estilos import (
    COR_CABECALHO, COR_TOPO, COR_ZEBRADO,
    ajustar_larguras, borda_padrao
)


def adicionar_aba_memoria(wb, nome_obra: str, itens: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title="05_Memória_Quantitativos")
    ws.freeze_panes = "A4"

    ws.merge_cells("A1:H1")
    ws["A1"].value = f"MEMÓRIA DE CÁLCULO E LEVANTAMENTO GEOMÉTRICO — {nome_obra.upper()}"
    ws["A1"].font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color=COR_TOPO, end_color=COR_TOPO, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    cabecalhos = [
        ("Item EAP", "center"), ("Descrição do Serviço", "left"), ("Disciplina", "left"),
        ("Unid", "center"), ("Equação Matemática Literal", "left"), ("Qtd Líquida Calculada", "right"),
        ("Prancha Referência", "center"), ("Status / RFI", "center"),
    ]

    for col_idx, (nome_col, alin) in enumerate(cabecalhos, start=1):
        cell = ws.cell(row=3, column=col_idx, value=nome_col)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type="solid")
        cell.alignment = Alignment(horizontal=alin, vertical="center")
        cell.border = borda_padrao()
    ws.row_dimensions[3].height = 24

    for idx, item in enumerate(itens):
        r_num = 4 + idx
        ws.row_dimensions[r_num].height = 20
        qtd = float(item.get("quantidade_liquida") or 0)

        valores = [
            (item.get("cod_eap", ""), "@", "center"),
            (item.get("descricao", ""), "@", "left"),
            (item.get("disciplina", ""), "@", "left"),
            (item.get("unidade", ""), "@", "center"),
            (item.get("expressao_matematica", ""), "@", "left"),
            (qtd, "#,##0.0000" if qtd % 1 != 0 else "#,##0", "right"),
            (item.get("prancha_referencia", ""), "@", "center"),
            (item.get("status", ""), "@", "center"),
        ]

        fill_cor = PatternFill(start_color=COR_ZEBRADO, end_color=COR_ZEBRADO, fill_type="solid") if idx % 2 == 1 else None

        for col_idx, (val, fmt, alin) in enumerate(valores, start=1):
            cell = ws.cell(row=r_num, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal=alin, vertical="center")
            cell.border = borda_padrao()
            if fill_cor:
                cell.fill = fill_cor

    ajustar_larguras(ws)
