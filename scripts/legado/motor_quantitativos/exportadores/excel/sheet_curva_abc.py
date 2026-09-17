from typing import Any
from openpyxl.styles import Alignment, Font, PatternFill
from motor_quantitativos.exportadores.excel.estilos import (
    COR_CABECALHO, COR_TOPO,
    ajustar_larguras, borda_padrao
)


def adicionar_aba_curva_abc(wb, nome_obra: str, itens: list[dict[str, Any]], linha_total_orcamento: int) -> None:
    ws = wb.create_sheet(title="04_Curva_ABC_Serviços")
    ws.freeze_panes = "A5"

    ws.merge_cells("A1:I1")
    ws["A1"].value = f"CURVA ABC DE SERVIÇOS (MATRIZ PARETO 80/20) — {nome_obra.upper()}"
    ws["A1"].font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color=COR_TOPO, end_color=COR_TOPO, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells("A2:I2")
    ws["A2"].value = "Classe A (até 80% do valor acumulado) | Classe B (80% a 95%) | Classe C (95% a 100%)"
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="64748B")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    cabecalhos = [
        ("Ranking", "center"), ("Item EAP", "center"), ("Descrição do Serviço", "left"),
        ("Disciplina", "left"), ("Unid", "center"), ("Custo Total (R$)", "right"),
        ("% Participação", "right"), ("% Acumulado", "right"), ("Classe ABC", "center"),
    ]

    for col_idx, (nome_col, alin) in enumerate(cabecalhos, start=1):
        cell = ws.cell(row=4, column=col_idx, value=nome_col)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type="solid")
        cell.alignment = Alignment(horizontal=alin, vertical="center")
        cell.border = borda_padrao()
    ws.row_dimensions[4].height = 24

    itens_ordenados = sorted(itens, key=lambda x: float(x.get("custo_total") or 0), reverse=True)

    for idx, item in enumerate(itens_ordenados):
        r_num = 5 + idx
        ws.row_dimensions[r_num].height = 20
        total_item = float(item.get("custo_total") or 0)

        formula_part = f"=F{r_num}/'01_Orçamento_Base'!$I${linha_total_orcamento}"
        formula_acum = f"=G{r_num}" if idx == 0 else f"=H{r_num-1}+G{r_num}"
        formula_classe = f'=IF(H{r_num}<=0.8001, "A", IF(H{r_num}<=0.9501, "B", "C"))'

        valores = [
            (idx + 1, "0", "center"),
            (item.get("cod_eap", ""), "@", "center"),
            (item.get("descricao", ""), "@", "left"),
            (item.get("disciplina", ""), "@", "left"),
            (item.get("unidade", ""), "@", "center"),
            (total_item, "R$ #,##0.00", "right"),
            (formula_part, "0.00%", "right"),
            (formula_acum, "0.00%", "right"),
            (formula_classe, "@", "center"),
        ]

        for col_idx, (val, fmt, alin) in enumerate(valores, start=1):
            cell = ws.cell(row=r_num, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal=alin, vertical="center")
            cell.border = borda_padrao()
            if col_idx == 9:
                cell.font = Font(name="Calibri", size=11, bold=True)

    ajustar_larguras(ws)
