from typing import Any
from openpyxl.styles import Alignment, Font, PatternFill
from motor_quantitativos.exportadores.excel.estilos import (
    COR_CABECALHO, COR_TOPO, COR_TOTAL, COR_ZEBRADO,
    ajustar_larguras, borda_padrao, borda_total
)


def adicionar_aba_medicao(wb, nome_obra: str, itens: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title="06_Boletim_Medição")
    ws.freeze_panes = "A5"

    ws.merge_cells("A1:L1")
    ws["A1"].value = f"BOLETIM DE MEDIÇÃO FÍSICO-FINANCEIRO — {nome_obra.upper()}"
    ws["A1"].font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color=COR_TOPO, end_color=COR_TOPO, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells("A2:L2")
    ws["A2"].value = "Preenchimento periódico de obra | Fórmulas automáticas de avanço %, saldo e valor liberado"
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="64748B")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    cabecalhos = [
        ("Item EAP", "center"), ("Descrição do Serviço", "left"), ("Unid", "center"),
        ("PU Contratado (R$)", "right"), ("Qtd Contratada", "right"), ("Valor Contratado (R$)", "right"),
        ("Qtd Medida Anterior", "right"), ("Qtd no Período", "right"), ("Qtd Acumulada", "right"),
        ("Saldo a Medir (Qtd)", "right"), ("% Avanço Físico", "right"), ("Valor Medido Período (R$)", "right"),
    ]

    for col_idx, (nome_col, alin) in enumerate(cabecalhos, start=1):
        cell = ws.cell(row=4, column=col_idx, value=nome_col)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type="solid")
        cell.alignment = Alignment(horizontal=alin, vertical="center")
        cell.border = borda_padrao()
    ws.row_dimensions[4].height = 24

    for idx, item in enumerate(itens):
        r_num = 5 + idx
        ws.row_dimensions[r_num].height = 20
        r_orc = 5 + idx

        f_pu = f"='01_Orçamento_Base'!H{r_orc}"
        f_qtd = f"='01_Orçamento_Base'!E{r_orc}"
        f_val_contratado = f"=ROUND(D{r_num}*E{r_num}, 2)"
        f_qtd_acum = f"=G{r_num}+H{r_num}"
        f_saldo_qtd = f"=E{r_num}-I{r_num}"
        f_avanco = f"=IF(E{r_num}>0, ROUND(I{r_num}/E{r_num}, 4), 0)"
        f_val_periodo = f"=ROUND(H{r_num}*D{r_num}, 2)"

        valores = [
            (item.get("cod_eap", ""), "@", "center"),
            (item.get("descricao", ""), "@", "left"),
            (item.get("unidade", ""), "@", "center"),
            (f_pu, "R$ #,##0.00", "right"),
            (f_qtd, "#,##0.00", "right"),
            (f_val_contratado, "R$ #,##0.00", "right"),
            (0.0, "#,##0.00", "right"),
            (0.0, "#,##0.00", "right"),
            (f_qtd_acum, "#,##0.00", "right"),
            (f_saldo_qtd, "#,##0.00", "right"),
            (f_avanco, "0.00%", "right"),
            (f_val_periodo, "R$ #,##0.00", "right"),
        ]

        fill_cor = PatternFill(start_color=COR_ZEBRADO, end_color=COR_ZEBRADO, fill_type="solid") if idx % 2 == 1 else None

        for col_idx, (val, fmt, alin) in enumerate(valores, start=1):
            cell = ws.cell(row=r_num, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal=alin, vertical="center")
            cell.border = borda_padrao()
            if fill_cor:
                cell.fill = fill_cor

    r_tot = 5 + len(itens)
    ws.row_dimensions[r_tot].height = 24
    ws.merge_cells(f"A{r_tot}:E{r_tot}")
    ws[f"A{r_tot}"].value = "TOTAIS DA MEDIÇÃO (R$):"
    ws[f"A{r_tot}"].font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    ws[f"A{r_tot}"].alignment = Alignment(horizontal="right", vertical="center")
    ws[f"A{r_tot}"].fill = PatternFill(start_color=COR_TOTAL, end_color=COR_TOTAL, fill_type="solid")

    for col in range(1, 6):
        ws.cell(row=r_tot, column=col).border = borda_total()

    ws[f"F{r_tot}"].value = f"=SUM(F5:F{r_tot-1})" if len(itens) > 0 else 0
    ws[f"F{r_tot}"].font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    ws[f"F{r_tot}"].number_format = "R$ #,##0.00"
    ws[f"F{r_tot}"].alignment = Alignment(horizontal="right", vertical="center")
    ws[f"F{r_tot}"].fill = PatternFill(start_color=COR_TOTAL, end_color=COR_TOTAL, fill_type="solid")
    ws[f"F{r_tot}"].border = borda_total()

    for col in range(7, 12):
        c = ws.cell(row=r_tot, column=col)
        c.fill = PatternFill(start_color=COR_TOTAL, end_color=COR_TOTAL, fill_type="solid")
        c.border = borda_total()

    ws[f"L{r_tot}"].value = f"=SUM(L5:L{r_tot-1})" if len(itens) > 0 else 0
    ws[f"L{r_tot}"].font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    ws[f"L{r_tot}"].number_format = "R$ #,##0.00"
    ws[f"L{r_tot}"].alignment = Alignment(horizontal="right", vertical="center")
    ws[f"L{r_tot}"].fill = PatternFill(start_color=COR_TOTAL, end_color=COR_TOTAL, fill_type="solid")
    ws[f"L{r_tot}"].border = borda_total()

    ajustar_larguras(ws)
