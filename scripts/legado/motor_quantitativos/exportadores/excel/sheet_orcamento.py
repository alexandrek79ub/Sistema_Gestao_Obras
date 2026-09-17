from typing import Any
from openpyxl.styles import Alignment, Font, PatternFill
from motor_quantitativos.exportadores.excel.estilos import (
    COR_CABECALHO, COR_TOPO, COR_TOTAL, COR_ZEBRADO,
    ajustar_larguras, borda_padrao, borda_total
)


def adicionar_aba_orcamento(wb, nome_obra: str, checksum: str, data_emissao: str, itens: list[dict[str, Any]]) -> int:
    ws = wb.active if wb.sheetnames == ["Sheet"] else wb.create_sheet(title="01_Orçamento_Base")
    ws.title = "01_Orçamento_Base"
    ws.freeze_panes = "A5"

    # Cabeçalho do documento
    ws.merge_cells("A1:N1")
    ws["A1"].value = f"PLANILHA ORÇAMENTÁRIA EXECUTIVA — {nome_obra.upper()}"
    ws["A1"].font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color=COR_TOPO, end_color=COR_TOPO, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:N2")
    ws["A2"].value = f"Emissão: {data_emissao} | Fonte: SQLite SSOT Oficial | Checksum SHA-256: {checksum[:16]}..."
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="64748B")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 8

    colunas = [
        ("Item EAP", "center"), ("Descrição do Serviço", "left"), ("Disciplina", "left"),
        ("Unid", "center"), ("Qtd Líquida", "right"), ("Custo Direto Unit. (R$)", "right"),
        ("BDI (%)", "right"), ("Preço Unit. c/ BDI (R$)", "right"), ("Custo Total (R$)", "right"),
        ("Prancha Referência", "center"), ("Fonte Preço", "left"), ("Status", "center"),
        ("Cód. SINAPI", "center"), ("Centro de Custo", "left"),
    ]

    for col_idx, (nome_col, alin) in enumerate(colunas, start=1):
        cell = ws.cell(row=4, column=col_idx, value=nome_col)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type="solid")
        cell.alignment = Alignment(horizontal=alin, vertical="center", wrap_text=True)
        cell.border = borda_padrao()
    ws.row_dimensions[4].height = 26

    linha_inicial = 5
    for idx, item in enumerate(itens):
        r_num = linha_inicial + idx
        ws.row_dimensions[r_num].height = 20
        
        qtd = float(item.get("quantidade_liquida") or 0)
        pu = float(item.get("preco_unitario") or 0)
        bdi = float(item.get("bdi_pct") or 0)
        
        valores = [
            (item.get("cod_eap", ""), "@", "center"),
            (item.get("descricao", ""), "@", "left"),
            (item.get("disciplina", ""), "@", "left"),
            (item.get("unidade", ""), "@", "center"),
            (qtd, "#,##0.0000" if qtd % 1 != 0 else "#,##0", "right"),
            (pu, "R$ #,##0.00", "right"),
            (bdi, "0.00", "right"),
            (f"=ROUND(F{r_num}*(1+G{r_num}/100), 2)", "R$ #,##0.00", "right"),
            (f"=ROUND(E{r_num}*H{r_num}, 2)", "R$ #,##0.00", "right"),
            (item.get("prancha_referencia", ""), "@", "center"),
            (item.get("fonte_preco", ""), "@", "left"),
            (item.get("status", ""), "@", "center"),
            (item.get("codigo_sinapi", ""), "@", "center"),
            (item.get("centro_custo", ""), "@", "left"),
        ]

        fill_cor = PatternFill(start_color=COR_ZEBRADO, end_color=COR_ZEBRADO, fill_type="solid") if idx % 2 == 1 else None

        for col_idx, (val, fmt, alin) in enumerate(valores, start=1):
            cell = ws.cell(row=r_num, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal=alin, vertical="center")
            cell.border = borda_padrao()
            if fill_cor:
                cell.fill = fill_cor

    linha_final_dados = linha_inicial + len(itens) - 1
    linha_total = linha_final_dados + 1

    ws.row_dimensions[linha_total].height = 24
    ws.merge_cells(f"A{linha_total}:H{linha_total}")
    lbl = ws[f"A{linha_total}"]
    lbl.value = "TOTAL GERAL ORÇADO (R$):"
    lbl.font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    lbl.alignment = Alignment(horizontal="right", vertical="center")
    lbl.fill = PatternFill(start_color=COR_TOTAL, end_color=COR_TOTAL, fill_type="solid")

    for col in range(1, 9):
        ws.cell(row=linha_total, column=col).border = borda_total()

    val_tot = ws[f"I{linha_total}"]
    val_tot.value = f"=SUM(I{linha_inicial}:I{linha_final_dados})" if len(itens) > 0 else 0
    val_tot.font = Font(name="Calibri", size=11, bold=True, color="0F172A")
    val_tot.number_format = "R$ #,##0.00"
    val_tot.alignment = Alignment(horizontal="right", vertical="center")
    val_tot.fill = PatternFill(start_color=COR_TOTAL, end_color=COR_TOTAL, fill_type="solid")
    val_tot.border = borda_total()

    for col in range(10, 15):
        c = ws.cell(row=linha_total, column=col)
        c.fill = PatternFill(start_color=COR_TOTAL, end_color=COR_TOTAL, fill_type="solid")
        c.border = borda_total()

    ajustar_larguras(ws)
    return linha_total
